"""Local FlowLab Pro browser MVP.

Run with: .venv\\Scripts\\python.exe web_app.py
Then open: http://127.0.0.1:5000
"""

import csv
import calendar
from datetime import date, timedelta
import io
import json
import hashlib
import os
from pathlib import Path
import secrets
import shutil
import smtplib
import time
from email.message import EmailMessage
from email.utils import parseaddr
from uuid import uuid4

from flask import Flask, abort, flash, jsonify, redirect, render_template, request, send_file, session, url_for

from app.database.database import get_connection
from app.database.init_db import initialize_database
from app.modules.uncertainty.engine import Budget2Resolver, RSSEngine, TypeAProcessor, UncertaintyInput
from app.modules.uncertainty.matrix import MatrixResolver
from app.modules.uncertainty.repository import UncertaintyRepository
from app.modules.uncertainty.web_service import (
    SOURCE_TEMPLATES, calculate_payload, create_cmc_revision, create_revision, equipment_dict,
    equipment_rows, pdf_report, record_payload, save_calculation, save_cmc, save_partial_draft, workflow,
)
from app.modules.auth.permissions import Permissions
from app.modules.auth.service import AuthService


app = Flask(__name__)
app.config["SECRET_KEY"] = "flowlab-local-mvp"
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024
REPORT_MAX_BYTES = 10 * 1024 * 1024
REPORT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".jpg", ".jpeg", ".png"}
PERSONNEL_DOCUMENT_TYPES = (
    "Laboratory Role Authorization",
    "Impartiality Assessment",
    "Confidentiality Policy Acknowledgement",
    "Job Description",
)


@app.context_processor
def permission_context():
    return {"can": lambda permission: app.config.get("TESTING") or
        Permissions.can(session.get("role_name", ""), permission)}


@app.context_processor
def notification_context():
    user_id = session.get("user_id")
    if not user_id:
        return {"notification_count": 0}
    connection = get_connection()
    try:
        if workflow_authority():
            count = connection.execute("""SELECT COUNT(*) FROM workflow_tasks
                WHERE status='Pending'""").fetchone()[0]
        else:
            count = connection.execute("""SELECT COUNT(*) FROM notifications
                WHERE user_id=? AND is_read=0""", (user_id,)).fetchone()[0]
        return {"notification_count": count}
    finally:
        connection.close()


def query(sql, params=()):
    connection = get_connection()
    try:
        return connection.execute(sql, params).fetchall()
    finally:
        connection.close()


def current_actor():
    return session.get("full_name") or session.get("username") or "Unauthenticated local session"


def email_recovery_configured():
    return all(os.environ.get(name, "").strip() for name in
        ("FLOWLAB_PUBLIC_URL", "FLOWLAB_SMTP_HOST", "FLOWLAB_SMTP_FROM"))


def send_password_reset_email(recipient, reset_url):
    message = EmailMessage()
    message["Subject"] = "FlowLab Pro password recovery"
    message["From"] = os.environ["FLOWLAB_SMTP_FROM"]
    message["To"] = recipient
    message.set_content(
        "A password reset was requested for your FlowLab Pro account.\n\n"
        f"Open this link within 30 minutes:\n{reset_url}\n\n"
        "If you did not request this reset, ignore this message. The link can be used only once."
    )
    host = os.environ["FLOWLAB_SMTP_HOST"]
    port = int(os.environ.get("FLOWLAB_SMTP_PORT", "587"))
    use_tls = os.environ.get("FLOWLAB_SMTP_STARTTLS", "1").lower() not in {"0", "false", "no"}
    username = os.environ.get("FLOWLAB_SMTP_USERNAME", "").strip()
    password = os.environ.get("FLOWLAB_SMTP_PASSWORD", "")
    with smtplib.SMTP(host, port, timeout=15) as server:
        if use_tls:
            server.starttls()
        if username:
            server.login(username, password)
        server.send_message(message)


def create_officer_reset_request(connection, user):
    pending = connection.execute("""SELECT id FROM password_reset_requests
        WHERE user_id=? AND status='Pending'""", (user[0],)).fetchone()
    if pending:
        return pending[0]
    request_id = connection.execute("""INSERT INTO password_reset_requests
        (user_id,status) VALUES (?,'Pending')""", (user[0],)).lastrowid
    officers = connection.execute("""SELECT u.id FROM users u JOIN roles r ON r.id=u.role_id
        WHERE u.is_active=1 AND r.name IN ('Chief Meteorologist','Administrator')""").fetchall()
    for officer in officers:
        add_notification(connection, officer[0], "Password reset requested",
            f"Account {user[1]} requested a password reset.", url_for("users"),
            "password_reset_request", request_id, "Security")
    audit_change(connection, "password_reset_request", request_id, "request",
        after={"username": user[1], "status": "Pending"})
    return request_id


def chief_auto_approval():
    return session.get("role_name") == "Chief Meteorologist"


def workflow_authority():
    return session.get("role_name") in {"Chief Meteorologist", "Administrator"}


def task_destination(connection, entity_type, entity_id, stored_link=None):
    if stored_link and stored_link.startswith("/") and not stored_link.startswith("//"):
        return stored_link
    if entity_type == "password_reset_request":
        return url_for("users")
    if entity_type in {"calibration_job", "job"}:
        return url_for("job_detail", job_id=entity_id)
    if entity_type == "project":
        return url_for("project_jobs", project=entity_id)
    if entity_type == "equipment":
        return url_for("equipment_detail", equipment_id=entity_id)
    if entity_type in {"calibration", "maintenance"}:
        table = "calibration_history" if entity_type == "calibration" else "maintenance_history"
        owner = connection.execute(f"SELECT equipment_id FROM {table} WHERE id=?", (entity_id,)).fetchone()
        if owner:
            return url_for("equipment_detail", equipment_id=owner[0], tab=entity_type)
    if entity_type in {"uncertainty_calculation", "uncertainty"}:
        return url_for("uncertainty", record=entity_id)
    if entity_type == "cmc_revision":
        return url_for("uncertainty")
    if entity_type == "capa":
        return url_for("capa_detail",record_id=entity_id)
    if entity_type in {"iso_clause_assessment", "iso_clause_evidence"}:
        if entity_type == "iso_clause_evidence":
            clause = connection.execute(
                "SELECT clause_id FROM iso_clause_evidence WHERE id=?", (entity_id,)).fetchone()
            entity_id = clause[0] if clause else entity_id
        return url_for("iso_clause_detail", clause_id=entity_id)
    if entity_type == "personnel_registration":
        return url_for("users", registration_id=entity_id)
    if entity_type == "personnel_document":
        owner = connection.execute("SELECT user_id FROM personnel_documents WHERE id=?", (entity_id,)).fetchone()
        return url_for("personnel_record", user_id=owner[0]) if owner else url_for("iso_clause_checklist")
    if entity_type == "controlled_document":
        return url_for("documents")
    if entity_type == "quality_record":
        record = connection.execute("SELECT record_type FROM quality_records WHERE id=?", (entity_id,)).fetchone()
        if record:
            for section, workspace_data in WORKSPACES.items():
                if any(item[0] == record[0] for item in workspace_data["modules"]):
                    return url_for("module_page", section=section, module=record[0])
    return url_for("reminders")


def add_notification(connection, user_id, title, message, link=None,
        entity_type=None, entity_id=None, notification_type="Workflow"):
    if not user_id:
        return
    connection.execute("""INSERT INTO notifications
        (user_id,notification_type,title,message,link,entity_type,entity_id,created_by)
        VALUES (?,?,?,?,?,?,?,?)""", (user_id, notification_type, title, message, link,
        entity_type, entity_id, current_actor()))


def user_id_for_actor(connection, actor):
    row = connection.execute("""SELECT id FROM users
        WHERE full_name=? OR username=? ORDER BY CASE WHEN full_name=? THEN 0 ELSE 1 END LIMIT 1""",
        (actor, actor, actor)).fetchone()
    return row[0] if row else None


def ensure_unique_document_number(connection, document_number, exclude_table=None, exclude_id=None):
    number = (document_number or "").strip()
    if not number: return
    sources = (("controlled_documents","Controlled document"),
        ("iso_clause_evidence","Clause evidence"),("personnel_documents","Personnel record"))
    for table,label in sources:
        sql=f"SELECT id FROM {table} WHERE LOWER(TRIM(document_number))=LOWER(TRIM(?))"
        params=[number]
        if table==exclude_table and exclude_id:
            sql+=" AND id<>?";params.append(exclude_id)
        if connection.execute(sql,tuple(params)).fetchone():
            raise ValueError(f"Document number {number} is already registered. Document numbers remain reserved after supersession or deletion.")


def active_reviewers(connection):
    return connection.execute("""SELECT u.id,u.full_name,r.name FROM users u
        JOIN roles r ON r.id=u.role_id WHERE u.is_active=1
        AND r.name IN ('Chief Meteorologist','Supervisor','Engineer','Administrator','HOD')
        AND u.id<>? ORDER BY CASE r.name WHEN 'Chief Meteorologist' THEN 1 WHEN 'Supervisor' THEN 2
        WHEN 'Engineer' THEN 3 ELSE 4 END,u.full_name""", (session.get("user_id") or -1,)).fetchall()


def require_permission(permission):
    if app.config.get("TESTING"):
        return
    if not Permissions.can(session.get("role_name", ""), permission):
        abort(403)


def setting_integer(key, fallback):
    rows = query("SELECT setting_value FROM system_settings WHERE setting_key=?", (key,))
    try:
        return max(0, int(rows[0][0])) if rows else fallback
    except (TypeError, ValueError):
        return fallback


def setting_text(key, fallback=""):
    rows = query("SELECT setting_value FROM system_settings WHERE setting_key=?", (key,))
    return str(rows[0][0]).strip() if rows and rows[0][0] is not None else fallback


def valid_email(value):
    parsed = parseaddr(value or "")[1]
    return parsed == (value or "").strip() and "@" in parsed and "." in parsed.rsplit("@", 1)[-1]


def next_reference(connection, table, column, prefix):
    year = date.today().year
    stem = f"{prefix}-{year}-"
    rows = connection.execute(f"SELECT {column} FROM {table} WHERE {column} LIKE ? ORDER BY {column} DESC LIMIT 1",
        (stem + "%",)).fetchone()
    try:
        number = int(rows[0].rsplit("-", 1)[1]) + 1 if rows else 1
    except (TypeError, ValueError, IndexError):
        number = 1
    candidate = f"{stem}{number:03d}"
    while connection.execute(f"SELECT 1 FROM {table} WHERE {column}=?", (candidate,)).fetchone():
        number += 1; candidate = f"{stem}{number:03d}"
    return candidate


def add_validity(start_date, value, unit):
    start = date.fromisoformat(start_date)
    amount = int(value)
    if amount <= 0 or unit not in {"Days", "Months", "Years"}:
        raise ValueError("Validity must be a positive duration in Days, Months, or Years.")
    if unit == "Days":
        from datetime import timedelta
        return (start + timedelta(days=amount)).isoformat()
    months = amount * (12 if unit == "Years" else 1)
    month_index = start.month - 1 + months
    year = start.year + month_index // 12
    month = month_index % 12 + 1
    day = min(start.day, calendar.monthrange(year, month)[1])
    return date(year, month, day).isoformat()


def controlled_status():
    role = session.get("role_name", "")
    return "Approved" if Permissions.can(role, "record_approve") or (app.config.get("TESTING") and not role) else "Draft"


def add_note(connection, entity_type, entity_id, note, visibility="Internal"):
    note = (note or "").strip()
    if note:
        connection.execute("""INSERT INTO record_notes
            (entity_type,entity_id,note,author,visibility) VALUES (?,?,?,?,?)""",
            (entity_type, entity_id, note, current_actor(), visibility))


def audit_change(connection, entity_type, entity_id, action, before=None, after=None):
    connection.execute("""INSERT INTO audit_trail
        (entity_type,entity_id,action,actor,before_state,after_state)
        VALUES (?,?,?,?,?,?)""", (
        entity_type, entity_id, action, current_actor(),
        json.dumps(before, default=str, sort_keys=True) if before is not None else None,
        json.dumps(after, default=str, sort_keys=True) if after is not None else None))


def recycle_record(connection, entity_type, entity_id, display_name, original_location=None):
    connection.execute("""INSERT INTO recycle_bin
        (entity_type,entity_id,display_name,original_location,deleted_by)
        VALUES (?,?,?,?,?) ON CONFLICT(entity_type,entity_id) DO UPDATE SET
        display_name=excluded.display_name,original_location=excluded.original_location,
        deleted_by=excluded.deleted_by,deleted_at=CURRENT_TIMESTAMP""",
        (entity_type, entity_id, display_name, original_location, current_actor()))


def validate_report(upload, required=False):
    if not upload or not upload.filename:
        if required:
            raise ValueError("A report file is required.")
        return
    extension = Path(upload.filename).suffix.lower()
    if extension not in REPORT_EXTENSIONS:
        raise ValueError("Report file type is not allowed.")
    upload.stream.seek(0, 2)
    size = upload.stream.tell()
    upload.stream.seek(0)
    if size > REPORT_MAX_BYTES:
        raise ValueError("Report files must not exceed 10 MB.")


@app.before_request
def ensure_database():
    initialize_database()
    if app.config.get("TESTING") or request.endpoint in {
            "login", "setup", "forgot_password", "email_password_reset", "static"}:
        return
    if not session.get("user_id"):
        return redirect(url_for("login", next=request.path))
    if session.get("must_change_password") and request.endpoint not in {"change_password", "logout"}:
        return redirect(url_for("change_password"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        user = AuthService().authenticate(
            request.form.get("username", "").strip(), request.form.get("password", ""))
        if user:
            session.clear()
            session.update(user_id=user.id, username=user.username, full_name=user.full_name,
                role_name=user.role_name, must_change_password=user.must_change_password)
            if user.must_change_password:
                return redirect(url_for("change_password"))
            return redirect(request.args.get("next") or url_for("dashboard"))
        flash("Username or password is invalid.", "error")
    needs_setup = query("SELECT COUNT(*) FROM users")[0][0] == 0
    return render_template("login.html", page_title="Sign In", needs_setup=needs_setup)


@app.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():
    if request.method == "POST":
        identifier = request.form.get("identifier", request.form.get("username", "")).strip()
        connection = get_connection()
        user = connection.execute("""SELECT u.id,u.username,u.email,r.name FROM users u
            JOIN roles r ON r.id=u.role_id
            WHERE (LOWER(u.username)=LOWER(?) OR LOWER(COALESCE(u.email,''))=LOWER(?))
            AND u.is_active=1""", (identifier,identifier)).fetchone()
        admin_recovery_email = setting_text("admin_recovery_email", "ikechukwuumezulike@gmail.com")
        if not user and identifier.lower() == admin_recovery_email.lower():
            user = connection.execute("""SELECT u.id,u.username,u.email,r.name FROM users u
                JOIN roles r ON r.id=u.role_id WHERE u.is_active=1
                AND r.name IN ('Administrator','Chief Meteorologist')
                ORDER BY CASE r.name WHEN 'Administrator' THEN 1 ELSE 2 END,u.id LIMIT 1""").fetchone()
        if user:
            delivered = False
            recipient = admin_recovery_email if user[3] in {"Administrator","Chief Meteorologist"} else user[2]
            if recipient and email_recovery_configured():
                token = secrets.token_urlsafe(32)
                token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
                connection.execute("""UPDATE password_reset_tokens SET used_at=CURRENT_TIMESTAMP
                    WHERE user_id=? AND used_at IS NULL""", (user[0],))
                connection.execute("""INSERT INTO password_reset_tokens
                    (user_id,token_hash,expires_at_epoch) VALUES (?,?,?)""",
                    (user[0],token_hash,int(time.time()) + 1800))
                connection.commit()
                reset_url = os.environ["FLOWLAB_PUBLIC_URL"].rstrip("/") + url_for(
                    "email_password_reset", token=token)
                try:
                    send_password_reset_email(recipient, reset_url)
                    delivered = True
                    audit_change(connection, "user", user[0], "email_password_reset_requested",
                        after={"username":user[1], "delivery":"email"})
                except Exception:
                    app.logger.exception("Password recovery email delivery failed")
                    connection.execute("""UPDATE password_reset_tokens SET used_at=CURRENT_TIMESTAMP
                        WHERE token_hash=?""", (token_hash,))
            if not delivered:
                create_officer_reset_request(connection, user)
            connection.commit()
        connection.close()
        flash("If that active account exists, recovery instructions have been sent to its registered email or forwarded for controlled support.", "success")
        return redirect(url_for("login"))
    return render_template("forgot_password.html", page_title="Forgot Password")


@app.route("/reset-password/<token>", methods=["GET", "POST"])
def email_password_reset(token):
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    connection = get_connection()
    reset = connection.execute("""SELECT t.id,t.user_id,t.expires_at_epoch,t.used_at,u.username
        FROM password_reset_tokens t JOIN users u ON u.id=t.user_id
        WHERE t.token_hash=? AND u.is_active=1""", (token_hash,)).fetchone()
    valid = bool(reset and not reset[3] and reset[2] >= int(time.time()))
    if request.method == "POST" and valid:
        password = request.form.get("password", "")
        confirmation = request.form.get("password_confirmation", "")
        if len(password) < 10:
            flash("Create a password of at least 10 characters.", "error")
        elif password != confirmation:
            flash("Password confirmation does not match.", "error")
        else:
            connection.execute("UPDATE users SET password_hash=?,must_change_password=0 WHERE id=?",
                (AuthService.hash_password(password),reset[1]))
            connection.execute("UPDATE password_reset_tokens SET used_at=CURRENT_TIMESTAMP WHERE id=?",
                (reset[0],))
            connection.execute("""UPDATE password_reset_requests SET status='Completed',
                resolved_by='Email recovery',resolved_at=CURRENT_TIMESTAMP
                WHERE user_id=? AND status IN ('Pending','Resolved')""", (reset[1],))
            audit_change(connection, "user", reset[1], "email_password_reset_completed",
                after={"username":reset[4]})
            connection.commit(); connection.close()
            flash("Your password has been reset. Sign in with the new password.", "success")
            return redirect(url_for("login"))
    connection.close()
    return render_template("reset_password.html", page_title="Reset Password", valid_token=valid)


@app.route("/change-password", methods=["GET", "POST"])
def change_password():
    if not session.get("user_id"):
        return redirect(url_for("login"))
    if request.method == "POST":
        password = request.form.get("password", "")
        confirmation = request.form.get("password_confirmation", "")
        recovery_email = request.form.get("recovery_email", "").strip()
        if not valid_email(recovery_email):
            flash("Enter a valid recovery email address.", "error")
        elif len(password) < 10:
            flash("Create a password of at least 10 characters.", "error")
        elif password != confirmation:
            flash("Password confirmation does not match.", "error")
        else:
            connection = get_connection()
            existing = connection.execute("SELECT password_hash FROM users WHERE id=?",
                (session["user_id"],)).fetchone()
            password_hash = AuthService.hash_password(password)
            if existing and existing[0] == password_hash:
                connection.close(); flash("Create a password different from the temporary password.", "error")
            else:
                connection.execute("""UPDATE users SET password_hash=?,email=?,must_change_password=0
                    WHERE id=?""", (password_hash,recovery_email,session["user_id"]))
                connection.execute("""UPDATE password_reset_requests SET status='Completed',
                    resolved_by=?,resolved_at=CURRENT_TIMESTAMP WHERE user_id=? AND status='Resolved'""",
                    (current_actor(),session["user_id"]))
                audit_change(connection, "user", session["user_id"], "password_change",
                    after={"username": session.get("username"), "recovery_email":recovery_email,
                        "must_change_password": 0})
                connection.commit(); connection.close(); session.clear()
                flash("Your new password has been created. Sign in with it to continue.", "success")
                return redirect(url_for("login"))
    recovery_email=query("SELECT COALESCE(email,'') FROM users WHERE id=?",(session["user_id"],))[0][0]
    return render_template("change_password.html", page_title="Create New Password",
        recovery_email=recovery_email)


@app.route("/setup", methods=["GET", "POST"])
def setup():
    if query("SELECT COUNT(*) FROM users")[0][0] > 0:
        return redirect(url_for("login"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        full_name = request.form.get("full_name", "").strip()
        password = request.form.get("password", "")
        confirmation = request.form.get("password_confirmation", "")
        if not username or not full_name or len(password) < 10:
            flash("Username, full name, and a password of at least 10 characters are required.", "error")
        elif password != confirmation:
            flash("Password confirmation does not match.", "error")
        else:
            connection = get_connection()
            role = connection.execute("SELECT id FROM roles WHERE name='Chief Meteorologist'").fetchone()
            cursor = connection.execute("""INSERT INTO users
                (username,password_hash,full_name,role_id,is_active) VALUES (?,?,?,?,1)""",
                (username, AuthService.hash_password(password), full_name, role[0]))
            audit_change(connection, "user", cursor.lastrowid, "initial_administrator_create",
                after={"username": username, "full_name": full_name, "role": "Chief Meteorologist"})
            connection.commit(); connection.close()
            flash("Administrator account created. Sign in to continue.", "success")
            return redirect(url_for("login"))
    return render_template("setup.html", page_title="Initial Setup")


@app.post("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.get("/")
def dashboard():
    due_soon_days = setting_integer("equipment_due_soon_days", 30)
    counts = {
        "equipment": query("""SELECT COUNT(*) FROM laboratory_equipment WHERE is_active=1
            AND TRIM(equipment_name)<>'' AND TRIM(asset_number)<>''""")[0][0],
        "programme": query("""SELECT COUNT(*) FROM laboratory_equipment
            WHERE include_in_calibration_programme=1 AND is_active=1""")[0][0],
        "documents": query("SELECT COUNT(*) FROM controlled_documents WHERE is_deleted=0")[0][0],
        "sessions": query("SELECT COUNT(*) FROM measurement_sessions")[0][0],
    }
    upcoming = query("""SELECT asset_number,equipment_name,next_calibration_date,
        CASE WHEN next_calibration_date IS NOT NULL AND next_calibration_date < date('now','localtime') THEN 'Expired'
             WHEN next_calibration_date IS NOT NULL AND next_calibration_date <= date('now','localtime',?) THEN 'Due Soon'
             ELSE 'Active' END
        FROM laboratory_equipment WHERE include_in_calibration_programme=1 AND is_active=1
        ORDER BY CASE WHEN next_calibration_date IS NULL THEN 1 ELSE 0 END,
        next_calibration_date LIMIT 6""", (f"+{due_soon_days} days",))
    overdue = query("""SELECT COUNT(*) FROM laboratory_equipment
        WHERE include_in_calibration_programme=1 AND is_active=1 AND next_calibration_date IS NOT NULL
        AND next_calibration_date < date('now')""")[0][0]
    due_today = query("""SELECT COUNT(*) FROM laboratory_equipment
        WHERE include_in_calibration_programme=1 AND is_active=1
        AND next_calibration_date = date('now')""")[0][0]
    due_soon = query("""SELECT COUNT(*) FROM laboratory_equipment
        WHERE include_in_calibration_programme=1 AND is_active=1
        AND next_calibration_date > date('now')
        AND next_calibration_date <= date('now', ?)""", (f"+{due_soon_days} days",))[0][0]
    review_tasks = query("""SELECT entity_type, entity_id, submitted_by,
        submitted_at, priority FROM workflow_tasks
        WHERE status='Pending' AND (assigned_user_id=? OR assigned_to=?)
        ORDER BY submitted_at LIMIT 5""", (session.get("user_id"), current_actor()))
    incomplete_jobs = query("""SELECT j.id,j.job_number,COALESCE(j.job_title,''),
        COALESCE(j.planned_start_date,''),COALESCE(j.required_date,''),j.status,
        COALESCE(u.full_name,'Unassigned') FROM calibration_jobs j
        LEFT JOIN users u ON u.id=j.assigned_user_id
        WHERE j.is_deleted=0 AND j.status NOT IN ('Completed','Cancelled')
        AND (?=1 OR j.assigned_user_id=? OR j.assigned_user_id IS NULL)
        ORDER BY CASE WHEN NULLIF(j.required_date,'') IS NULL THEN 1 ELSE 0 END,
        j.required_date,j.id LIMIT 12""", (int(workflow_authority()),session.get("user_id")))
    standards = query("SELECT COUNT(*) FROM laboratory_equipment WHERE is_reference_standard=1 AND is_active=1")[0][0]
    kpis = (
        {"value": counts["equipment"], "label": "Equipment Assets"},
        {"value": counts["programme"], "label": "In Programme"},
        {"value": standards, "label": "Reference Standards"},
        {"value": overdue, "label": "Overdue"},
    )
    readiness = [
        {"name": "Equipment register", "detail": f'{counts["equipment"]} controlled assets', "status": "Ready" if counts["equipment"] else "Warning", "warning": not counts["equipment"]},
        {"name": "Calibration programme", "detail": f'{counts["programme"]} assets under schedule control', "status": "Ready" if counts["programme"] and not overdue else "Warning", "warning": not counts["programme"] or bool(overdue)},
        {"name": "Controlled documents", "detail": f'{counts["documents"]} registered documents', "status": "Ready" if counts["documents"] else "Warning", "warning": not counts["documents"]},
        {"name": "Measurement evidence", "detail": f'{counts["sessions"]} recorded sessions', "status": "Ready" if counts["sessions"] else "Warning", "warning": not counts["sessions"]},
    ]
    attention = [
        {"title": "Calibration due-date review", "detail": f"{overdue} overdue programme assets", "label": "Action" if overdue else "Clear", "tone": "overdue" if overdue else "ready"},
        {"title": "Equipment due today", "detail": f"{due_today} programme assets", "label": "Today" if due_today else "Clear", "tone": "overdue" if due_today else "ready"},
        {"title": "Equipment due soon", "detail": f"{due_soon} due within {due_soon_days} days", "label": "Due soon" if due_soon else "Clear", "tone": "due" if due_soon else "ready"},
    ]
    attention += [{"title": f"{task[0].replace('_', ' ').title()} awaiting review",
        "detail": f"Record {task[1]} · submitted by {task[2]} at {task[3]}",
        "label": task[4] or "Normal", "tone": "due"} for task in review_tasks]
    return render_template("dashboard.html", counts=counts, kpis=kpis,
        readiness=readiness, upcoming=upcoming, attention=attention,incomplete_jobs=incomplete_jobs,
        today=date.today().isoformat(), page_title=f"Welcome back, {current_actor()}",
        page_subtitle="Today's calibration and compliance picture", active_nav="dashboard")


@app.get("/reminders")
def reminders():
    user_id = session.get("user_id")
    connection = get_connection()
    try:
        notifications = connection.execute("""SELECT id,notification_type,title,message,link,
            entity_type,entity_id,is_read,created_at FROM notifications
            WHERE user_id=? AND is_read=0 ORDER BY created_at DESC LIMIT 100""", (user_id,)).fetchall()
        equipment_reminders = connection.execute("""SELECT id,asset_number,equipment_name,
            next_calibration_date,CASE WHEN next_calibration_date<date('now') THEN 'Expired'
            WHEN next_calibration_date=date('now') THEN 'Due Today' ELSE 'Due Soon' END
            FROM laboratory_equipment WHERE is_active=1 AND include_in_calibration_programme=1
            AND next_calibration_date IS NOT NULL AND next_calibration_date<=date('now','+30 days')
            ORDER BY next_calibration_date""").fetchall()
        job_reminders = connection.execute("""SELECT j.id,j.job_number,COALESCE(j.job_title,''),
            COALESCE(j.planned_start_date,''),COALESCE(j.required_date,''),j.status,
            COALESCE(u.full_name,'Unassigned') FROM calibration_jobs j
            LEFT JOIN users u ON u.id=j.assigned_user_id WHERE j.is_deleted=0
            AND j.status NOT IN ('Completed','Cancelled')
            AND COALESCE(NULLIF(j.required_date,''),NULLIF(j.planned_start_date,'')) IS NOT NULL
            AND COALESCE(NULLIF(j.required_date,''),NULLIF(j.planned_start_date,''))<=date('now','+14 days')
            AND (j.assigned_user_id=? OR j.assigned_user_id IS NULL OR ? IN ('Chief Meteorologist','Supervisor','Administrator'))
            ORDER BY COALESCE(NULLIF(j.required_date,''),NULLIF(j.planned_start_date,''))""",
            (user_id, session.get("role_name", ""))).fetchall()
        task_rows = connection.execute("""SELECT id,entity_type,entity_id,task_type,submitted_by,
            submitted_at,assigned_to,priority,due_date FROM workflow_tasks
            WHERE status='Pending' AND (?=1 OR assigned_user_id=? OR assigned_to=?)
            ORDER BY CASE priority WHEN 'Urgent' THEN 1 WHEN 'High' THEN 2 ELSE 3 END,submitted_at""",
            (int(workflow_authority()),user_id,current_actor())).fetchall()
        tasks = [{"id":row[0], "entity_type":row[1], "entity_id":row[2],
            "task_type":row[3], "submitted_by":row[4], "submitted_at":row[5],
            "assigned_to":row[6], "priority":row[7] or "Normal", "due_date":row[8],
            "url":task_destination(connection,row[1],row[2])} for row in task_rows]
        connection.commit()
        return render_template("reminders.html", notifications=notifications,
            equipment_reminders=equipment_reminders, job_reminders=job_reminders, tasks=tasks,
            has_workflow_authority=workflow_authority(),
            page_title="Notifications & Reminders", page_subtitle="Assigned reviews, due equipment and scheduled Jobs")
    finally:
        connection.close()


@app.get("/workflow/tasks/<int:task_id>")
def workflow_task_overview(task_id):
    if not workflow_authority():
        abort(403)
    connection = get_connection()
    try:
        task = connection.execute("""SELECT id,entity_type,entity_id,task_type,status,
            submitted_by,submitted_at,assigned_to,priority,due_date,COALESCE(comment,'')
            FROM workflow_tasks WHERE id=? AND status='Pending'""",(task_id,)).fetchone()
        if not task:
            abort(404)
        entity_type,entity_id = task[1],task[2]
        title = f"{entity_type.replace('_',' ').title()} · Record {entity_id}"
        facts,files = [],[]
        if entity_type == "controlled_document":
            row = connection.execute("""SELECT title,document_type,version,status,
                original_filename FROM controlled_documents WHERE id=? AND is_deleted=0""",(entity_id,)).fetchone()
            if row:
                title=row[0]; facts=[("Document type",row[1]),("Version",row[2]),("Status",row[3])]
                files=[(row[4] or "Associated document",url_for("document_file",document_id=entity_id))]
        elif entity_type in {"uncertainty_calculation","uncertainty"}:
            row=connection.execute("""SELECT calculation_uid,method_name,status,revision,
                expanded_uncertainty FROM uncertainty_calculations WHERE id=?""",(entity_id,)).fetchone()
            if row:
                title=row[0]; facts=[("Method",row[1]),("Status",row[2]),("Revision",row[3]),("Expanded uncertainty",row[4])]
        elif entity_type in {"calibration","maintenance"}:
            table="calibration_history" if entity_type=="calibration" else "maintenance_history"
            path_column="certificate_path" if entity_type=="calibration" else "document_path"
            row=connection.execute(f"""SELECT r.status,e.asset_number,e.equipment_name,e.id,r.{path_column}
                FROM {table} r JOIN laboratory_equipment e ON e.id=r.equipment_id WHERE r.id=?""",(entity_id,)).fetchone()
            if row:
                title=f"{row[1]} · {row[2]}";facts=[("Record type",entity_type.title()),("Status",row[0])]
                if row[4]: files=[("Associated report",url_for("equipment_file",equipment_id=row[3],kind=entity_type,record_id=entity_id))]
        elif entity_type in {"calibration_job","job"}:
            row=connection.execute("SELECT job_number,job_title,status FROM calibration_jobs WHERE id=?",(entity_id,)).fetchone()
            if row:
                title=row[0];facts=[("Title",row[1]),("Status",row[2])]
        elif entity_type=="capa":
            row=connection.execute("""SELECT ncr_number,title,source,clause_reference,status,
                target_close_date,issued_ncr_path,closeout_report_path FROM capa_records WHERE id=?""",(entity_id,)).fetchone()
            if row:
                title=f"{row[0]} · {row[1]}";facts=[("Source",row[2]),("ISO/IEC 17025 clause",row[3]),
                    ("Status",row[4]),("Target close-out",row[5])]
                files=[("Issued NCR",url_for("capa_file",record_id=entity_id,kind="issued"))]
                if row[7]: files.append(("Close-out report",url_for("capa_file",record_id=entity_id,kind="closeout")))
        elif entity_type == "iso_clause_assessment":
            row=connection.execute("""SELECT c.clause_code,c.title,a.compliance_status,a.status,
                COALESCE(a.finding,''),COALESCE(a.planned_action,'') FROM iso_clause_assessments a
                JOIN iso_clauses c ON c.id=a.clause_id WHERE c.id=?""",(entity_id,)).fetchone()
            if row:
                title=f"Clause {row[0]} · {row[1]}";facts=[("Compliance",row[2]),("Status",row[3]),
                    ("Finding",row[4] or "—"),("Planned action",row[5] or "—")]
        elif entity_type == "iso_clause_evidence":
            row=connection.execute("""SELECT c.clause_code,c.title,e.title,e.document_number,e.revision,
                e.status,e.file_name FROM iso_clause_evidence e JOIN iso_clauses c ON c.id=e.clause_id
                WHERE e.id=? AND e.is_deleted=0""",(entity_id,)).fetchone()
            if row:
                title=f"Clause {row[0]} · {row[2]}";facts=[("Requirement area",row[1]),
                    ("Document number",row[3] or "—"),("Revision",row[4] or "—"),("Status",row[5])]
                files=[(row[6],url_for("iso_clause_evidence_file",evidence_id=entity_id))]
        destination=task_destination(connection,entity_type,entity_id)
        return render_template("workflow_task.html",task=task,title=title,facts=facts,files=files,
            destination=destination,page_title="Approval Task",page_subtitle="Review the record and its evidence before deciding",active_nav="database")
    finally:
        connection.close()


@app.get("/notifications/<int:notification_id>/open")
def notification_open(notification_id):
    connection = get_connection()
    try:
        notification = connection.execute("""SELECT link,entity_type,entity_id FROM notifications
            WHERE id=? AND user_id=?""", (notification_id,session.get("user_id"))).fetchone()
        if not notification:
            abort(404)
        destination = task_destination(connection,notification[1],notification[2],notification[0])
        connection.execute("""UPDATE notifications SET is_read=1,read_at=CURRENT_TIMESTAMP
            WHERE id=?""", (notification_id,))
        connection.commit()
        return redirect(destination)
    finally:
        connection.close()


def _save_capa_file(upload, ncr_number, category):
    if not upload or not upload.filename:
        return None
    if Path(upload.filename).suffix.lower() not in REPORT_EXTENSIONS:
        raise ValueError("Evidence must be PDF, DOCX, XLSX, JPG or PNG.")
    destination=Path("data/documents/capa")/ncr_number/category
    destination.mkdir(parents=True,exist_ok=True)
    target=destination/f"{uuid4().hex}_{Path(upload.filename).name}"
    upload.save(target)
    return str(target)


@app.route("/capa",methods=["GET","POST"])
def capa():
    connection=get_connection()
    try:
        if request.method=="POST":
            require_permission("equipment_add")
            required={"NCR number":request.form.get("ncr_number","").strip(),
                "Title":request.form.get("title","").strip(),"Source":request.form.get("source","").strip(),
                "Description":request.form.get("description","").strip(),"Owner":request.form.get("owner","").strip(),
                "Issued date":request.form.get("issued_date"),"Target close-out date":request.form.get("target_close_date")}
            missing=[name for name,value in required.items() if not value]
            if missing: raise ValueError("Required NCR fields: "+", ".join(missing)+".")
            issued_path=_save_capa_file(request.files.get("issued_ncr"),required["NCR number"],"issued")
            if not issued_path: raise ValueError("Upload the issued NCR before saving the record.")
            cursor=connection.execute("""INSERT INTO capa_records
                (ncr_number,title,source,clause_reference,description,owner,issued_date,
                 target_close_date,status,issued_ncr_path,created_by)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)""",(required["NCR number"],required["Title"],required["Source"],
                request.form.get("clause_reference","").strip(),required["Description"],required["Owner"],
                required["Issued date"],required["Target close-out date"],"Open",issued_path,current_actor()))
            audit_change(connection,"capa",cursor.lastrowid,"create",after={"ncr_number":required["NCR number"],"status":"Open"})
            connection.commit();flash("NCR opened. Complete corrective-action evidence before submitting closure.","success")
            return redirect(url_for("capa_detail",record_id=cursor.lastrowid))
        records=connection.execute("""SELECT id,ncr_number,title,source,owner,issued_date,
            target_close_date,status,approved_by FROM capa_records WHERE is_deleted=0
            ORDER BY CASE status WHEN 'Submitted for Closure Review' THEN 1 WHEN 'Open' THEN 2
            WHEN 'Reverted' THEN 3 ELSE 4 END,target_close_date""").fetchall()
        return render_template("capa.html",records=records,page_title="CAPA & NCR",page_subtitle="ISO/IEC 17025 nonconforming work and corrective action",active_nav="iso17025")
    except ValueError as error:
        connection.rollback();flash(str(error),"error");return redirect(url_for("capa"))
    finally: connection.close()


@app.route("/capa/<int:record_id>",methods=["GET","POST"])
def capa_detail(record_id):
    connection=get_connection()
    try:
        row=connection.execute("""SELECT id,ncr_number,title,source,clause_reference,description,
            immediate_correction,root_cause,corrective_action,owner,issued_date,target_close_date,
            status,issued_ncr_path,closeout_report_path,created_by,assigned_reviewer_id,
            submitted_by,submitted_at,review_comment,approved_at,approved_by,created_at,closed_at,
            updated_at,is_deleted FROM capa_records WHERE id=? AND is_deleted=0""",(record_id,)).fetchone()
        if not row: abort(404)
        if request.method=="POST":
            require_permission("equipment_edit")
            if row[12]=="Closed": raise ValueError("Closed NCR records are locked.")
            closeout=_save_capa_file(request.files.get("closeout_report"),row[1],"closeout") or row[14]
            connection.execute("""UPDATE capa_records SET immediate_correction=?,root_cause=?,
                corrective_action=?,owner=?,target_close_date=?,closeout_report_path=?,
                status=CASE WHEN status='Reverted' THEN 'Open' ELSE status END,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                (request.form.get("immediate_correction","").strip(),request.form.get("root_cause","").strip(),
                 request.form.get("corrective_action","").strip(),request.form.get("owner","").strip(),
                 request.form.get("target_close_date"),closeout,record_id))
            audit_change(connection,"capa",record_id,"update",after={"closeout_report":bool(closeout)})
            connection.commit();flash("Corrective-action record updated.","success")
            return redirect(url_for("capa_detail",record_id=record_id))
        reviewers=connection.execute("""SELECT u.id,u.full_name,r.name FROM users u JOIN roles r ON r.id=u.role_id
            WHERE u.is_active=1 AND r.name IN ('Chief Meteorologist','Administrator') ORDER BY u.full_name""").fetchall()
        history=connection.execute("""SELECT action,actor,notes,occurred_at FROM approval_history
            WHERE entity_type='capa' AND entity_id=? ORDER BY id""",(record_id,)).fetchall()
        return render_template("capa_detail.html",record=row,reviewers=reviewers,history=history,
            page_title=f"NCR · {row[1]}",page_subtitle=row[2],active_nav="iso17025")
    except ValueError as error:
        connection.rollback();flash(str(error),"error");return redirect(url_for("capa_detail",record_id=record_id))
    finally: connection.close()


@app.post("/capa/<int:record_id>/submit")
def capa_submit(record_id):
    require_permission("equipment_edit")
    connection=get_connection()
    try:
        row=connection.execute("""SELECT ncr_number,status,immediate_correction,root_cause,
            corrective_action,issued_ncr_path,closeout_report_path,created_by FROM capa_records
            WHERE id=? AND is_deleted=0""",(record_id,)).fetchone()
        reviewer_id=request.form.get("reviewer_user_id",type=int)
        reviewer=connection.execute("""SELECT u.full_name FROM users u JOIN roles r ON r.id=u.role_id
            WHERE u.id=? AND u.is_active=1 AND r.name IN ('Chief Meteorologist','Administrator')""",(reviewer_id,)).fetchone()
        if not row or row[1] not in {"Open","Reverted"}: raise ValueError("This NCR is not ready for closure submission.")
        if not all(row[index] for index in range(2,7)): raise ValueError("Correction, root cause, corrective action, issued NCR and close-out report are mandatory.")
        if not reviewer: raise ValueError("Select a Chief Meteorologist or Administrator to review closure.")
        connection.execute("""UPDATE capa_records SET status='Submitted for Closure Review',
            assigned_reviewer_id=?,submitted_by=?,submitted_at=CURRENT_TIMESTAMP,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
            (reviewer_id,current_actor(),record_id))
        connection.execute("""INSERT INTO workflow_tasks(entity_type,entity_id,task_type,status,
            submitted_by,assigned_to,submitted_user_id,assigned_user_id,priority,due_date)
            VALUES (?,?,?,'Pending',?,?,?,?,?,(SELECT target_close_date FROM capa_records WHERE id=?))""",
            ("capa",record_id,"CAPA Closure Review",current_actor(),reviewer[0],session.get("user_id"),reviewer_id,"High",record_id))
        audit_change(connection,"capa",record_id,"submit",after={"status":"Submitted for Closure Review","reviewer":reviewer[0]})
        connection.commit();flash("NCR closure submitted for Chief/Admin review.","success")
    except ValueError as error: connection.rollback();flash(str(error),"error")
    finally: connection.close()
    return redirect(url_for("capa_detail",record_id=record_id))


@app.post("/capa/<int:record_id>/workflow")
def capa_workflow(record_id):
    if not workflow_authority(): abort(403)
    action=request.form.get("action");comment=request.form.get("comment","").strip()
    connection=get_connection()
    try:
        row=connection.execute("SELECT status,created_by FROM capa_records WHERE id=? AND is_deleted=0",(record_id,)).fetchone()
        if not row or row[0]!="Submitted for Closure Review": raise ValueError("This NCR is not awaiting closure review.")
        if action=="revert":
            if not comment: raise ValueError("A review comment is required when reverting closure.")
            status="Reverted";connection.execute("""UPDATE capa_records SET status=?,review_comment=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",(status,comment,record_id))
        elif action=="approve":
            status="Closed";connection.execute("""UPDATE capa_records SET status=?,approved_by=?,approved_at=CURRENT_TIMESTAMP,
                closed_at=CURRENT_TIMESTAMP,review_comment=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",(status,current_actor(),comment,record_id))
        else: raise ValueError("Select Approve Closure or Revert.")
        connection.execute("""UPDATE workflow_tasks SET status='Completed',reviewed_at=CURRENT_TIMESTAMP,
            decision=?,comment=? WHERE entity_type='capa' AND entity_id=? AND status='Pending'""",(action,comment,record_id))
        creator_id=user_id_for_actor(connection,row[1]);add_notification(connection,creator_id,
            f"NCR closure {status.lower()}",comment or f"Decision by {current_actor()}.",url_for("capa_detail",record_id=record_id),"capa",record_id)
        audit_change(connection,"capa",record_id,action,after={"status":status,"comment":comment})
        connection.commit();flash(f"NCR status changed to {status}.","success")
    except ValueError as error: connection.rollback();flash(str(error),"error")
    finally: connection.close()
    return redirect(url_for("capa_detail",record_id=record_id))


@app.get("/capa/<int:record_id>/file/<kind>")
def capa_file(record_id,kind):
    column="issued_ncr_path" if kind=="issued" else "closeout_report_path" if kind=="closeout" else None
    if not column: abort(404)
    rows=query(f"SELECT {column} FROM capa_records WHERE id=? AND is_deleted=0",(record_id,))
    if not rows or not rows[0][0] or not Path(rows[0][0]).is_file(): abort(404)
    return send_file(Path(rows[0][0]).resolve(),as_attachment=False)


@app.route("/api/drafts/<path:draft_key>", methods=["GET", "POST", "DELETE"])
def user_draft(draft_key):
    user_id = session.get("user_id")
    connection = get_connection()
    try:
        if request.method == "GET":
            row = connection.execute("SELECT payload_json,updated_at FROM user_drafts WHERE user_id=? AND draft_key=?",
                (user_id,draft_key)).fetchone()
            return jsonify(ok=True, draft=json.loads(row[0]) if row else None,
                updatedAt=row[1] if row else None)
        if request.method == "DELETE":
            connection.execute("DELETE FROM user_drafts WHERE user_id=? AND draft_key=?", (user_id,draft_key))
            connection.commit(); return jsonify(ok=True)
        payload = request.get_json(silent=True) or {}
        connection.execute("""INSERT INTO user_drafts(user_id,draft_key,payload_json)
            VALUES (?,?,?) ON CONFLICT(user_id,draft_key) DO UPDATE SET
            payload_json=excluded.payload_json,updated_at=CURRENT_TIMESTAMP""",
            (user_id,draft_key,json.dumps(payload)))
        connection.commit(); return jsonify(ok=True)
    finally:
        connection.close()


WORKSPACES = {
    "analytics": {
        "title": "Analytics", "description": "Live analysis sourced from controlled records—never duplicate data entry.",
        "modules": (
            ("calibration", "Calibration Analytics", "Throughput, turnaround and calibration status."),
            ("uncertainty", "Uncertainty Analytics", "Budgets, contributors and result trends."),
            ("cmc", "CMC vs Actual", "Compare declared capability with achieved results."),
            ("equipment", "Equipment Performance", "Reliability, drift and calibration performance."),
            ("capa", "CAPA Trends", "Nonconformity, cause and closure trends."),
            ("audit-pt", "Audit / PT Performance", "Findings and proficiency-testing outcomes."),
        )},
    "iso17025": {
        "title": "ISO 17025 Control Centre", "description": "",
        "modules": (
            ("clause-checklist", "Clause Checklist", "Requirement, evidence, owner, findings and action."),
            ("clause-scope", "Accreditation Scope & Clause Applicability", "Chief-controlled selection of clauses released to the laboratory checklist.", "iso_clause_scope"),
            ("uncertainty", "Uncertainty Management", "Measurement budgets, CMC, approvals and retained records.", "uncertainty"),
            ("calibration-programme", "Calibration Programme", "Master equipment lifecycle and uncertainty source.", "equipment"),
            ("traceability", "Traceability Register", "Equipment, certificate, calibration chain and evidence."),
            ("capa", "CAPA", "Nonconformity through effectiveness review and closure.", "capa"),
            ("audit-pt", "Audit & PT Programme", "Schedule audits, PT and ILC; capture findings."),
            ("methods", "Methods & Validation", "Controlled revisions, verification and approval."),
            ("personnel", "Personnel & Competence", "Training, assessment and authorization records."),
            ("facilities", "Facilities & Environment", "Limits, readings, deviations and investigations."),
            ("validity", "Validity of Results", "QC checks, charts, PT/ILC and investigations."),
            ("documents", "Controlled Documents & Records", "Revision, review, approval and archive.", "documents"),
            ("management-review", "Management Review", "Inputs, decisions, actions and closure evidence."),
        )},
    "projects": {"title": "Projects", "description": "Master project records linking clients, work, evidence and closeout.",
        "modules": (("project-register", "Project Register & New Job", "Client, scope, schedule, personnel, equipment, jobs, documents and reports.", "project_jobs"),)},
    "laboratory": {"title": "E&I Laboratory", "description": "Discipline workspaces for calibration activities and assigned equipment.",
        "modules": (("mass", "Mass Lab", "Mass calibration activities and equipment."),("pressure", "Pressure Lab", "Pressure calibration activities and equipment."),("low-voltage", "Low Voltage Lab", "Electrical calibration and testing."),("dimension", "Dimension Lab", "Dimensional measurement and calibration."))},
    "standards": {"title": "Standards & Procedures", "description": "Controlled standards and procedure repository.",
        "modules": (("standards", "Standards", "External and internal controlled standards.", "documents"),("procedures", "Procedures", "Approved laboratory procedures.", "documents"),("sops", "SOPs", "Standard operating procedures.", "documents"),("work-instructions", "Work Instructions", "Task-level controlled instructions.", "documents"),("revision-approval", "Revision & Approval", "Review, approval, effective and archive lifecycle."))},
    "database": {"title": "Database & User Management", "description": "Shared governance, approvals, audit history, files and master data.",
        "modules": (("users", "User Management", "Users, roles, permissions and approval authority.", "users"),("approvals", "Approval Management", "Technical review, approval, revert and record locking."),("audit-trail", "Audit Trail", "Every controlled action with user and timestamp."),("files", "File Repository", "Files linked to their originating records."),("administration", "Data Administration", "Matrix uploads, master data and import/export."))},
    "flowpro": {"title": "FlowPro_Wiz", "description": "Integration workspace reserved for ongoing development.", "modules": ()},
    "settings": {"title": "Settings", "description": "System configuration only.",
        "modules": (("laboratory", "Laboratory Configuration", "Laboratory identity and defaults."),("units", "Units", "Approved engineering units."),("notifications", "Notifications", "Attention and reminder rules."),("approval-rules", "Approval Rules", "Shared review and approval routing."),("numbering", "Numbering", "Controlled record numbering schemes."),("system", "System Configuration", "Application-level configuration."))},
}

DEFAULT_WORKFLOW = ("Create or open record", "Capture required data", "Attach supporting evidence", "Assign owner or reviewer", "Save status and audit event")
WORKFLOWS = {
    "clause-checklist": ("Review requirement and status", "Link objective evidence", "Record responsible person and finding", "Create linked CAPA when required", "Save review status"),
    "capa": ("Record NCR or originating finding", "Immediate correction and root cause", "Define corrective action and owner", "Attach effectiveness evidence", "Approve and close"),
    "audit-pt": ("Choose activity type", "Schedule, assign and set recurrence", "Record result and findings", "Upload evidence", "Create linked CAPA when required"),
    "approvals": ("Submit record for technical review", "Assign HOD or engineer", "Review controlled record", "Approve and lock, or revert with comment", "Correct and resubmit if reverted"),
    "revision-approval": ("Create controlled revision", "Complete technical review", "Approve revision", "Set effective date", "Archive superseded revision"),
}


@app.get("/workspace/<section>")
def workspace(section):
    if section not in WORKSPACES:
        return redirect(url_for("dashboard"))
    if section == "settings":
        return redirect(url_for("settings"))
    if section == "database":
        require_permission("user_management")
    workspace_data = WORKSPACES[section]
    modules = []
    for index, item in enumerate(workspace_data["modules"]):
        slug, title, description, *endpoint = item
        if slug == "clause-scope" and session.get("role_name") != "Chief Meteorologist":
            continue
        target = url_for(endpoint[0]) if endpoint else url_for("module_page", section=section, module=slug)
        modules.append({"title": title, "description": description, "url": target,
            "icon": f"{index + 1:02d}", "action": "Open module"})
    return render_template("workspace.html", page_title=workspace_data["title"],
        page_subtitle=workspace_data["description"], description=workspace_data["description"],
        active_nav=section, modules=modules, eyebrow="Functional workspace")


@app.route("/users", methods=["GET", "POST"])
def users():
    require_permission("user_management")
    allowed_roles = ("Chief Meteorologist", "Supervisor", "Engineer", "Technician", "Operator", "Guest")
    connection = get_connection()
    try:
        roles = connection.execute(
            f"SELECT id,name FROM roles WHERE name IN ({','.join('?' for _ in allowed_roles)}) "
            f"ORDER BY CASE name {' '.join(f'WHEN ? THEN {index}' for index, _ in enumerate(allowed_roles, 1))} END",
            (*allowed_roles, *allowed_roles)).fetchall()
        if request.method == "POST":
            action = request.form.get("action", "")
            if action == "create":
                username = request.form.get("username", "").strip()
                full_name = request.form.get("full_name", "").strip()
                email = request.form.get("email", "").strip() or None
                password = request.form.get("password", "")
                role_id = request.form.get("role_id", type=int)
                role = next((item for item in roles if item[0] == role_id), None)
                if not username or not full_name or len(password) < 10 or not role:
                    raise ValueError("Username, full name, an approved role, and a password of at least 10 characters are required.")
                if connection.execute("SELECT 1 FROM users WHERE LOWER(username)=LOWER(?)", (username,)).fetchone():
                    raise ValueError("Username already exists.")
                cursor = connection.execute("""INSERT INTO users
                    (username,password_hash,full_name,email,role_id,is_active,must_change_password)
                    VALUES (?,?,?,?,?,1,1)""",
                    (username, AuthService.hash_password(password), full_name, email, role_id))
                audit_change(connection, "user", cursor.lastrowid, "create", after={
                    "username": username, "full_name": full_name, "role": role[1], "is_active": 1})
                flash(f"User {username} created.", "success")
            elif action == "approve_personnel":
                registration_id = request.form.get("registration_id", type=int)
                registration = connection.execute("""SELECT id,full_name,email,requested_role,status,initiated_by_user_id
                    FROM personnel_registrations WHERE id=?""", (registration_id,)).fetchone()
                username = request.form.get("username", "").strip()
                full_name = request.form.get("full_name", "").strip()
                email = request.form.get("email", "").strip() or None
                password = request.form.get("password", "")
                role_id = request.form.get("role_id", type=int)
                role = next((item for item in roles if item[0] == role_id), None)
                if not registration or registration[4] != "Pending": raise ValueError("Pending personnel registration was not found.")
                if not username or not full_name or len(password) < 10 or not role:
                    raise ValueError("Complete the username, full name, role and temporary password of at least 10 characters.")
                if connection.execute("SELECT 1 FROM users WHERE LOWER(username)=LOWER(?)", (username,)).fetchone():
                    raise ValueError("Username already exists.")
                user_id = connection.execute("""INSERT INTO users
                    (username,password_hash,full_name,email,role_id,is_active,must_change_password)
                    VALUES (?,?,?,?,?,1,1)""", (username,AuthService.hash_password(password),full_name,email,role_id)).lastrowid
                connection.execute("""UPDATE personnel_registrations SET status='Approved',completed_user_id=?,
                    review_comment=?,reviewed_at=CURRENT_TIMESTAMP WHERE id=?""",
                    (user_id,request.form.get("comment","").strip() or None,registration_id))
                connection.execute("""UPDATE workflow_tasks SET status='Completed',decision='approve',reviewed_at=CURRENT_TIMESTAMP
                    WHERE entity_type='personnel_registration' AND entity_id=? AND status='Pending'""", (registration_id,))
                add_notification(connection,registration[5],"Personnel registration approved",
                    f"{full_name} is now an active FlowLab user.",url_for("personnel_record",user_id=user_id),"personnel_registration",registration_id)
                audit_change(connection,"personnel_registration",registration_id,"approve",after={"user_id":user_id,"role":role[1]})
                flash(f"Personnel account for {full_name} created and approved.","success")
            elif action == "update":
                user_id = request.form.get("user_id", type=int)
                full_name = request.form.get("full_name", "").strip()
                email = request.form.get("email", "").strip() or None
                role_id = request.form.get("role_id", type=int)
                role = next((item for item in roles if item[0] == role_id), None)
                existing = connection.execute("""SELECT u.username,u.full_name,u.email,r.name
                    FROM users u JOIN roles r ON r.id=u.role_id WHERE u.id=?""", (user_id,)).fetchone()
                if not existing or not full_name or not role:
                    raise ValueError("Select an existing user, full name, and approved role.")
                if existing[3] == "Administrator":
                    raise ValueError("The legacy Administrator role cannot be reassigned from this module.")
                connection.execute("UPDATE users SET full_name=?,email=?,role_id=? WHERE id=?",
                    (full_name, email, role_id, user_id))
                audit_change(connection, "user", user_id, "update",
                    before={"full_name": existing[1], "email": existing[2], "role": existing[3]},
                    after={"full_name": full_name, "email": email, "role": role[1]})
                flash(f"User {existing[0]} updated.", "success")
            elif action == "reset_password":
                user_id = request.form.get("user_id", type=int)
                password = request.form.get("password", "")
                existing = connection.execute("SELECT username FROM users WHERE id=?", (user_id,)).fetchone()
                if not existing or len(password) < 10:
                    raise ValueError("Select an existing user and enter a password of at least 10 characters.")
                connection.execute("UPDATE users SET password_hash=?,must_change_password=1 WHERE id=?",
                    (AuthService.hash_password(password), user_id))
                connection.execute("""UPDATE password_reset_requests SET status='Resolved',resolved_by=?,
                    resolved_at=CURRENT_TIMESTAMP WHERE user_id=? AND status='Pending'""",
                    (current_actor(),user_id))
                audit_change(connection, "user", user_id, "password_reset", after={"username": existing[0]})
                flash(f"Temporary password set for {existing[0]}. The user must create a new password at sign-in.", "success")
            elif action == "toggle_active":
                user_id = request.form.get("user_id", type=int)
                existing = connection.execute("SELECT username,is_active FROM users WHERE id=?", (user_id,)).fetchone()
                if not existing:
                    raise ValueError("User account was not found.")
                if user_id == session.get("user_id") and existing[1]:
                    raise ValueError("You cannot deactivate your own signed-in account.")
                new_state = 0 if existing[1] else 1
                connection.execute("UPDATE users SET is_active=? WHERE id=?", (new_state, user_id))
                audit_change(connection, "user", user_id, "activate" if new_state else "deactivate",
                    before={"is_active": existing[1]}, after={"is_active": new_state})
                flash(f"User {existing[0]} {'activated' if new_state else 'deactivated'}.", "success")
            else:
                raise ValueError("Select a valid user-management action.")
            connection.commit()
            return redirect(url_for("users"))
        rows = connection.execute("""SELECT u.id,u.username,u.full_name,COALESCE(u.email,''),
            r.name,u.role_id,u.is_active,COALESCE(u.last_login,''),u.created_at,u.must_change_password
            FROM users u JOIN roles r ON r.id=u.role_id
            ORDER BY CASE r.name WHEN 'Chief Meteorologist' THEN 1 WHEN 'Supervisor' THEN 2
            WHEN 'Engineer' THEN 3 WHEN 'Technician' THEN 4 WHEN 'Operator' THEN 5
            WHEN 'Guest' THEN 6 ELSE 7 END,u.full_name""").fetchall()
        reset_requests = connection.execute("""SELECT pr.id,u.id,u.username,u.full_name,
            pr.requested_at FROM password_reset_requests pr JOIN users u ON u.id=pr.user_id
            WHERE pr.status='Pending' ORDER BY pr.requested_at""").fetchall()
        pending_personnel = connection.execute("""SELECT p.id,p.full_name,COALESCE(p.email,''),p.requested_role,
            p.initiated_by,p.created_at FROM personnel_registrations p WHERE p.status='Pending' ORDER BY p.created_at""").fetchall()
        selected_registration = None
        registration_id = request.args.get("registration_id", type=int)
        if registration_id:
            selected_registration = connection.execute("""SELECT id,full_name,COALESCE(email,''),requested_role,initiated_by
                FROM personnel_registrations WHERE id=? AND status='Pending'""", (registration_id,)).fetchone()
        return render_template("users.html", users=rows, roles=roles, reset_requests=reset_requests,
            pending_personnel=pending_personnel,selected_registration=selected_registration,
            page_title="User Management", page_subtitle="Controlled accounts, roles and access",
            active_nav="database")
    except ValueError as error:
        connection.rollback(); flash(str(error), "error")
        return redirect(url_for("users"))
    finally:
        connection.close()


@app.route("/workspace/<section>/<module>", methods=["GET", "POST"])
def module_page(section, module):
    if section == "iso17025" and module == "clause-checklist":
        return redirect(url_for("iso_clause_checklist"))
    workspace_data = WORKSPACES.get(section)
    if not workspace_data:
        return redirect(url_for("dashboard"))
    match = next((item for item in workspace_data["modules"] if item[0] == module), None)
    if not match:
        return redirect(url_for("workspace", section=section))
    if request.method == "POST":
        require_permission("equipment_edit")
        title = request.form.get("title", "").strip()
        if not title:
            flash("Title is required.", "error")
        else:
            connection = get_connection()
            try:
                cursor = connection.execute("""INSERT INTO quality_records
                    (record_type,reference_number,title,status,owner,due_date,notes)
                    VALUES (?,?,?,?,?,?,?)""", (
                    module, request.form.get("reference_number", "").strip() or None,
                    title, request.form.get("status", "Open"),
                    request.form.get("owner", "").strip() or None,
                    request.form.get("due_date") or None,
                    request.form.get("notes", "").strip() or None))
                add_note(connection, "quality_record", cursor.lastrowid, request.form.get("notes"))
                audit_change(connection, "quality_record", cursor.lastrowid, "create", after={
                    "record_type": module, "title": title, "status": request.form.get("status", "Open")})
                connection.commit()
                flash("Record saved.", "success")
                return redirect(url_for("module_page", section=section, module=module))
            except Exception as error:
                connection.rollback()
                flash(str(error), "error")
            finally:
                connection.close()
    records = query("""SELECT id,COALESCE(reference_number,''),title,status,
        COALESCE(owner,''),due_date,created_at FROM quality_records
        WHERE record_type=? AND is_deleted=0 ORDER BY id DESC""", (module,))
    return render_template("module.html", section=section,
        parent_title=workspace_data["title"], page_title=match[1],
        page_subtitle=match[2], description=match[2], active_nav=section,
        records=records)


@app.get("/iso17025/clause-checklist")
def iso_clause_checklist():
    status = request.args.get("status", "").strip()
    category = request.args.get("category", "").strip()
    sql = """SELECT c.id,c.clause_code,c.title,c.category,c.guidance,
        CASE WHEN a.applicability_set_by IS NULL THEN 'Pending determination' ELSE a.applicability END,
        COALESCE(a.compliance_status,'Not Assessed'),
        COALESCE(a.status,'Draft'),COALESCE(u.full_name,''),a.next_review_date,
        (SELECT COUNT(*) FROM iso_clause_evidence_requirements r WHERE r.clause_id=c.id AND r.is_required=1 AND r.is_active=1),
        (SELECT COUNT(*) FROM iso_clause_evidence_requirements r WHERE r.clause_id=c.id AND r.is_required=1 AND r.is_active=1
            AND EXISTS (SELECT 1 FROM iso_clause_evidence e WHERE e.requirement_id=r.id AND e.status='Approved' AND e.is_deleted=0)),
        (SELECT COUNT(*) FROM iso_clause_evidence_requirements r WHERE r.clause_id=c.id AND r.is_required=1 AND r.is_active=1
            AND EXISTS (SELECT 1 FROM iso_clause_evidence e WHERE e.requirement_id=r.id AND e.is_deleted=0)),
        (SELECT COUNT(*) FROM iso_clause_evidence_requirements r WHERE r.clause_id=c.id AND r.is_required=1 AND r.is_active=1
            AND EXISTS (SELECT 1 FROM iso_clause_evidence e WHERE e.requirement_id=r.id AND e.status='Reverted' AND e.is_deleted=0))
        FROM iso_clauses c LEFT JOIN iso_clause_assessments a ON a.clause_id=c.id
        LEFT JOIN users u ON u.id=a.owner_user_id WHERE c.is_active=1
        AND a.applicability='Applicable' AND a.applicability_set_by IS NOT NULL"""
    parameters = []
    if status:
        sql += " AND COALESCE(a.compliance_status,'Not Assessed')=?"; parameters.append(status)
    if category:
        sql += " AND c.category=?"; parameters.append(category)
    sql += " ORDER BY c.sort_order"
    clauses = query(sql, tuple(parameters))
    clauses = [tuple(row) + (("Compliant" if row[10] and row[11] == row[10] else
        "Action Required" if row[13] else "Awaiting Approval" if row[10] and row[12] == row[10] else "Pending Records"),)
        for row in clauses]
    if any(row[1] == "6.2" for row in clauses):
        personnel_count = query("SELECT COUNT(*) FROM users WHERE is_active=1")[0][0]
        approved_count = query("""SELECT COUNT(*) FROM personnel_documents
            WHERE is_deleted=0 AND is_current=1 AND status='Approved'""")[0][0]
        personnel_pending = query("""SELECT COUNT(*) FROM personnel_documents
            WHERE is_deleted=0 AND status IN ('Submitted for Review','Reverted')""")[0][0]
        expected = personnel_count * len(PERSONNEL_DOCUMENT_TYPES)
        personnel_state = ("Compliant" if expected and approved_count >= expected else
            "Action Required" if personnel_pending else "Pending Records")
        clauses = [row[:-1] + (personnel_state,) if row[1] == "6.2" else row for row in clauses]
    totals = {"total": len(clauses), "compliant": sum(1 for row in clauses if row[14] == "Compliant"),
        "partial": sum(1 for row in clauses if row[6] == "Partially Compliant"),
        "nonconforming": sum(1 for row in clauses if row[6] == "Nonconforming"),
        "not_assessed": sum(1 for row in clauses if row[6] == "Not Assessed")}
    return render_template("clause_checklist.html", clauses=clauses, totals=totals,
        selected_status=status, selected_category=category,
        can_manage_scope=session.get("role_name")=="Chief Meteorologist",
        page_title="ISO 17025 Clause Checklist", page_subtitle="Clause assessment and objective evidence",
        active_nav="iso17025")


@app.route("/iso17025/clause-scope", methods=["GET", "POST"])
def iso_clause_scope():
    if session.get("role_name") != "Chief Meteorologist": abort(403)
    connection=get_connection()
    try:
        if request.method == "POST":
            clauses=connection.execute("SELECT id,clause_code FROM iso_clauses WHERE is_active=1 ORDER BY sort_order").fetchall()
            for clause_id,code in clauses:
                decision=request.form.get(f"decision_{clause_id}","Pending determination")
                justification=request.form.get(f"justification_{clause_id}","").strip()
                if decision not in {"Pending determination","Applicable","Not Applicable"}:
                    raise ValueError(f"Clause {code}: select a valid applicability decision.")
                if decision=="Not Applicable" and not justification:
                    raise ValueError(f"Clause {code}: justify the Not Applicable decision.")
                existing=connection.execute("SELECT id,applicability FROM iso_clause_assessments WHERE clause_id=?",(clause_id,)).fetchone()
                set_by=current_actor() if decision!="Pending determination" else None
                if existing:
                    timestamp_sql="NULL" if decision=="Pending determination" else "CURRENT_TIMESTAMP"
                    connection.execute(f"""UPDATE iso_clause_assessments SET applicability=?,applicability_reason=?,
                        applicability_set_by=?,applicability_set_at={timestamp_sql},updated_by=?,
                        updated_at=CURRENT_TIMESTAMP WHERE clause_id=?""",
                        (decision,justification or None,set_by,current_actor(),clause_id))
                    assessment_id=existing[0]
                else:
                    timestamp_sql="NULL" if decision=="Pending determination" else "CURRENT_TIMESTAMP"
                    assessment_id=connection.execute(f"""INSERT INTO iso_clause_assessments
                        (clause_id,applicability,applicability_reason,applicability_set_by,applicability_set_at,
                        compliance_status,status,created_by,updated_by) VALUES (?,?,?,?,{timestamp_sql},
                        'Not Assessed','Draft',?,?)""",
                        (clause_id,decision,justification or None,set_by,current_actor(),current_actor())).lastrowid
                audit_change(connection,"iso_clause_assessment",assessment_id,"set_applicability",
                    after={"clause":code,"decision":decision,"justification":justification or None})
            connection.commit();flash("Accreditation scope and clause applicability updated.","success")
            return redirect(url_for("iso_clause_scope"))
        rows=connection.execute("""SELECT c.id,c.clause_code,c.title,c.category,c.guidance,
            CASE WHEN a.applicability_set_by IS NULL THEN 'Pending determination' ELSE a.applicability END,
            COALESCE(a.applicability_reason,''),COALESCE(a.applicability_set_by,''),a.applicability_set_at
            FROM iso_clauses c LEFT JOIN iso_clause_assessments a ON a.clause_id=c.id
            WHERE c.is_active=1 ORDER BY c.sort_order""").fetchall()
        return render_template("clause_scope.html",clauses=rows,page_title="Accreditation Scope & Clause Applicability",
            page_subtitle="Chief Meteorologist controlled ISO/IEC 17025 scope decisions",active_nav="iso17025")
    except ValueError as error:
        connection.rollback();flash(str(error),"error");return redirect(url_for("iso_clause_scope"))
    finally: connection.close()


def _save_clause_evidence(upload, clause_code):
    validate_report(upload)
    destination = Path("data/documents/iso17025/clause_evidence") / clause_code.replace(".", "-")
    destination.mkdir(parents=True, exist_ok=True)
    safe_name = Path(upload.filename).name
    target = destination / f"{uuid4().hex}_{safe_name}"
    upload.save(target)
    return safe_name, str(target)


def _sync_clause_compliance(connection, clause_id):
    counts = connection.execute("""SELECT
        COUNT(*),
        SUM(CASE WHEN EXISTS (SELECT 1 FROM iso_clause_evidence e WHERE e.requirement_id=r.id
            AND e.status='Approved' AND e.is_deleted=0) THEN 1 ELSE 0 END),
        SUM(CASE WHEN EXISTS (SELECT 1 FROM iso_clause_evidence e WHERE e.requirement_id=r.id
            AND e.is_deleted=0) THEN 1 ELSE 0 END),
        SUM(CASE WHEN EXISTS (SELECT 1 FROM iso_clause_evidence e WHERE e.requirement_id=r.id
            AND e.status='Reverted' AND e.is_deleted=0) THEN 1 ELSE 0 END)
        FROM iso_clause_evidence_requirements r
        WHERE r.clause_id=? AND r.is_required=1 AND r.is_active=1""", (clause_id,)).fetchone()
    required, approved, uploaded, reverted = (int(value or 0) for value in counts)
    if required and approved == required:
        clause_status = "Compliant"
        connection.execute("""UPDATE iso_clause_assessments SET compliance_status='Compliant',status='Approved',
            approved_by=?,approved_at=CURRENT_TIMESTAMP,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE clause_id=?""",
            (current_actor(), current_actor(), clause_id))
    else:
        clause_status = "Action Required" if reverted else ("Awaiting Approval" if required and uploaded == required else "Pending Records")
        connection.execute("""UPDATE iso_clause_assessments SET compliance_status=?,status='Draft',
            approved_by=NULL,approved_at=NULL,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE clause_id=?""",
            (clause_status, current_actor(), clause_id))
    return clause_status


@app.route("/iso17025/clause-checklist/<int:clause_id>", methods=["GET", "POST"])
def iso_clause_detail(clause_id):
    connection = get_connection()
    try:
        clause = connection.execute("""SELECT c.id,c.clause_code,c.title,c.category,c.guidance FROM iso_clauses c
            JOIN iso_clause_assessments a ON a.clause_id=c.id WHERE c.id=? AND c.is_active=1
            AND a.applicability='Applicable' AND a.applicability_set_by IS NOT NULL""", (clause_id,)).fetchone()
        if not clause: abort(404)
        if clause[1] == "6.2":
            return redirect(url_for("personnel_clause",clause_id=clause_id))
        if request.method == "POST":
            require_permission("equipment_edit")
            chief_controls_applicability = False
            existing = connection.execute("""SELECT id,status,applicability,applicability_reason,
                applicability_set_by FROM iso_clause_assessments WHERE clause_id=?""", (clause_id,)).fetchone()
            applicability = existing[2]
            reason = existing[3] or ""
            compliance = request.form.get("compliance_status", "Not Assessed")
            if applicability not in {"Pending determination", "Applicable", "Not Applicable"} or compliance not in {"Not Assessed", "Compliant", "Partially Compliant", "Nonconforming"}:
                raise ValueError("Select valid applicability and compliance values.")
            if not chief_controls_applicability and applicability == "Pending determination":
                raise ValueError("The Chief Meteorologist must determine clause applicability before assessment can begin.")
            owner_id = request.form.get("owner_user_id", type=int)
            owner = connection.execute("SELECT 1 FROM users WHERE id=? AND is_active=1", (owner_id,)).fetchone() if owner_id else None
            if owner_id and not owner: raise ValueError("Select an active responsible officer.")
            if applicability == "Not Applicable" and not reason:
                raise ValueError("A justification is required when a clause is marked Not Applicable.")
            values = (applicability,reason or None,compliance,owner_id,
                request.form.get("finding", "").strip() or None,
                request.form.get("planned_action", "").strip() or None,
                request.form.get("target_date") or None,request.form.get("last_review_date") or None,
                request.form.get("next_review_date") or None,current_actor())
            if existing:
                if existing[1] == "Approved" and not workflow_authority():
                    raise ValueError("An approved assessment is locked. Chief/Admin authority is required to revise it.")
                connection.execute("""UPDATE iso_clause_assessments SET applicability=?,applicability_reason=?,
                    compliance_status=?,owner_user_id=?,finding=?,planned_action=?,target_date=?,last_review_date=?,
                    next_review_date=?,updated_by=?,status='Draft',approved_by=NULL,approved_at=NULL,
                    applicability_set_by=CASE WHEN ? THEN ? ELSE applicability_set_by END,
                    applicability_set_at=CASE WHEN ? THEN CURRENT_TIMESTAMP ELSE applicability_set_at END,
                    updated_at=CURRENT_TIMESTAMP WHERE clause_id=?""",
                    (*values,int(chief_controls_applicability),current_actor(),int(chief_controls_applicability),clause_id))
                assessment_id = existing[0]
            else:
                assessment_id = connection.execute("""INSERT INTO iso_clause_assessments
                    (clause_id,applicability,applicability_reason,compliance_status,owner_user_id,finding,
                    planned_action,target_date,last_review_date,next_review_date,created_by,updated_by,
                    applicability_set_by,applicability_set_at)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,CASE WHEN ? IS NOT NULL THEN CURRENT_TIMESTAMP ELSE NULL END)""",
                    (clause_id,*values[:-1],current_actor(),current_actor(),
                     current_actor() if chief_controls_applicability else None,
                     current_actor() if chief_controls_applicability else None)).lastrowid
            audit_change(connection,"iso_clause_assessment",assessment_id,"save_draft",
                after={"clause":clause[1],"compliance_status":compliance,"applicability":applicability})
            connection.commit(); flash(f"Clause {clause[1]} assessment saved as Draft.","success")
            return redirect(url_for("iso_clause_detail",clause_id=clause_id))
        assessment = connection.execute("""SELECT a.id,a.applicability,a.applicability_reason,a.compliance_status,
            a.owner_user_id,a.finding,a.planned_action,a.target_date,a.last_review_date,a.next_review_date,a.status,
            a.created_by,a.updated_by,a.submitted_by,a.submitted_at,a.review_comment,a.approved_by,a.approved_at,
            a.applicability_set_by,a.applicability_set_at
            FROM iso_clause_assessments a WHERE a.clause_id=?""", (clause_id,)).fetchone()
        requirements = connection.execute("""SELECT r.id,r.evidence_name,r.is_required,
            (SELECT COUNT(*) FROM iso_clause_evidence e WHERE e.requirement_id=r.id AND e.status='Approved' AND e.is_deleted=0)
            ,COALESCE(r.activity_scope,''),COALESCE(r.applicability_rule,''),COALESCE(r.required_fields,''),
            (SELECT COUNT(*) FROM iso_clause_evidence e WHERE e.requirement_id=r.id AND e.is_deleted=0),
            (SELECT COUNT(*) FROM iso_clause_evidence e WHERE e.requirement_id=r.id AND e.status='Submitted for Review' AND e.is_deleted=0),
            (SELECT COUNT(*) FROM iso_clause_evidence e WHERE e.requirement_id=r.id AND e.status='Reverted' AND e.is_deleted=0),
            COALESCE(r.subclause_code,c.clause_code)
            FROM iso_clause_evidence_requirements r JOIN iso_clauses c ON c.id=r.clause_id WHERE r.clause_id=? AND r.is_active=1
            ORDER BY r.is_required DESC,r.id""", (clause_id,)).fetchall()
        evidence = connection.execute("""SELECT e.id,e.title,COALESCE(e.document_number,''),COALESCE(e.revision,''),
            e.effective_date,e.review_date,e.retention_until,e.file_name,e.status,e.uploaded_by,
            COALESCE(e.review_comment,''),COALESCE(e.approved_by,''),e.created_at,
            COALESCE(r.evidence_name,'Additional evidence'),COALESCE(e.document_owner,''),COALESCE(r.subclause_code,c.clause_code) FROM iso_clause_evidence e
            LEFT JOIN iso_clause_evidence_requirements r ON r.id=e.requirement_id
            JOIN iso_clauses c ON c.id=e.clause_id
            WHERE e.clause_id=? AND e.is_deleted=0 ORDER BY e.id DESC""", (clause_id,)).fetchall()
        users = connection.execute("""SELECT u.id,u.full_name,r.name FROM users u JOIN roles r ON r.id=u.role_id
            WHERE u.is_active=1 ORDER BY u.full_name""").fetchall()
        reviewers = [row for row in users if row[2] in {"Chief Meteorologist","Administrator"}]
        capas = connection.execute("SELECT id,ncr_number,title,status FROM capa_records WHERE is_deleted=0 ORDER BY id DESC").fetchall()
        applicability = assessment[1] if assessment and assessment[18] else "Pending determination"
        missing_required = sum(1 for requirement in requirements if requirement[2] and not requirement[3])
        required = [requirement for requirement in requirements if requirement[2]]
        if required and all(requirement[3] for requirement in required): clause_status = "Compliant"
        elif any(requirement[9] for requirement in required): clause_status = "Action Required"
        elif required and all(requirement[7] for requirement in required): clause_status = "Awaiting Approval"
        else: clause_status = "Pending Records"
        return render_template("clause_detail.html",clause=clause,assessment=assessment,
            requirements=requirements,evidence=evidence,users=users,reviewers=reviewers,capas=capas,
            missing_required=missing_required,clause_status=clause_status,
            applicability=applicability,can_set_applicability=session.get("role_name")=="Chief Meteorologist",
            page_title=f"Clause {clause[1]} · {clause[2]}",page_subtitle=clause[4],active_nav="iso17025")
    except ValueError as error:
        connection.rollback(); flash(str(error),"error")
        return redirect(url_for("iso_clause_detail",clause_id=clause_id))
    finally:
        connection.close()


def _save_personnel_document(upload, user_id):
    suffix = Path(upload.filename).suffix.lower()
    if suffix not in REPORT_EXTENSIONS:
        raise ValueError("Upload a PDF, DOCX, XLSX, JPG or PNG file.")
    destination = Path("data/documents/iso17025/personnel") / str(user_id)
    destination.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid4().hex}{suffix}"
    stored_path = destination / stored_name
    upload.save(stored_path)
    return Path(upload.filename).name, str(stored_path)


@app.get("/iso17025/clause-checklist/<int:clause_id>/personnel")
def personnel_clause(clause_id):
    connection = get_connection()
    try:
        clause = connection.execute("""SELECT c.id,c.clause_code,c.title,c.guidance FROM iso_clauses c
            JOIN iso_clause_assessments a ON a.clause_id=c.id WHERE c.id=? AND c.clause_code='6.2'
            AND c.is_active=1 AND a.applicability='Applicable' AND a.applicability_set_by IS NOT NULL""", (clause_id,)).fetchone()
        if not clause: abort(404)
        users = connection.execute("""SELECT u.id,u.full_name,u.username,COALESCE(u.email,''),r.name,u.is_active
            FROM users u JOIN roles r ON r.id=u.role_id ORDER BY u.is_active DESC,u.full_name""").fetchall()
        documents = connection.execute("""SELECT d.id,d.user_id,d.document_type,d.document_number,d.revision,
            d.issued_date,d.file_name,d.status,d.uploaded_by,COALESCE(d.approved_by,''),
            COALESCE(d.review_comment,''),d.created_at,d.is_current,d.assigned_reviewer_id
            FROM personnel_documents d WHERE d.is_deleted=0 ORDER BY d.user_id,d.document_type,d.id DESC""").fetchall()
        by_user = {}
        for document in documents:
            by_user.setdefault(document[1], []).append(document)
        personnel = []
        for user in users:
            current = {}
            pending = {}
            for doc in by_user.get(user[0], []):
                if doc[7] == "Approved" and doc[12]: current.setdefault(doc[2], doc)
                if doc[7] in {"Submitted for Review","Reverted"}: pending.setdefault(doc[2], doc)
            personnel.append({"user":user,"current":current,"pending":pending,
                "documents":by_user.get(user[0], []),
                "complete":all(kind in current for kind in PERSONNEL_DOCUMENT_TYPES)})
        pending_registrations = connection.execute("SELECT COUNT(*) FROM personnel_registrations WHERE status='Pending'").fetchone()[0]
        selected_user_id = request.args.get("user_id", type=int)
        selected = next((item for item in personnel if item["user"][0] == selected_user_id), None)
        role = session.get("role_name")
        can_manage_all = role in {"Supervisor","Chief Meteorologist","Administrator"}
        if selected and not (can_manage_all or selected_user_id == session.get("user_id")):
            selected = None
        reviewers = connection.execute("""SELECT u.id,u.full_name,r.name FROM users u JOIN roles r ON r.id=u.role_id
            WHERE u.is_active=1 AND r.name IN ('Supervisor','Chief Meteorologist','Administrator')
            ORDER BY CASE r.name WHEN 'Supervisor' THEN 1 WHEN 'Chief Meteorologist' THEN 2 ELSE 3 END,u.full_name""").fetchall()
        return render_template("personnel_clause.html",clause=clause,personnel=personnel,
            document_types=PERSONNEL_DOCUMENT_TYPES,selected=selected,reviewers=reviewers,
            can_manage_all=can_manage_all,pending_registrations=pending_registrations,
            page_title="Clause 6.2 · Personnel",page_subtitle=clause[3],active_nav="iso17025")
    finally:
        connection.close()


@app.post("/iso17025/clause-checklist/<int:clause_id>/personnel/register")
def personnel_register(clause_id):
    connection = get_connection()
    try:
        full_name = request.form.get("full_name","").strip()
        email = request.form.get("email","").strip() or None
        requested_role = request.form.get("requested_role","").strip()
        allowed = {"Supervisor","Engineer","Technician","Operator","Guest"}
        if not full_name or requested_role not in allowed: raise ValueError("Full name and a valid laboratory role are required.")
        if email and not valid_email(email): raise ValueError("Enter a valid email address.")
        chief = connection.execute("""SELECT u.id,u.full_name FROM users u JOIN roles r ON r.id=u.role_id
            WHERE u.is_active=1 AND r.name='Chief Meteorologist' ORDER BY u.id LIMIT 1""").fetchone()
        if not chief: raise ValueError("An active Chief Meteorologist is required to review personnel registration.")
        registration_id = connection.execute("""INSERT INTO personnel_registrations
            (full_name,email,requested_role,status,initiated_by_user_id,initiated_by,assigned_reviewer_id)
            VALUES (?,?,?,'Pending',?,?,?)""",(full_name,email,requested_role,session.get("user_id"),current_actor(),chief[0])).lastrowid
        initiated_by_chief = session.get("role_name") == "Chief Meteorologist"
        if not initiated_by_chief:
            connection.execute("""INSERT INTO workflow_tasks(entity_type,entity_id,task_type,status,submitted_by,
                assigned_to,submitted_user_id,assigned_user_id,priority) VALUES (?,?,'Personnel registration','Pending',?,?,?,?,?)""",
                ("personnel_registration",registration_id,current_actor(),chief[1],session.get("user_id"),chief[0],"Normal"))
            add_notification(connection,chief[0],"Personnel registration awaiting completion",full_name,
                url_for("users",registration_id=registration_id),"personnel_registration",registration_id)
        audit_change(connection,"personnel_registration",registration_id,"submit",after={"name":full_name,"role":requested_role})
        connection.commit()
        flash("Complete the personnel account details." if initiated_by_chief else
            "Personnel registration sent to the Chief Meteorologist for completion.","success")
        destination = url_for("users",registration_id=registration_id) if initiated_by_chief else url_for("personnel_clause",clause_id=clause_id)
    except ValueError as error:
        connection.rollback(); flash(str(error),"error")
    finally:
        connection.close()
    return redirect(destination if 'destination' in locals() else url_for("personnel_clause",clause_id=clause_id))


@app.post("/iso17025/personnel/<int:user_id>/documents")
def personnel_document_upload(user_id):
    connection = get_connection()
    try:
        target = connection.execute("SELECT id,full_name FROM users WHERE id=? AND is_active=1",(user_id,)).fetchone()
        if not target: raise ValueError("Active personnel record was not found.")
        role = session.get("role_name")
        if user_id != session.get("user_id") and role not in {"Supervisor","Chief Meteorologist","Administrator"}:
            abort(403)
        document_type = request.form.get("document_type","")
        document_number = request.form.get("document_number","").strip()
        revision = request.form.get("revision","").strip()
        issued_date = request.form.get("issued_date") or None
        upload = request.files.get("document_file")
        if document_type not in PERSONNEL_DOCUMENT_TYPES or not all((document_number,revision,issued_date,upload and upload.filename)):
            raise ValueError("Document type, number, revision, issued date and file are required.")
        ensure_unique_document_number(connection,document_number)
        reviewer_id = request.form.get("reviewer_user_id",type=int)
        auto_approved = role in {"Chief Meteorologist","Administrator"}
        if not auto_approved:
            permitted_roles = ("Chief Meteorologist","Administrator") if role == "Supervisor" else ("Supervisor","Chief Meteorologist","Administrator")
            reviewer = connection.execute(f"""SELECT u.id,u.full_name,r.name FROM users u JOIN roles r ON r.id=u.role_id
                WHERE u.id=? AND u.is_active=1 AND r.name IN ({','.join('?' for _ in permitted_roles)})""",
                (reviewer_id,*permitted_roles)).fetchone()
            if not reviewer: raise ValueError("Select an authorized reviewer for this personnel document.")
        else: reviewer = None
        previous = connection.execute("""SELECT id FROM personnel_documents WHERE user_id=? AND document_type=?
            AND status='Approved' AND is_current=1 AND is_deleted=0 ORDER BY id DESC LIMIT 1""",(user_id,document_type)).fetchone()
        open_version = connection.execute("""SELECT id FROM personnel_documents WHERE user_id=? AND document_type=?
            AND status IN ('Submitted for Review','Reverted') AND is_deleted=0 ORDER BY id DESC LIMIT 1""",
            (user_id,document_type)).fetchone()
        if open_version:
            raise ValueError("An editable version of this personnel record already exists. Open and edit that version instead of adding a duplicate.")
        file_name,file_path = _save_personnel_document(upload,user_id)
        status = "Approved" if auto_approved else "Submitted for Review"
        document_id = connection.execute("""INSERT INTO personnel_documents
            (user_id,document_type,document_number,revision,issued_date,file_name,file_path,status,
            uploaded_by_user_id,uploaded_by,assigned_reviewer_id,approved_by_user_id,approved_by,approved_at,
            supersedes_id,is_current) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,CASE WHEN ?='Approved' THEN CURRENT_TIMESTAMP ELSE NULL END,?,?)""",
            (user_id,document_type,document_number,revision,issued_date,file_name,file_path,status,
             session.get("user_id"),current_actor(),reviewer[0] if reviewer else None,
             session.get("user_id") if auto_approved else None,current_actor() if auto_approved else None,status,
             previous[0] if previous else None,int(auto_approved))).lastrowid
        if auto_approved and previous:
            connection.execute("UPDATE personnel_documents SET is_current=0,superseded_by_id=? WHERE id=?",(document_id,previous[0]))
        if reviewer:
            connection.execute("""INSERT INTO workflow_tasks(entity_type,entity_id,task_type,status,submitted_by,
                assigned_to,submitted_user_id,assigned_user_id,priority) VALUES (?,?,'Personnel document review','Pending',?,?,?,?,?)""",
                ("personnel_document",document_id,current_actor(),reviewer[1],session.get("user_id"),reviewer[0],"Normal"))
            add_notification(connection,reviewer[0],"Personnel document awaiting review",
                f"{target[1]} · {document_type}",url_for("personnel_record",user_id=user_id),"personnel_document",document_id)
        audit_change(connection,"personnel_document",document_id,"upload",after={"person":target[1],"type":document_type,"status":status})
        connection.commit(); flash("Personnel document approved." if auto_approved else "Personnel document submitted for review.","success")
    except ValueError as error:
        connection.rollback(); flash(str(error),"error")
    finally:
        connection.close()
    return redirect(url_for("personnel_record",user_id=user_id))


def _personnel_document_reviewer(connection, reviewer_id, uploader_role):
    permitted = ("Chief Meteorologist","Administrator") if uploader_role == "Supervisor" else (
        "Supervisor","Chief Meteorologist","Administrator")
    reviewer = connection.execute(f"""SELECT u.id,u.full_name,r.name FROM users u JOIN roles r ON r.id=u.role_id
        WHERE u.id=? AND u.is_active=1 AND r.name IN ({','.join('?' for _ in permitted)})""",
        (reviewer_id,*permitted)).fetchone()
    if not reviewer: raise ValueError("Select an authorized reviewer.")
    return reviewer


@app.post("/iso17025/personnel-documents/<int:document_id>/edit")
def personnel_document_edit(document_id):
    connection = get_connection()
    try:
        row = connection.execute("""SELECT d.id,d.user_id,d.document_type,d.document_number,d.revision,d.issued_date,
            d.file_name,d.file_path,d.status,d.uploaded_by_user_id,d.assigned_reviewer_id,d.is_current,r.name
            FROM personnel_documents d JOIN users u ON u.id=d.uploaded_by_user_id JOIN roles r ON r.id=u.role_id
            WHERE d.id=? AND d.is_deleted=0""",(document_id,)).fetchone()
        if not row: abort(404)
        role = session.get("role_name")
        if session.get("user_id") not in {row[1],row[9]} and role not in {"Supervisor","Chief Meteorologist","Administrator"}:
            abort(403)
        document_number=request.form.get("document_number","").strip()
        revision=request.form.get("revision","").strip()
        issued_date=request.form.get("issued_date") or None
        if not all((document_number,revision,issued_date)): raise ValueError("Document number, revision and issued date are required.")
        ensure_unique_document_number(connection,document_number,"personnel_documents",document_id)
        upload=request.files.get("document_file")
        file_name,file_path=(row[6],row[7])
        if upload and upload.filename: file_name,file_path=_save_personnel_document(upload,row[1])
        auto_approved=role in {"Chief Meteorologist","Administrator"}
        reviewer=None
        if not auto_approved:
            reviewer=_personnel_document_reviewer(connection,request.form.get("reviewer_user_id",type=int),role)
        if row[8]=="Approved":
            status="Approved" if auto_approved else "Submitted for Review"
            new_id=connection.execute("""INSERT INTO personnel_documents
                (user_id,document_type,document_number,revision,issued_date,file_name,file_path,status,
                uploaded_by_user_id,uploaded_by,assigned_reviewer_id,approved_by_user_id,approved_by,approved_at,
                supersedes_id,is_current) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,CASE WHEN ?='Approved' THEN CURRENT_TIMESTAMP ELSE NULL END,?,?)""",
                (row[1],row[2],document_number,revision,issued_date,file_name,file_path,status,session.get("user_id"),
                 current_actor(),reviewer[0] if reviewer else None,session.get("user_id") if auto_approved else None,
                 current_actor() if auto_approved else None,status,row[0],int(auto_approved))).lastrowid
            if auto_approved:
                connection.execute("UPDATE personnel_documents SET is_current=0,superseded_by_id=? WHERE id=?",(new_id,row[0]))
            document_id=new_id
        else:
            status="Approved" if auto_approved else "Submitted for Review"
            connection.execute("""UPDATE personnel_documents SET document_number=?,revision=?,issued_date=?,file_name=?,file_path=?,
                status=?,assigned_reviewer_id=?,review_comment=NULL,approved_by_user_id=?,approved_by=?,
                approved_at=CASE WHEN ?='Approved' THEN CURRENT_TIMESTAMP ELSE NULL END,is_current=? WHERE id=?""",
                (document_number,revision,issued_date,file_name,file_path,status,reviewer[0] if reviewer else None,
                 session.get("user_id") if auto_approved else None,current_actor() if auto_approved else None,
                 status,int(auto_approved),document_id))
        if not auto_approved:
            connection.execute("""UPDATE workflow_tasks SET status='Cancelled',decision='superseded by edit',reviewed_at=CURRENT_TIMESTAMP
                WHERE entity_type='personnel_document' AND entity_id=? AND status='Pending'""",(row[0],))
            connection.execute("""INSERT INTO workflow_tasks(entity_type,entity_id,task_type,status,submitted_by,
                assigned_to,submitted_user_id,assigned_user_id,priority) VALUES (?,?,'Personnel document review','Pending',?,?,?,?,?)""",
                ("personnel_document",document_id,current_actor(),reviewer[1],session.get("user_id"),reviewer[0],"Normal"))
            add_notification(connection,reviewer[0],"Revised personnel document awaiting review",row[2],
                url_for("personnel_record",user_id=row[1]),"personnel_document",document_id)
        audit_change(connection,"personnel_document",document_id,"edit",after={"revision":revision,"status":status})
        connection.commit();flash("Personnel document updated" + (" and approved." if auto_approved else " and submitted for review."),"success")
    except ValueError as error:
        connection.rollback();flash(str(error),"error")
    finally: connection.close()
    return redirect(url_for("personnel_record",user_id=row[1] if 'row' in locals() and row else session.get("user_id")))


@app.post("/iso17025/personnel-documents/<int:document_id>/delete")
def personnel_document_delete(document_id):
    if session.get("role_name") not in {"Chief Meteorologist","Administrator"}: abort(403)
    connection=get_connection()
    try:
        row=connection.execute("SELECT user_id,document_type,is_current,supersedes_id FROM personnel_documents WHERE id=? AND is_deleted=0",(document_id,)).fetchone()
        if not row: abort(404)
        connection.execute("UPDATE personnel_documents SET is_deleted=1,is_current=0 WHERE id=?",(document_id,))
        connection.execute("""UPDATE workflow_tasks SET status='Cancelled',decision='deleted',reviewed_at=CURRENT_TIMESTAMP
            WHERE entity_type='personnel_document' AND entity_id=? AND status='Pending'""",(document_id,))
        if row[2]:
            previous=connection.execute("""SELECT id FROM personnel_documents WHERE user_id=? AND document_type=?
                AND status='Approved' AND is_deleted=0 AND id<>? ORDER BY id DESC LIMIT 1""",(row[0],row[1],document_id)).fetchone()
            if previous:
                connection.execute("UPDATE personnel_documents SET is_current=1,superseded_by_id=NULL WHERE id=?",(previous[0],))
        audit_change(connection,"personnel_document",document_id,"delete",before={"type":row[1]},after={"is_deleted":1})
        connection.commit();flash("Personnel document deleted. The most recent approved predecessor was restored where available.","success")
    finally: connection.close()
    return redirect(url_for("personnel_record",user_id=row[0]))


@app.get("/iso17025/personnel/<int:user_id>")
def personnel_record(user_id):
    clause = query("SELECT id FROM iso_clauses WHERE clause_code='6.2'")
    if not clause: abort(404)
    return redirect(url_for("personnel_clause",clause_id=clause[0][0],user_id=user_id))


@app.post("/iso17025/personnel-documents/<int:document_id>/workflow")
def personnel_document_workflow(document_id):
    connection = get_connection()
    try:
        row = connection.execute("""SELECT d.id,d.user_id,d.document_type,d.status,d.uploaded_by_user_id,
            d.assigned_reviewer_id,d.supersedes_id,r.name,u.full_name FROM personnel_documents d
            JOIN users u ON u.id=d.uploaded_by_user_id JOIN roles r ON r.id=u.role_id
            WHERE d.id=? AND d.is_deleted=0""",(document_id,)).fetchone()
        if not row or row[3] != "Submitted for Review": raise ValueError("Only submitted personnel documents can be reviewed.")
        role = session.get("role_name")
        if role not in {"Supervisor","Chief Meteorologist","Administrator"}: abort(403)
        if row[7] == "Supervisor" and role not in {"Chief Meteorologist","Administrator"}:
            raise ValueError("A document changed by a Supervisor must be approved by the Chief Meteorologist.")
        if session.get("user_id") != row[5] and role not in {"Chief Meteorologist","Administrator"}:
            raise ValueError("This review is assigned to another authorized reviewer.")
        action = request.form.get("action")
        comment = request.form.get("comment","").strip()
        if action not in {"approve","revert"}: raise ValueError("Select Approve or Revert.")
        if action == "revert" and not comment: raise ValueError("A revert comment is required.")
        status = "Approved" if action == "approve" else "Reverted"
        if status == "Approved":
            connection.execute("""UPDATE personnel_documents SET is_current=0,superseded_by_id=?
                WHERE user_id=? AND document_type=? AND is_current=1 AND status='Approved' AND id<>?""",
                (document_id,row[1],row[2],document_id))
        connection.execute("""UPDATE personnel_documents SET status=?,review_comment=?,approved_by_user_id=?,approved_by=?,
            approved_at=CASE WHEN ?='Approved' THEN CURRENT_TIMESTAMP ELSE NULL END,is_current=? WHERE id=?""",
            (status,comment or None,session.get("user_id") if status=="Approved" else None,
             current_actor() if status=="Approved" else None,status,int(status=="Approved"),document_id))
        connection.execute("""UPDATE workflow_tasks SET status='Completed',decision=?,comment=?,reviewed_at=CURRENT_TIMESTAMP
            WHERE entity_type='personnel_document' AND entity_id=? AND status='Pending'""",(action,comment or None,document_id))
        add_notification(connection,row[4],f"Personnel document {status.lower()}",
            comment or f"{row[2]} was {status.lower()} by {current_actor()}.",url_for("personnel_record",user_id=row[1]),"personnel_document",document_id)
        audit_change(connection,"personnel_document",document_id,action,after={"status":status,"comment":comment or None})
        connection.commit(); flash(f"Personnel document {status.lower()}.","success")
    except ValueError as error:
        connection.rollback(); flash(str(error),"error")
    finally:
        connection.close()
    return redirect(url_for("personnel_record",user_id=row[1] if 'row' in locals() and row else session.get("user_id")))


@app.get("/iso17025/personnel-documents/<int:document_id>/file")
def personnel_document_file(document_id):
    connection = get_connection()
    try:
        row = connection.execute("SELECT user_id,file_path,file_name FROM personnel_documents WHERE id=? AND is_deleted=0",(document_id,)).fetchone()
        if not row: abort(404)
        if row[0] != session.get("user_id") and session.get("role_name") not in {"Supervisor","Chief Meteorologist","Administrator"}:
            abort(403)
        path = Path(row[1])
        if not path.is_file(): abort(404)
        return send_file(path.resolve(),download_name=row[2],as_attachment=False)
    finally:
        connection.close()


@app.post("/iso17025/clause-checklist/<int:clause_id>/evidence")
def iso_clause_evidence_upload(clause_id):
    require_permission("equipment_edit")
    connection=get_connection()
    try:
        clause=connection.execute("""SELECT c.clause_code FROM iso_clauses c JOIN iso_clause_assessments a
            ON a.clause_id=c.id WHERE c.id=? AND a.applicability='Applicable'
            AND a.applicability_set_by IS NOT NULL""",(clause_id,)).fetchone()
        if not clause: abort(404)
        upload=request.files.get("evidence_file")
        requirement_id=request.form.get("requirement_id",type=int)
        requirement=connection.execute("SELECT evidence_name FROM iso_clause_evidence_requirements WHERE id=? AND clause_id=? AND is_active=1",(requirement_id,clause_id)).fetchone() if requirement_id else None
        if not requirement: raise ValueError("Select a document or record required for this clause.")
        title=requirement[0]
        document_owner=request.form.get("document_owner","").strip() or current_actor()
        document_number=request.form.get("document_number","").strip()
        revision=request.form.get("revision","").strip()
        effective_date=request.form.get("effective_date") or None
        retention_until=request.form.get("retention_until") or None
        if not upload or not upload.filename or not all((document_number,revision,effective_date,retention_until)):
            raise ValueError("Document/record number, revision, issued date, retention date and file are required.")
        ensure_unique_document_number(connection,document_number)
        file_name,file_path=_save_clause_evidence(upload,clause[0])
        status="Approved" if chief_auto_approval() else "Draft"
        evidence_id=connection.execute("""INSERT INTO iso_clause_evidence
            (clause_id,requirement_id,title,document_number,revision,effective_date,review_date,
            retention_until,document_owner,file_name,file_path,status,uploaded_by,approved_by,approved_at)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,CASE WHEN ?='Approved' THEN CURRENT_TIMESTAMP ELSE NULL END)""",
            (clause_id,requirement_id,title,document_number,
             revision,effective_date,
             request.form.get("review_date") or None,request.form.get("retention_until") or None,
             document_owner,file_name,file_path,status,current_actor(),current_actor() if status=="Approved" else None,status)).lastrowid
        audit_change(connection,"iso_clause_evidence",evidence_id,"upload",after={"title":title,"status":status})
        _sync_clause_compliance(connection, clause_id)
        connection.commit();flash("Evidence uploaded" + (" and approved." if status=="Approved" else " as Draft."),"success")
    except Exception as error:
        connection.rollback();flash(str(error),"error")
    finally: connection.close()
    return redirect(url_for("iso_clause_detail",clause_id=clause_id))


@app.post("/iso17025/clause-evidence/<int:evidence_id>/workflow")
def iso_clause_evidence_workflow(evidence_id):
    connection=get_connection()
    try:
        row=connection.execute("SELECT clause_id,status,uploaded_by,title FROM iso_clause_evidence WHERE id=? AND is_deleted=0",(evidence_id,)).fetchone()
        if not row: abort(404)
        action=request.form.get("action")
        comment=request.form.get("comment","").strip()
        if action=="submit":
            reviewer_id=request.form.get("reviewer_user_id",type=int)
            reviewer=connection.execute("""SELECT u.id,u.full_name FROM users u JOIN roles r ON r.id=u.role_id
                WHERE u.id=? AND u.is_active=1 AND r.name IN ('Chief Meteorologist','Administrator')""",(reviewer_id,)).fetchone()
            if not reviewer: raise ValueError("Select an authorized reviewer.")
            connection.execute("UPDATE iso_clause_evidence SET status='Submitted for Review',assigned_reviewer_id=?,submitted_at=CURRENT_TIMESTAMP WHERE id=?",(reviewer_id,evidence_id))
            connection.execute("""INSERT INTO workflow_tasks(entity_type,entity_id,task_type,status,submitted_by,
                assigned_to,submitted_user_id,assigned_user_id,priority) VALUES (?,?,?,'Pending',?,?,?,?,?)""",
                ("iso_clause_evidence",evidence_id,"Clause evidence review",current_actor(),reviewer[1],session.get("user_id"),reviewer_id,"Normal"))
            add_notification(connection,reviewer_id,"Clause evidence awaiting review",row[3],url_for("iso_clause_detail",clause_id=row[0]),"iso_clause_evidence",evidence_id)
            status="Submitted for Review"
        else:
            require_permission("record_approve")
            if session.get("role_name") not in {"Chief Meteorologist","Administrator"}: raise ValueError("Only the Chief Meteorologist or Administrator can approve clause records.")
            if row[1]!="Submitted for Review": raise ValueError("Only submitted evidence can be reviewed.")
            if action=="revert" and not comment: raise ValueError("A revert comment is required.")
            if action not in {"approve","revert"}: raise ValueError("Select Approve or Revert.")
            status="Approved" if action=="approve" else "Reverted"
            connection.execute("""UPDATE iso_clause_evidence SET status=?,review_comment=?,approved_by=?,
                approved_at=CASE WHEN ?='Approved' THEN CURRENT_TIMESTAMP ELSE NULL END WHERE id=?""",
                (status,comment or None,current_actor() if status=="Approved" else None,status,evidence_id))
            connection.execute("""UPDATE workflow_tasks SET status='Completed',decision=?,comment=?,reviewed_at=CURRENT_TIMESTAMP
                WHERE entity_type='iso_clause_evidence' AND entity_id=? AND status='Pending'""",(action,comment or None,evidence_id))
            submitter_id=user_id_for_actor(connection,row[2])
            add_notification(connection,submitter_id,f"Clause evidence {status.lower()}",
                comment or f"{row[3]} was {status.lower()} by {current_actor()}.",
                url_for("iso_clause_detail",clause_id=row[0]),"iso_clause_evidence",evidence_id)
        audit_change(connection,"iso_clause_evidence",evidence_id,action,after={"status":status,"comment":comment or None})
        _sync_clause_compliance(connection, row[0])
        connection.commit();flash(f"Evidence status changed to {status}.","success")
    except ValueError as error:
        connection.rollback();flash(str(error),"error")
    finally: connection.close()
    return redirect(url_for("iso_clause_detail",clause_id=row[0] if 'row' in locals() and row else 1))


@app.get("/iso17025/clause-evidence/<int:evidence_id>/file")
def iso_clause_evidence_file(evidence_id):
    rows=query("SELECT file_path,file_name FROM iso_clause_evidence WHERE id=? AND is_deleted=0",(evidence_id,))
    if not rows or not Path(rows[0][0]).is_file(): abort(404)
    return send_file(Path(rows[0][0]).resolve(),download_name=rows[0][1],as_attachment=False)


@app.post("/iso17025/clause-checklist/<int:clause_id>/workflow")
def iso_clause_assessment_workflow(clause_id):
    connection=get_connection()
    try:
        row=connection.execute("SELECT id,status,compliance_status,created_by FROM iso_clause_assessments WHERE clause_id=?",(clause_id,)).fetchone()
        if not row: raise ValueError("Save the clause assessment before submitting it.")
        action=request.form.get("action");comment=request.form.get("comment","").strip()
        if action in {"submit", "approve"}:
            missing_rows=connection.execute("""SELECT r.evidence_name FROM iso_clause_evidence_requirements r
                WHERE r.clause_id=? AND r.is_required=1 AND r.is_active=1 AND NOT EXISTS (SELECT 1 FROM iso_clause_evidence e
                WHERE e.requirement_id=r.id AND e.status='Approved' AND e.is_deleted=0)
                ORDER BY r.id""",(clause_id,)).fetchall()
            if missing_rows:
                missing_names=", ".join(item[0] for item in missing_rows)
                raise ValueError(f"Approval is blocked. Upload and approve all required evidence first: {missing_names}.")
        if action=="submit":
            if chief_auto_approval():
                status="Approved";connection.execute("UPDATE iso_clause_assessments SET status=?,approved_by=?,approved_at=CURRENT_TIMESTAMP WHERE id=?",(status,current_actor(),row[0]))
            else:
                reviewer_id=request.form.get("reviewer_user_id",type=int)
                reviewer=connection.execute("""SELECT u.id,u.full_name FROM users u JOIN roles r ON r.id=u.role_id
                    WHERE u.id=? AND u.is_active=1 AND r.name IN ('Chief Meteorologist','Administrator','Supervisor')""",(reviewer_id,)).fetchone()
                if not reviewer: raise ValueError("Select an authorized reviewer.")
                status="Submitted for Review";connection.execute("""UPDATE iso_clause_assessments SET status=?,assigned_reviewer_id=?,submitted_by=?,submitted_at=CURRENT_TIMESTAMP WHERE id=?""",(status,reviewer_id,current_actor(),row[0]))
                connection.execute("""INSERT INTO workflow_tasks(entity_type,entity_id,task_type,status,submitted_by,assigned_to,
                    submitted_user_id,assigned_user_id,priority) VALUES (?,?,?,'Pending',?,?,?,?,?)""",
                    ("iso_clause_assessment",clause_id,"ISO clause assessment review",current_actor(),reviewer[1],session.get("user_id"),reviewer_id,"Normal"))
                add_notification(connection,reviewer_id,"ISO clause assessment awaiting review",
                    f"Clause {clause_id} was submitted by {current_actor()}.",
                    url_for("iso_clause_detail",clause_id=clause_id),"iso_clause_assessment",clause_id)
        else:
            require_permission("record_approve")
            if row[1]!="Submitted for Review": raise ValueError("Only submitted assessments can be reviewed.")
            if action=="revert" and not comment: raise ValueError("A revert comment is required.")
            if action not in {"approve","revert"}: raise ValueError("Select Approve or Revert.")
            status="Approved" if action=="approve" else "Reverted"
            connection.execute("""UPDATE iso_clause_assessments SET status=?,review_comment=?,approved_by=?,
                approved_at=CASE WHEN ?='Approved' THEN CURRENT_TIMESTAMP ELSE NULL END WHERE id=?""",
                (status,comment or None,current_actor() if status=="Approved" else None,status,row[0]))
            connection.execute("""UPDATE workflow_tasks SET status='Completed',decision=?,comment=?,reviewed_at=CURRENT_TIMESTAMP
                WHERE entity_type='iso_clause_assessment' AND entity_id=? AND status='Pending'""",(action,comment or None,clause_id))
            submitter_id=user_id_for_actor(connection,row[3])
            add_notification(connection,submitter_id,f"ISO clause assessment {status.lower()}",
                comment or f"The assessment was {status.lower()} by {current_actor()}.",
                url_for("iso_clause_detail",clause_id=clause_id),"iso_clause_assessment",clause_id)
        audit_change(connection,"iso_clause_assessment",row[0],action,after={"status":status,"comment":comment or None})
        connection.commit();flash(f"Assessment status changed to {status}.","success")
    except ValueError as error:
        connection.rollback();flash(str(error),"error")
    finally: connection.close()
    return redirect(url_for("iso_clause_detail",clause_id=clause_id))


@app.post("/quality-records/<int:record_id>/delete")
def quality_record_delete(record_id):
    require_permission("equipment_delete")
    section = request.form.get("section", "iso17025")
    module = request.form.get("module", "")
    connection = get_connection()
    record = connection.execute("SELECT record_type,title,status FROM quality_records WHERE id=? AND is_deleted=0",
        (record_id,)).fetchone()
    if not record:
        connection.close(); abort(404)
    connection.execute("""UPDATE quality_records SET is_deleted=1,deleted_at=CURRENT_TIMESTAMP,
        deleted_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""", (current_actor(), record_id))
    audit_change(connection, "quality_record", record_id, "soft_delete",
        before={"record_type": record[0], "title": record[1], "status": record[2]},
        after={"is_deleted": 1})
    connection.commit(); connection.close()
    flash("Record removed. Its audit history was preserved.", "success")
    return redirect(url_for("module_page", section=section, module=module or record[0]))


@app.route("/settings", methods=["GET", "POST"])
def settings():
    if request.method == "POST":
        require_permission("settings")
        raw_days = request.form.get("equipment_due_soon_days", "").strip()
        existing_recovery_email = setting_text(
            "admin_recovery_email", "ikechukwuumezulike@gmail.com")
        recovery_email = (request.form.get("admin_recovery_email", "").strip()
            if session.get("role_name") == "Administrator" else existing_recovery_email)
        try:
            days = int(raw_days)
            if not 0 <= days <= 365:
                raise ValueError
            if not valid_email(recovery_email):
                raise ValueError
        except ValueError:
            flash("Enter a whole-number warning period from 0 to 365 and a valid Administrator recovery email.", "error")
        else:
            connection = get_connection()
            before = connection.execute("SELECT setting_value FROM system_settings WHERE setting_key='equipment_due_soon_days'").fetchone()
            connection.execute("""INSERT INTO system_settings
                (setting_key,setting_value,value_type,updated_by,updated_at)
                VALUES ('equipment_due_soon_days',?,'integer',?,CURRENT_TIMESTAMP)
                ON CONFLICT(setting_key) DO UPDATE SET setting_value=excluded.setting_value,
                updated_by=excluded.updated_by,updated_at=CURRENT_TIMESTAMP""", (str(days), current_actor()))
            before_email = connection.execute("""SELECT setting_value FROM system_settings
                WHERE setting_key='admin_recovery_email'""").fetchone()
            connection.execute("""INSERT INTO system_settings
                (setting_key,setting_value,value_type,updated_by,updated_at)
                VALUES ('admin_recovery_email',?,'email',?,CURRENT_TIMESTAMP)
                ON CONFLICT(setting_key) DO UPDATE SET setting_value=excluded.setting_value,
                updated_by=excluded.updated_by,updated_at=CURRENT_TIMESTAMP""",
                (recovery_email,current_actor()))
            audit_change(connection, "system_setting", None, "update",
                before={"equipment_due_soon_days": before[0] if before else None,
                    "admin_recovery_email":before_email[0] if before_email else None},
                after={"equipment_due_soon_days": days,
                    "admin_recovery_email":recovery_email})
            connection.commit(); connection.close()
            flash("Settings saved.", "success")
            return redirect(url_for("settings"))
    return render_template("settings.html",
        due_soon_days=setting_integer("equipment_due_soon_days", 30),
        admin_recovery_email=setting_text("admin_recovery_email", "ikechukwuumezulike@gmail.com"),
        email_delivery_configured=email_recovery_configured(),
        can_change_admin_recovery_email=session.get("role_name") == "Administrator",
        page_title="Settings", page_subtitle="", active_nav="settings")


@app.route("/projects/jobs", methods=["GET", "POST"])
def project_jobs():
    if request.method == "POST":
        action = request.form.get("action", "add_job")
        if action == "add_project":
            require_permission("project_create")
            connection = get_connection()
            try:
                project_client = request.form.get("project_client_name", "").strip()
                if not project_client:
                    raise ValueError("Project: select or enter a client.")
                customer = connection.execute("SELECT id FROM customers WHERE customer_name=? COLLATE NOCASE", (project_client,)).fetchone()
                if customer:
                    customer_id = customer[0]
                else:
                    require_permission("client_create")
                    customer_id = connection.execute("INSERT INTO customers(customer_name,status) VALUES (?,'Active')", (project_client,)).lastrowid
                    audit_change(connection, "customer", customer_id, "quick_create", after={"customer_name":project_client})
                project_name = request.form.get("project_name", "").strip()
                if not project_name:
                    raise ValueError("Project: enter the project name.")
                project_number = request.form.get("project_number", "").strip() or next_reference(
                    connection, "projects", "project_number", "PRJ")
                cursor = connection.execute("""INSERT INTO projects
                    (project_number,customer_id,project_name,description,status,target_start_date,
                     target_completion_date,project_manager,notes,created_by,purchase_order)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?)""", (project_number, customer_id, project_name,
                    request.form.get("project_description", "").strip() or None,
                    request.form.get("project_status", "Draft"), request.form.get("project_start_date") or None,
                    request.form.get("project_completion_date") or None,
                    request.form.get("project_manager", "").strip() or None,
                    request.form.get("project_notes", "").strip() or None, current_actor(),
                    request.form.get("project_purchase_order", "").strip() or None))
                audit_change(connection, "project", cursor.lastrowid, "create", after={
                    "project_number": project_number, "project_name": project_name,
                    "status": request.form.get("project_status", "Draft")})
                connection.commit(); flash(f"Project {project_number} created.", "success")
                return redirect(url_for("project_jobs", project=cursor.lastrowid))
            except Exception as error:
                connection.rollback(); flash(str(error), "error")
            finally:
                connection.close()
        else:
            require_permission("calibration")
            required = ("job_title", "project_id", "mut_description", "meter_type",
                        "method_id", "required_date", "flow_min",
                        "flow_max", "flow_unit", "flow_point_count", "job_type",
                        "fluid_medium", "calibration_quantity")
            missing = [field.replace("_", " ").title() for field in required
                       if not request.form.get(field, "").strip()]
            if missing:
                flash("Required: " + ", ".join(missing), "error")
            else:
                attachment_path = None
                attachment = request.files.get("attachment")
                connection = get_connection()
                try:
                    project_id = request.form.get("project_id", type=int)
                    project = connection.execute("""SELECT p.customer_id,p.project_name,
                        COALESCE(p.purchase_order,''),c.customer_name FROM projects p
                        JOIN customers c ON c.id=p.customer_id WHERE p.id=? AND p.status<>'Closed' AND p.is_deleted=0""",
                        (project_id,)).fetchone()
                    if not project:
                        raise ValueError("Select an active parent Project for this Job.")
                    customer_id = project[0]
                    flow_min = float(request.form["flow_min"]); flow_max = float(request.form["flow_max"])
                    flow_point_count = int(request.form["flow_point_count"])
                    if flow_min > flow_max:
                        raise ValueError("Minimum flow cannot be greater than maximum flow.")
                    if not 1 <= flow_point_count <= 50:
                        raise ValueError("Number of flow points must be between 1 and 50.")
                    if attachment and attachment.filename:
                        validate_report(attachment); destination = Path("data/documents/jobs"); destination.mkdir(parents=True, exist_ok=True)
                        target = destination / f"{uuid4().hex}_{Path(attachment.filename).name}"; attachment.save(target); attachment_path = str(target)
                    job_number = request.form.get("job_number", "").strip() or next_reference(
                        connection, "calibration_jobs", "job_number", "JOB")
                    may_assign = app.config.get("TESTING") or Permissions.can(session.get("role_name", ""), "job_assign")
                    assigned_user_id = request.form.get("assigned_user_id", type=int) if may_assign else None
                    assigned_user = connection.execute("SELECT full_name FROM users WHERE id=? AND is_active=1",
                        (assigned_user_id,)).fetchone() if assigned_user_id else None
                    if may_assign and not assigned_user and not app.config.get("TESTING"):
                        raise ValueError("Chief Meteorologist or Supervisor must select an active assignee.")
                    assigned_name = assigned_user[0] if assigned_user else (
                        request.form.get("engineer_operator","").strip() if app.config.get("TESTING") else "Unassigned")
                    assignment_event = (json.dumps({"assigned_to":assigned_name,"assigned_by":current_actor(),
                        "date":date.today().isoformat()}) + "\n") if assigned_user else None
                    cursor = connection.execute("""INSERT INTO calibration_jobs
                        (job_number,job_title,job_type,customer_id,equipment_id,method_id,status,flow_range,
                         meter_type,mut_description,mut_asset_id,mut_serial_number,manufacturer,model_number,
                         meter_size,engineer_operator,planned_start_date,required_date,flow_min,flow_max,
                         flow_unit,flow_point_count,fluid_medium,calibration_quantity,attachment_path,
                         created_by,notes,attachment_uploaded_by,assigned_user_id,assigned_by,assigned_at,
                         assignment_history,purchase_order_override)
                         VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
                        job_number, request.form["job_title"].strip(), request.form["job_type"].strip(), customer_id,
                        int(request.form["equipment_id"]) if request.form.get("equipment_id") else None,
                        int(request.form["method_id"]), request.form.get("status", "Draft"),
                        f'{request.form["flow_min"]}–{request.form["flow_max"]} {request.form["flow_unit"]}',
                        request.form["meter_type"].strip(), request.form["mut_description"].strip(),
                        request.form.get("mut_asset_id", "").strip() or None,
                        request.form.get("mut_serial_number", "").strip() or None,
                        request.form.get("manufacturer", "").strip() or None,
                        request.form.get("model_number", "").strip() or None,
                        request.form.get("meter_size", "").strip() or None,
                        assigned_name,
                        request.form.get("planned_start_date") or None,
                        request.form["required_date"], flow_min, flow_max, request.form["flow_unit"],
                        flow_point_count, request.form["fluid_medium"].strip(), request.form["calibration_quantity"],
                        attachment_path, current_actor(), request.form.get("notes", "").strip() or None,
                        current_actor() if attachment_path else None,assigned_user_id,
                        current_actor() if assigned_user else None,None,assignment_event,
                        request.form.get("job_purchase_order", "").strip() or None))
                    if assigned_user:
                        connection.execute("UPDATE calibration_jobs SET assigned_at=CURRENT_TIMESTAMP WHERE id=?",
                            (cursor.lastrowid,))
                    connection.execute("INSERT INTO project_jobs (project_id,job_id) VALUES (?,?)",
                        (project_id, cursor.lastrowid))
                    add_note(connection, "calibration_job", cursor.lastrowid, request.form.get("notes"))
                    audit_change(connection, "calibration_job", cursor.lastrowid, "create", after={
                        "job_number": job_number, "project_id": project_id or None,
                        "status": request.form.get("status", "Draft")})
                    connection.commit(); flash(f"Job {job_number} created.", "success")
                    return redirect(url_for("job_detail", job_id=cursor.lastrowid))
                except Exception as error:
                    connection.rollback(); flash(str(error), "error")
                finally:
                    connection.close()
    customers = query("SELECT id, customer_name FROM customers WHERE status='Active' ORDER BY customer_name")
    projects = query("""SELECT p.id,p.project_number,p.project_name,COALESCE(p.purchase_order,''),
        c.customer_name FROM projects p JOIN customers c ON c.id=p.customer_id
        WHERE p.status<>'Closed' AND p.is_deleted=0 ORDER BY p.project_number""")
    equipment_rows = query("SELECT id, equipment_name, asset_number, serial_number FROM laboratory_equipment WHERE is_active=1 ORDER BY equipment_name")
    methods = query("SELECT id, method_name, revision, meter_type FROM calibration_methods WHERE is_active=1 OR status='Approved' ORDER BY method_name")
    assignees = query("""SELECT u.id,u.full_name,r.name FROM users u JOIN roles r ON r.id=u.role_id
        WHERE u.is_active=1 ORDER BY CASE r.name WHEN 'Chief Meteorologist' THEN 1 WHEN 'Supervisor' THEN 2
        WHEN 'Engineer' THEN 3 WHEN 'Technician' THEN 4 WHEN 'Operator' THEN 5 WHEN 'Guest' THEN 6 ELSE 7 END,u.full_name""")
    selected_project = request.args.get("project", type=int)
    jobs_sql = """SELECT j.id,j.job_number,c.customer_name,j.mut_description,
        COALESCE(j.mut_serial_number,''),m.method_name,j.flow_range,
        COALESCE(j.flow_point_count,1),j.status,j.required_date,COALESCE(j.job_title,''),
        p.project_number,p.project_name,COALESCE(p.purchase_order,''),
        COALESCE(j.purchase_order_override,''),COALESCE(j.purchase_order_override,p.purchase_order,''),
        COALESCE(j.attachment_path,''),
        (SELECT u.id FROM uncertainty_calculations u WHERE u.job_id=j.id AND u.status='Approved'
         ORDER BY u.revision DESC,u.approved_at DESC,u.id DESC LIMIT 1)
        FROM calibration_jobs j LEFT JOIN customers c ON c.id=j.customer_id
        LEFT JOIN calibration_methods m ON m.id=j.method_id
        LEFT JOIN project_jobs pj ON pj.job_id=j.id LEFT JOIN projects p ON p.id=pj.project_id
        WHERE j.is_deleted=0 AND COALESCE(p.is_deleted,0)=0"""
    jobs_parameters = ()
    if selected_project is not None:
        jobs_sql += " AND p.id=?"
        jobs_parameters = (selected_project,)
    jobs = query(jobs_sql + " ORDER BY j.id DESC", jobs_parameters)
    numbering_connection = get_connection()
    next_job_number = next_reference(numbering_connection, "calibration_jobs", "job_number", "JOB")
    numbering_connection.close()
    selected_project_row = next((project for project in projects if project[0] == selected_project), None)
    return render_template("jobs.html", customers=customers, projects=projects,
        equipment_rows=equipment_rows, methods=methods, jobs=jobs, assignees=assignees,
        selected_project=selected_project_row,
        next_job_number=next_job_number,
        page_title="Projects · Calibration Jobs",
        page_subtitle="Shared project and MUT calibration workflow", active_nav="projects")


@app.get("/projects/jobs/<int:job_id>")
def job_detail(job_id):
    rows = query("""SELECT j.id,j.job_number,j.job_title,j.job_type,j.status,c.customer_name,
        p.id,p.project_number,p.project_name,j.mut_description,j.mut_asset_id,j.mut_serial_number,
        j.manufacturer,j.model_number,j.meter_type,j.meter_size,j.flow_min,j.flow_max,j.flow_unit,
        j.fluid_medium,j.calibration_quantity,m.method_name,j.engineer_operator,j.created_at,
        j.planned_start_date,j.required_date,j.notes,COALESCE(p.purchase_order,''),
        COALESCE(j.purchase_order_override,''),COALESCE(j.purchase_order_override,p.purchase_order,''),
        COALESCE(j.attachment_path,''),COALESCE(j.flow_point_count,1)
        FROM calibration_jobs j LEFT JOIN customers c ON c.id=j.customer_id
        LEFT JOIN project_jobs pj ON pj.job_id=j.id LEFT JOIN projects p ON p.id=pj.project_id
        LEFT JOIN calibration_methods m ON m.id=j.method_id
        WHERE j.id=? AND j.is_deleted=0""", (job_id,))
    if not rows:
        abort(404)
    calculations = query("""SELECT id,calculation_uid,method_name,status,revision,calculated_at,
        expanded_uncertainty,approved_by,approved_at FROM uncertainty_calculations
        WHERE job_id=? AND (status NOT IN ('Draft','Reverted') OR calculated_by=?)
        ORDER BY revision DESC,id DESC""", (job_id,current_actor()))
    approved_uncertainty_rows = query("""SELECT id,calculation_uid,method_name,status,revision,
        calculated_at,expanded_uncertainty,approved_by,approved_at
        FROM uncertainty_calculations WHERE job_id=? AND status='Approved'
        ORDER BY revision DESC,approved_at DESC,id DESC LIMIT 1""", (job_id,))
    approved_uncertainty = approved_uncertainty_rows[0] if approved_uncertainty_rows else None
    reopen_requests = query("""SELECT id,reason,status,requested_by,requested_at,authorized_by,
        authorized_at,decision_comment FROM job_reopen_requests WHERE job_id=? ORDER BY id DESC""", (job_id,))
    assignees = query("""SELECT u.id,u.full_name,r.name FROM users u JOIN roles r ON r.id=u.role_id
        WHERE u.is_active=1 AND r.name IN ('Chief Meteorologist','Supervisor','Engineer','Technician','Operator')
        ORDER BY u.full_name""")
    return render_template("job_detail.html", job=rows[0], calculations=calculations,
        approved_uncertainty=approved_uncertainty,
        reopen_requests=reopen_requests,assignees=assignees, page_title=f"Job · {rows[0][1]}",
        page_subtitle=rows[0][2] or rows[0][9], active_nav="projects")


@app.get("/projects/jobs/<int:job_id>/report")
def job_report_file(job_id):
    rows = query("SELECT attachment_path FROM calibration_jobs WHERE id=? AND is_deleted=0", (job_id,))
    if not rows or not rows[0][0] or not Path(rows[0][0]).is_file():
        abort(404)
    return send_file(Path(rows[0][0]).resolve(), as_attachment=False)


@app.post("/projects/jobs/<int:job_id>/assign")
def job_assign(job_id):
    require_permission("job_assign")
    user_id = request.form.get("assigned_user_id",type=int)
    reason = request.form.get("reason","").strip()
    connection = get_connection()
    job = connection.execute("SELECT engineer_operator,assigned_user_id FROM calibration_jobs WHERE id=? AND is_deleted=0",
        (job_id,)).fetchone()
    user = connection.execute("""SELECT u.full_name,r.name FROM users u JOIN roles r ON r.id=u.role_id
        WHERE u.id=? AND u.is_active=1 AND r.name IN ('Chief Meteorologist','Supervisor','Engineer','Technician','Operator')""",
        (user_id,)).fetchone()
    if not job or not user:
        connection.close(); flash("Select an active permitted assignee.","error")
        return redirect(url_for("job_detail",job_id=job_id))
    if job[1] and not reason:
        connection.close(); flash("A reason is required when reassigning a Job.","error")
        return redirect(url_for("job_detail",job_id=job_id))
    event = json.dumps({"assigned_to":user[0],"role":user[1],"assigned_by":current_actor(),
        "date":date.today().isoformat(),"reason":reason or None}) + "\n"
    connection.execute("""UPDATE calibration_jobs SET assigned_user_id=?,engineer_operator=?,assigned_by=?,
        assigned_at=CURRENT_TIMESTAMP,assignment_history=COALESCE(assignment_history,'') || ?,
        updated_at=CURRENT_TIMESTAMP WHERE id=?""",(user_id,user[0],current_actor(),event,job_id))
    audit_change(connection,"calibration_job",job_id,"reassign" if job[1] else "assign",
        before={"assigned_to":job[0]},after={"assigned_to":user[0],"assigned_by":current_actor(),"reason":reason or None})
    connection.commit(); connection.close(); flash(f"Job assigned to {user[0]}.","success")
    return redirect(url_for("job_detail",job_id=job_id))


@app.post("/projects/jobs/<int:job_id>/complete")
def job_complete(job_id):
    require_permission("calibration")
    connection = get_connection()
    row = connection.execute("SELECT job_number,status FROM calibration_jobs WHERE id=? AND is_deleted=0", (job_id,)).fetchone()
    if not row:
        connection.close(); abort(404)
    if not connection.execute("SELECT 1 FROM uncertainty_calculations WHERE job_id=? AND status='Approved'", (job_id,)).fetchone():
        connection.close(); flash("The Job cannot be completed until its required uncertainty calculation is approved.", "error")
        return redirect(url_for("job_detail", job_id=job_id))
    connection.execute("""UPDATE calibration_jobs SET status='Completed',completed_at=CURRENT_TIMESTAMP,
        completion_history=COALESCE(completion_history,'') || ?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
        (json.dumps({"action":"completed","actor":current_actor(),"timestamp":date.today().isoformat()}) + "\n", job_id))
    audit_change(connection, "calibration_job", job_id, "complete", before={"status":row[1]}, after={"status":"Completed"})
    connection.commit(); connection.close(); flash("Job completed and locked.", "success")
    return redirect(url_for("job_detail", job_id=job_id))


@app.post("/projects/jobs/<int:job_id>/request-reopen")
def job_request_reopen(job_id):
    require_permission("calibration")
    reason = request.form.get("reason", "").strip()
    if not reason:
        flash("A reason is required to request reopening a completed Job.", "error")
        return redirect(url_for("job_detail", job_id=job_id))
    connection = get_connection()
    job = connection.execute("SELECT status FROM calibration_jobs WHERE id=? AND is_deleted=0", (job_id,)).fetchone()
    if not job or job[0] != "Completed":
        connection.close(); flash("Only a completed Job can request HOD reopening authorization.", "error")
        return redirect(url_for("job_detail", job_id=job_id))
    if connection.execute("SELECT 1 FROM job_reopen_requests WHERE job_id=? AND status='Pending'", (job_id,)).fetchone():
        connection.close(); flash("A reopening request is already awaiting HOD authorization.", "error")
        return redirect(url_for("job_detail", job_id=job_id))
    cursor = connection.execute("INSERT INTO job_reopen_requests(job_id,reason,requested_by) VALUES (?,?,?)",
        (job_id, reason, current_actor()))
    audit_change(connection, "calibration_job", job_id, "request_reopen", after={"request_id":cursor.lastrowid,"reason":reason})
    connection.commit(); connection.close(); flash("Reopening request sent for HOD authorization.", "success")
    return redirect(url_for("job_detail", job_id=job_id))


@app.post("/projects/jobs/<int:job_id>/authorize-reopen")
def job_authorize_reopen(job_id):
    require_permission("job_reopen")
    request_id = request.form.get("request_id", type=int)
    comment = request.form.get("comment", "").strip()
    if not comment:
        flash("HOD authorization requires a decision comment.", "error")
        return redirect(url_for("job_detail", job_id=job_id))
    connection = get_connection()
    pending = connection.execute("SELECT id FROM job_reopen_requests WHERE id=? AND job_id=? AND status='Pending'",
        (request_id, job_id)).fetchone()
    if not pending:
        connection.close(); flash("The reopening request is no longer pending.", "error")
        return redirect(url_for("job_detail", job_id=job_id))
    connection.execute("""UPDATE job_reopen_requests SET status='Authorized',authorized_by=?,
        authorized_at=CURRENT_TIMESTAMP,decision_comment=? WHERE id=?""", (current_actor(), comment, request_id))
    connection.execute("UPDATE calibration_jobs SET status='In Progress',updated_at=CURRENT_TIMESTAMP WHERE id=?", (job_id,))
    audit_change(connection, "calibration_job", job_id, "authorize_reopen",
        before={"status":"Completed"}, after={"status":"In Progress","request_id":request_id,"comment":comment})
    connection.commit(); connection.close(); flash("HOD authorized the Job reopening. Status is now In Progress.", "success")
    return redirect(url_for("job_detail", job_id=job_id))


@app.post("/projects/jobs/<int:job_id>/delete")
def project_job_delete(job_id):
    require_permission("equipment_delete")
    connection = get_connection()
    job = connection.execute("SELECT job_number,status FROM calibration_jobs WHERE id=? AND is_deleted=0",
        (job_id,)).fetchone()
    if not job:
        connection.close(); abort(404)
    if job[1] in {"Approved", "Completed"}:
        connection.close()
        flash("Completed or approved Jobs cannot be deleted.", "error")
        return redirect(url_for("project_jobs"))
    connection.execute("""UPDATE calibration_jobs SET is_deleted=1,deleted_at=CURRENT_TIMESTAMP,
        deleted_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""", (current_actor(), job_id))
    recycle_record(connection, "calibration_job", job_id, job[0], "Projects & Jobs")
    audit_change(connection, "calibration_job", job_id, "soft_delete",
        before={"job_number": job[0], "status": job[1]}, after={"is_deleted": 1})
    connection.commit(); connection.close()
    flash("Job removed. Related calculations and audit history were preserved.", "success")
    return redirect(url_for("project_jobs"))


@app.post("/projects/<int:project_id>/delete")
def project_delete(project_id):
    if not app.config.get("TESTING") and session.get("role_name") not in {"Chief Meteorologist", "Administrator"}:
        abort(403)
    connection = get_connection()
    project = connection.execute("""SELECT project_number,project_name,status FROM projects
        WHERE id=? AND is_deleted=0""", (project_id,)).fetchone()
    if not project:
        connection.close(); abort(404)
    connection.execute("""UPDATE projects SET is_deleted=1,deleted_at=CURRENT_TIMESTAMP,
        deleted_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""", (current_actor(), project_id))
    recycle_record(connection, "project", project_id, f"{project[0]} · {project[1]}", "Projects")
    audit_change(connection, "project", project_id, "soft_delete",
        before={"project_number": project[0], "status": project[2]}, after={"is_deleted": 1})
    connection.commit(); connection.close()
    flash(f"Project {project[0]} moved to the Recycle Bin. Its Jobs and records were retained.", "success")
    return redirect(url_for("project_jobs"))


@app.get("/recycle-bin")
def recycle_bin():
    rows = query("""SELECT id,entity_type,entity_id,display_name,COALESCE(original_location,''),
        deleted_by,deleted_at FROM recycle_bin ORDER BY deleted_at DESC,id DESC""")
    return render_template("recycle_bin.html", rows=rows,
        can_manage=session.get("role_name") in {"Chief Meteorologist", "Administrator"} or app.config.get("TESTING"),
        page_title="Recycle Bin", page_subtitle="Recoverable deleted Projects, Jobs, files and folders",
        active_nav="recycle")


@app.post("/recycle-bin/<int:item_id>/restore")
def recycle_restore(item_id):
    if not app.config.get("TESTING") and session.get("role_name") not in {"Chief Meteorologist", "Administrator"}:
        abort(403)
    connection = get_connection()
    item = connection.execute("SELECT entity_type,entity_id,display_name FROM recycle_bin WHERE id=?",
        (item_id,)).fetchone()
    if not item:
        connection.close(); abort(404)
    targets = {"project": ("projects",), "calibration_job": ("calibration_jobs",),
        "controlled_document": ("controlled_documents",)}
    target = targets.get(item[0])
    if not target:
        connection.close(); flash("This item type cannot be restored automatically.", "error")
        return redirect(url_for("recycle_bin"))
    connection.execute(f"""UPDATE {target[0]} SET is_deleted=0,deleted_at=NULL,deleted_by=NULL,
        updated_at=CURRENT_TIMESTAMP WHERE id=?""", (item[1],))
    connection.execute("DELETE FROM recycle_bin WHERE id=?", (item_id,))
    audit_change(connection, item[0], item[1], "restore", after={"restored": True})
    connection.commit(); connection.close()
    flash(f"{item[2]} restored.", "success")
    return redirect(url_for("recycle_bin"))


@app.post("/recycle-bin/<int:item_id>/purge")
def recycle_purge(item_id):
    if not app.config.get("TESTING") and session.get("role_name") not in {"Chief Meteorologist", "Administrator"}:
        abort(403)
    connection = get_connection()
    try:
        item = connection.execute("SELECT entity_type,entity_id,display_name FROM recycle_bin WHERE id=?",
            (item_id,)).fetchone()
        if not item:
            abort(404)
        if item[0] == "project":
            linked = connection.execute("SELECT COUNT(*) FROM project_jobs WHERE project_id=?", (item[1],)).fetchone()[0]
            if linked:
                raise ValueError("This Project contains retained Jobs and cannot be permanently deleted.")
            connection.execute("DELETE FROM projects WHERE id=? AND is_deleted=1", (item[1],))
        elif item[0] == "calibration_job":
            linked = connection.execute("SELECT COUNT(*) FROM uncertainty_calculations WHERE job_id=?", (item[1],)).fetchone()[0]
            if linked:
                raise ValueError("This Job contains retained uncertainty records and cannot be permanently deleted.")
            connection.execute("DELETE FROM project_jobs WHERE job_id=?", (item[1],))
            connection.execute("DELETE FROM calibration_jobs WHERE id=? AND is_deleted=1", (item[1],))
        elif item[0] == "controlled_document":
            document = connection.execute("SELECT file_path FROM controlled_documents WHERE id=? AND is_deleted=1",
                (item[1],)).fetchone()
            if document and document[0]:
                file_path = Path(document[0]).resolve()
                document_root = Path("data/documents").resolve()
                if document_root not in file_path.parents:
                    raise ValueError("The stored document path is outside the controlled file repository.")
                if file_path.is_file():
                    file_path.unlink()
            connection.execute("DELETE FROM controlled_documents WHERE id=? AND is_deleted=1", (item[1],))
        else:
            raise ValueError("This item type cannot be permanently deleted automatically.")
        connection.execute("DELETE FROM recycle_bin WHERE id=?", (item_id,))
        audit_change(connection, item[0], item[1], "permanent_delete",
            before={"display_name": item[2], "is_deleted": 1}, after={"purged": True})
        connection.commit(); flash(f"{item[2]} permanently deleted.", "success")
    except ValueError as error:
        connection.rollback(); flash(str(error), "error")
    finally:
        connection.close()
    return redirect(url_for("recycle_bin"))


@app.route("/uncertainty/job/<int:job_id>", methods=["GET", "POST"])
def uncertainty_job(job_id):
    return redirect(url_for("uncertainty", job=job_id))


@app.route("/equipment", methods=["GET", "POST"])
def equipment():
    due_soon_days = setting_integer("equipment_due_soon_days", 30)
    if request.method == "POST":
        require_permission("equipment_add")
        if request.form.get("action") == "upload":
            upload = request.files.get("equipment_list")
            if not upload or not upload.filename:
                flash("Select a CSV or XLSX equipment list.", "error")
                return redirect(url_for("equipment"))
            try:
                imported, updated, skipped = _import_equipment_list(upload)
                flash(f"Equipment list processed: {imported} added, {updated} updated, {skipped} skipped.", "success")
            except Exception as error:
                flash(str(error), "error")
            return redirect(url_for("equipment"))
        asset_number = request.form.get("asset_number", "").strip()
        equipment_name = request.form.get("equipment_name", "").strip()
        validity_value = request.form.get("validity_value", type=int)
        validity_unit = request.form.get("validity_unit", "")
        if not asset_number or not equipment_name or not validity_value or validity_unit not in {"Days", "Months", "Years"}:
            flash("Equipment, Equipment ID, and a positive Validity in Days, Months, or Years are required.", "error")
        else:
            connection = get_connection()
            try:
                serial_number = request.form.get("serial_number", "").strip()
                if serial_number and connection.execute("""SELECT 1 FROM laboratory_equipment
                    WHERE LOWER(serial_number)=LOWER(?) AND is_active=1""", (serial_number,)).fetchone():
                    raise ValueError("Serial number is already assigned to another equipment record.")
                cursor = connection.execute("""INSERT INTO laboratory_equipment
                    (asset_number, equipment_name, equipment_type, manufacturer, model,
                     serial_number,laboratory_location,is_reference_standard,include_in_calibration_programme,
                     validity_value,validity_unit,record_status,notes,created_by,updated_by)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (asset_number, equipment_name, request.form.get("equipment_type", "").strip(),
                     request.form.get("manufacturer", "").strip(), request.form.get("model", "").strip(),
                     serial_number,request.form.get("laboratory_location", "").strip(),
                     int(bool(request.form.get("is_primary"))),
                     int(bool(request.form.get("include_programme"))),
                     validity_value,validity_unit,controlled_status(),
                     request.form.get("notes", "").strip() or None, current_actor(), current_actor()))
                add_note(connection, "equipment", cursor.lastrowid, request.form.get("notes"))
                audit_change(connection, "equipment", cursor.lastrowid, "create", after={
                    "asset_number": asset_number, "equipment_name": equipment_name,
                    "serial_number": serial_number})
                connection.commit(); flash("Equipment added to the register." if controlled_status()=="Approved"
                    else "Equipment saved as Draft. Submit it for review before controlled use.", "success")
            except Exception as error:
                flash(str(error), "error")
            finally:
                connection.close()
            return redirect(url_for("equipment"))
    all_rows = query("""SELECT id, equipment_name, asset_number, COALESCE(serial_number, ''),
        COALESCE(model, ''), last_calibration_date, next_calibration_date,
        CASE WHEN next_calibration_date IS NOT NULL AND next_calibration_date < date('now','localtime') THEN 'Expired'
             WHEN next_calibration_date IS NOT NULL AND next_calibration_date <= date('now','localtime',?) THEN 'Due Soon'
             ELSE 'Active' END AS derived_status,
        COALESCE(certificate_path, ''), COALESCE(equipment_type, ''),
        COALESCE(manufacturer, ''), is_reference_standard, include_in_calibration_programme,
        validity_value,COALESCE(validity_unit,''),record_status
        FROM laboratory_equipment WHERE is_active=1
        AND TRIM(equipment_name) <> '' AND TRIM(asset_number) <> ''
        AND LOWER(TRIM(equipment_name)) NOT IN ('equipment','sample','sample equipment','test','test equipment','placeholder')
        ORDER BY equipment_name""", (f"+{due_soon_days} days",))
    search_term = request.args.get("q", "").strip()
    status_filter = request.args.get("status", "").strip()
    rows = [row for row in all_rows if not search_term or search_term.casefold() in " ".join(
        str(value or "") for value in row[1:5]).casefold()]
    if status_filter:
        rows = [row for row in rows if row[7] == status_filter]
    stats = {"All": len(all_rows), "Active": 0, "Due Soon": 0, "Expired": 0}
    for row in all_rows:
        stats[row[7]] += 1
    return render_template("equipment.html", rows=rows, stats=stats,
        search_term=search_term, status_filter=status_filter, page_title="Equipment",
        page_subtitle="", active_nav="iso17025")


@app.get("/equipment/export")
def equipment_export():
    due_soon_days = setting_integer("equipment_due_soon_days", 30)
    rows = query("""SELECT equipment_name, asset_number, COALESCE(serial_number,''),
        COALESCE(model,''), COALESCE(last_calibration_date,''), COALESCE(next_calibration_date,''),
        CASE WHEN next_calibration_date IS NOT NULL AND next_calibration_date < date('now','localtime') THEN 'Expired'
             WHEN next_calibration_date IS NOT NULL AND next_calibration_date <= date('now','localtime',?) THEN 'Due Soon'
             ELSE 'Active' END
        FROM laboratory_equipment WHERE is_active=1
        AND TRIM(equipment_name) <> '' AND TRIM(asset_number) <> ''
        AND LOWER(TRIM(equipment_name)) NOT IN ('equipment','sample','sample equipment','test','test equipment','placeholder')
        ORDER BY equipment_name""", (f"+{due_soon_days} days",))
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(("Equipment", "Equipment ID", "Serial No.", "Model", "Cal Date", "Due Date", "Status"))
    writer.writerows(rows)
    return app.response_class(output.getvalue(), mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=flowlab-equipment.csv"})


def _equipment_header(value):
    return " ".join(str(value or "").strip().lower().replace("_", " ").replace(".", "").split())


def _equipment_value(value):
    if value is None:
        return ""
    if hasattr(value, "date") and hasattr(value, "isoformat"):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _import_equipment_list(upload):
    extension = Path(upload.filename).suffix.lower()
    content = upload.read()
    if extension == ".csv":
        text_stream = io.StringIO(content.decode("utf-8-sig"))
        raw_rows = list(csv.reader(text_stream))
    elif extension == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise ValueError("XLSX import is unavailable. Use CSV.") from error
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        raw_rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
    else:
        raise ValueError("Only CSV and XLSX equipment lists are accepted.")
    if not raw_rows:
        raise ValueError("The equipment list is empty.")
    headers = [_equipment_header(value) for value in raw_rows[0]]
    aliases = {
        "equipment_name": ("equipment", "equipment name", "name", "description"),
        "asset_number": ("equipment id", "equipment no", "equipment number", "asset", "asset number", "id", "tag"),
        "serial_number": ("serial no", "serial number", "serial"),
        "model": ("model", "model no", "model number"),
        "manufacturer": ("manufacturer", "make"),
        "equipment_type": ("equipment type", "type", "category"),
        "last_calibration_date": ("cal date", "calibration date", "last calibration", "last calibration date"),
        "next_calibration_date": ("due date", "next due", "next calibration", "next calibration date"),
        "laboratory_location": ("location", "laboratory location"),
        "validity": ("validity", "calibration validity", "calibration interval",
                     "calibration interval months", "calibration interval (months)"),
        "validity_value": ("validity value",),
        "validity_unit": ("validity unit",),
        "classification": ("classification", "class"),
        "programme": ("programme", "calibration programme", "include in programme"),
        "notes": ("notes", "note"),
    }
    positions = {}
    for field, names in aliases.items():
        positions[field] = next((headers.index(name) for name in names if name in headers), None)
    if positions["equipment_name"] is None or positions["asset_number"] is None:
        raise ValueError("The list must contain Equipment and Equipment ID columns.")

    def cell(row, field):
        position = positions[field]
        return _equipment_value(row[position]) if position is not None and position < len(row) else ""

    connection = get_connection()
    imported = updated = skipped = 0
    try:
        for row in raw_rows[1:]:
            equipment_name = cell(row, "equipment_name")
            asset_number = cell(row, "asset_number")
            if (not equipment_name or not asset_number or
                    equipment_name.casefold() in {"equipment", "sample", "sample equipment", "test", "test equipment", "placeholder"}):
                skipped += 1
                continue
            validity_text = cell(row, "validity")
            value_text = cell(row, "validity_value")
            unit_text = cell(row, "validity_unit").title()
            if validity_text and not value_text:
                parts = validity_text.split()
                value_text = parts[0] if parts else ""
                unit_text = parts[1].title().rstrip("s") + "s" if len(parts) > 1 else unit_text
                validity_header = headers[positions["validity"]] if positions["validity"] is not None else ""
                if not unit_text and "month" in validity_header:
                    unit_text = "Months"
            if not value_text:
                start_text = cell(row, "last_calibration_date")
                due_text = cell(row, "next_calibration_date")
                try:
                    start_date = date.fromisoformat(start_text)
                    due_date = date.fromisoformat(due_text)
                    if due_date <= start_date:
                        raise ValueError
                    if start_date.day == due_date.day and start_date.month == due_date.month:
                        value_text,unit_text = str(due_date.year-start_date.year),"Years"
                    elif start_date.day == due_date.day:
                        value_text,unit_text = str((due_date.year-start_date.year)*12+due_date.month-start_date.month),"Months"
                    else:
                        value_text,unit_text = str((due_date-start_date).days),"Days"
                except (TypeError, ValueError):
                    pass
            try:
                validity_value = int(float(value_text))
            except (TypeError, ValueError):
                raise ValueError(f"{asset_number}: a positive Validity value is required.")
            unit_text = {"Day":"Days","Month":"Months","Year":"Years"}.get(unit_text,unit_text)
            if validity_value <= 0 or unit_text not in {"Days","Months","Years"}:
                raise ValueError(f"{asset_number}: Validity unit must be Days, Months, or Years.")
            classification = cell(row, "classification").casefold()
            programme = cell(row, "programme").casefold()
            serial_number = cell(row, "serial_number")
            duplicate_serial = (connection.execute("""SELECT asset_number FROM laboratory_equipment
                WHERE LOWER(serial_number)=LOWER(?) AND asset_number<>? AND is_active=1""",
                (serial_number, asset_number)).fetchone() if serial_number else None)
            if duplicate_serial:
                raise ValueError(f"Serial number for {asset_number} is already assigned to {duplicate_serial[0]}.")
            exists = connection.execute("SELECT id FROM laboratory_equipment WHERE asset_number=?", (asset_number,)).fetchone()
            connection.execute("""INSERT INTO laboratory_equipment
                (asset_number,equipment_name,serial_number,model,manufacturer,equipment_type,
                 last_calibration_date,next_calibration_date,laboratory_location,
                 validity_value,validity_unit,is_reference_standard,
                 include_in_calibration_programme,record_status,is_active,updated_at)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,1,CURRENT_TIMESTAMP)
                ON CONFLICT(asset_number) DO UPDATE SET
                 equipment_name=excluded.equipment_name,
                 serial_number=excluded.serial_number,model=excluded.model,
                 manufacturer=excluded.manufacturer,equipment_type=excluded.equipment_type,
                 last_calibration_date=excluded.last_calibration_date,
                 next_calibration_date=excluded.next_calibration_date,
                 laboratory_location=excluded.laboratory_location,
                 validity_value=excluded.validity_value,validity_unit=excluded.validity_unit,
                 is_reference_standard=excluded.is_reference_standard,
                 include_in_calibration_programme=excluded.include_in_calibration_programme,
                 is_active=1,updated_at=CURRENT_TIMESTAMP""", (
                asset_number, equipment_name, serial_number, cell(row, "model"),
                cell(row, "manufacturer"), cell(row, "equipment_type"),
                cell(row, "last_calibration_date") or None,
                cell(row, "next_calibration_date") or None, cell(row, "laboratory_location"),
                validity_value,unit_text,
                int(classification in {"primary", "reference", "reference standard", "yes", "true", "1"}),
                int(programme in {"included", "include", "yes", "true", "1"}),controlled_status()))
            equipment_id = exists[0] if exists else connection.execute("SELECT last_insert_rowid()").fetchone()[0]
            add_note(connection, "equipment", equipment_id, cell(row, "notes"))
            audit_change(connection, "equipment", equipment_id, "import_update" if exists else "import_create",
                after={"asset_number": asset_number, "equipment_name": equipment_name,
                       "serial_number": serial_number, "source_file": Path(upload.filename).name})
            updated += int(exists is not None)
            imported += int(exists is None)
        import_directory = Path("data/documents/equipment/imports")
        import_directory.mkdir(parents=True, exist_ok=True)
        (import_directory / f"{uuid4().hex}_{Path(upload.filename).name}").write_bytes(content)
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    return imported, updated, skipped


def _save_equipment_upload(upload, asset_number, category):
    if not upload or not upload.filename:
        return None
    destination = Path("data/documents/equipment") / asset_number / category
    destination.mkdir(parents=True, exist_ok=True)
    target = destination / f"{uuid4().hex}_{Path(upload.filename).name}"
    upload.save(target)
    return str(target)


def _activate_calibration(connection, equipment_id, calibration_id, actor):
    record = connection.execute("""SELECT calibration_date,next_due_date,certificate_path,
        uncertainty_value,uncertainty_unit,coverage_factor,certificate_number,calibrated_by,
        uncertainty_input_type,classification,source_name,sensitivity_coefficient,degrees_of_freedom,
        calibration_range,range_unit,result FROM calibration_history WHERE id=? AND equipment_id=?""",
        (calibration_id,equipment_id)).fetchone()
    if not record:
        raise ValueError("Calibration record was not found.")
    previous = connection.execute("""SELECT id FROM calibration_history WHERE equipment_id=?
        AND id<>? AND status='Approved' AND is_deleted=0 ORDER BY calibration_date DESC,id DESC""",
        (equipment_id,calibration_id)).fetchall()
    for item in previous:
        connection.execute("UPDATE calibration_history SET status='Superseded',superseded_by=? WHERE id=?",
            (calibration_id,item[0]))
        audit_change(connection,"calibration",item[0],"supersede",
            before={"status":"Approved"},after={"status":"Superseded","superseded_by":calibration_id})
    connection.execute("""UPDATE calibration_history SET status='Approved',approved_by=?,
        approved_at=CURRENT_TIMESTAMP WHERE id=?""", (actor,calibration_id))
    connection.execute("""UPDATE laboratory_equipment SET last_calibration_date=?,next_calibration_date=?,
        certificate_path=COALESCE(?,certificate_path),is_reference_standard=?,status=?,updated_by=?,
        updated_at=CURRENT_TIMESTAMP WHERE id=?""",
        (record[0],record[1],record[2],int(record[9]=="Primary Equipment"),
         "Active" if record[15] in {"Pass","Adjusted"} else "Out of Service",actor,equipment_id))
    if record[15] == "Fail":
        connection.execute("UPDATE uncertainty_profiles SET is_active=0 WHERE equipment_id=?",(equipment_id,))
        return
    if record[3] is not None:
        rule = Budget2Resolver(UncertaintyRepository().approved_budget2_rules()).resolve(
            record[8] or "Expanded / Certificate Uncertainty",record[5])
        standard_uncertainty = record[3] / rule.divisor
        expanded_uncertainty = record[3] if rule.requires_coverage_factor else record[3] * record[5]
        basis = "expanded" if rule.requires_coverage_factor else "standard"
        connection.execute("UPDATE uncertainty_profiles SET is_active=0 WHERE equipment_id=?",(equipment_id,))
        connection.execute("""INSERT INTO uncertainty_profiles
            (equipment_id,version,certificate_number,calibration_laboratory,calibration_date,
             next_due_date,coverage_factor,expanded_uncertainty,standard_uncertainty,
             certificate_path,source_name,uncertainty_unit,evaluation_basis,distribution,divisor,
             sensitivity,degrees_of_freedom,is_active) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,1)""",(
            equipment_id,f"CAL-{record[0]}-{uuid4().hex[:6]}",record[6],record[7],record[0],
            record[1],record[5],expanded_uncertainty,standard_uncertainty,record[2],
            record[10],record[4],basis,rule.distribution,rule.divisor,record[11],record[12]))
        if record[9] == "Primary Equipment":
            range_parts = (record[13] or "").split("–",1)
            range_min = float(range_parts[0]) if len(range_parts)==2 else None
            range_max = float(range_parts[1]) if len(range_parts)==2 else None
            connection.execute("""INSERT INTO equipment_capabilities
                (equipment_id,operating_range_min,operating_range_max,operating_unit,standard_uncertainty)
                VALUES (?,?,?,?,?) ON CONFLICT(equipment_id) DO UPDATE SET
                operating_range_min=excluded.operating_range_min,
                operating_range_max=excluded.operating_range_max,
                operating_unit=excluded.operating_unit,
                standard_uncertainty=excluded.standard_uncertainty,updated_at=CURRENT_TIMESTAMP""",
                (equipment_id,range_min,range_max,record[14],standard_uncertainty))


@app.route("/equipment/<int:equipment_id>/edit", methods=["GET", "POST"])
def equipment_edit(equipment_id):
    require_permission("equipment_edit")
    connection = get_connection()
    row = connection.execute("""SELECT id,equipment_name,asset_number,COALESCE(serial_number,''),
        COALESCE(manufacturer,''),COALESCE(model,''),COALESCE(equipment_type,''),
        validity_value,COALESCE(validity_unit,''),is_reference_standard,
        include_in_calibration_programme,record_status,COALESCE(notes,''),
        COALESCE(laboratory_location,'')
        FROM laboratory_equipment WHERE id=? AND is_active=1""", (equipment_id,)).fetchone()
    if not row:
        connection.close(); abort(404)
    if row[11] == "Approved" and not (app.config.get("TESTING") or
            Permissions.can(session.get("role_name", ""), "edit_approved")):
        connection.close(); abort(403)
    if request.method == "POST":
        try:
            reason = request.form.get("reason", "").strip()
            if row[11] == "Approved" and not reason:
                raise ValueError("A reason is required to edit approved equipment.")
            validity_value = request.form.get("validity_value", type=int)
            validity_unit = request.form.get("validity_unit", "")
            if not request.form.get("equipment_name", "").strip() or not validity_value or validity_unit not in {"Days","Months","Years"}:
                raise ValueError("Equipment and a positive Validity in Days, Months, or Years are required.")
            serial = request.form.get("serial_number", "").strip()
            if serial and connection.execute("""SELECT 1 FROM laboratory_equipment WHERE id<>?
                AND LOWER(serial_number)=LOWER(?) AND is_active=1""", (equipment_id,serial)).fetchone():
                raise ValueError("Serial number is already assigned to another equipment record.")
            before = {"equipment_name":row[1],"serial_number":row[3],"manufacturer":row[4],
                "model":row[5],"equipment_type":row[6],"validity_value":row[7],
                "validity_unit":row[8],"programme":row[10],"status":row[11],
                "laboratory_location":row[13]}
            connection.execute("""UPDATE laboratory_equipment SET equipment_name=?,serial_number=?,
                manufacturer=?,model=?,equipment_type=?,validity_value=?,validity_unit=?,
                is_reference_standard=?,include_in_calibration_programme=?,notes=?,laboratory_location=?,updated_by=?,
                updated_at=CURRENT_TIMESTAMP WHERE id=?""", (
                request.form["equipment_name"].strip(),serial,
                request.form.get("manufacturer", "").strip(),request.form.get("model", "").strip(),
                request.form.get("equipment_type", "").strip(),validity_value,validity_unit,
                int(bool(request.form.get("is_primary"))),int(bool(request.form.get("include_programme"))),
                request.form.get("notes", "").strip() or None,
                request.form.get("laboratory_location", "").strip(),current_actor(),equipment_id))
            after = {"equipment_name":request.form["equipment_name"].strip(),"serial_number":serial,
                "validity_value":validity_value,"validity_unit":validity_unit,
                "laboratory_location":request.form.get("laboratory_location", "").strip(),
                "reason":reason,"status":row[11]}
            audit_change(connection,"equipment",equipment_id,"edit_approved" if row[11]=="Approved" else "edit",
                before=before,after=after)
            connection.commit(); connection.close(); flash("Equipment updated.","success")
            return redirect(url_for("equipment"))
        except Exception as error:
            connection.rollback(); flash(str(error),"error")
    connection.close()
    return render_template("equipment_edit.html",item=row,page_title="Edit Equipment",
        page_subtitle="",active_nav="iso17025")


@app.route("/equipment/<int:equipment_id>", methods=["GET", "POST"])
def equipment_detail(equipment_id):
    equipment_row = query("""SELECT id, equipment_name, asset_number, COALESCE(serial_number,''),
        COALESCE(model,''), COALESCE(manufacturer,''), COALESCE(equipment_type,''),
        COALESCE(laboratory_location,''), validity_value,COALESCE(validity_unit,''),
        last_calibration_date, next_calibration_date, is_reference_standard,
        include_in_calibration_programme, COALESCE(notes,''),record_status
        FROM laboratory_equipment WHERE id=? AND is_active=1""", (equipment_id,))
    if not equipment_row:
        abort(404)
    item = equipment_row[0]
    if request.method == "POST":
        action = request.form.get("action")
        require_permission("calibration" if action == "calibration" else "maintenance")
        connection = get_connection()
        try:
            if action == "calibration":
                calibration_date = request.form.get("calibration_date")
                calibrated_by = request.form.get("calibrated_by", "").strip()
                certificate_number = request.form.get("certificate_number", "").strip()
                validity_value = request.form.get("validity_value", type=int)
                validity_unit = request.form.get("validity_unit", "")
                required = {"Calibration Date":calibration_date,"Calibrated By":calibrated_by,
                    "Certificate Number":certificate_number,"Uncertainty":request.form.get("uncertainty_value", "").strip(),
                    "Uncertainty Unit":request.form.get("uncertainty_unit", "").strip(),
                    "Range Minimum":request.form.get("range_min", "").strip(),
                    "Range Maximum":request.form.get("range_max", "").strip(),
                    "Range Unit":request.form.get("range_unit", "").strip(),"Validity":validity_value,
                    "Validity Unit":validity_unit,"Classification":request.form.get("classification"),
                    "Input Type":request.form.get("uncertainty_input_type"),
                    "Coverage Factor":request.form.get("coverage_factor"),"Result":request.form.get("result")}
                missing = [name for name,value in required.items() if value in (None,"")]
                if missing:
                    raise ValueError("Required calibration fields: " + ", ".join(missing) + ".")
                next_due_date = add_validity(calibration_date,validity_value,validity_unit)
                duplicate_date = connection.execute("""SELECT 1 FROM calibration_history WHERE equipment_id=?
                    AND calibration_date=? AND status IN ('Approved','Superseded') AND is_deleted=0""",
                    (equipment_id,calibration_date)).fetchone()
                if duplicate_date:
                    raise ValueError("A calibration for this equipment already exists with this calibration date.")
                duplicate_certificate = connection.execute("""SELECT 1 FROM calibration_history WHERE equipment_id=?
                    AND LOWER(certificate_number)=LOWER(?) AND status IN ('Approved','Superseded') AND is_deleted=0""",
                    (equipment_id,certificate_number)).fetchone()
                if duplicate_certificate:
                    raise ValueError("This certificate number already exists for this equipment.")
                certificate = request.files.get("certificate")
                validate_report(certificate, required=True)
                certificate_path = _save_equipment_upload(
                    certificate, item[2], "calibration")
                uncertainty_value = float(request.form["uncertainty_value"])
                if uncertainty_value < 0:
                    raise ValueError("Uncertainty cannot be negative.")
                coverage_factor = float(request.form.get("coverage_factor") or 2)
                range_min = float(request.form["range_min"]); range_max = float(request.form["range_max"])
                if range_min > range_max:
                    raise ValueError("Range minimum cannot be greater than range maximum.")
                calibration_range = f"{range_min:g}–{range_max:g}"
                classification = request.form.get("classification")
                classification = {"Primary": "Primary Equipment", "Secondary": "Secondary Equipment"}.get(
                    classification, classification)
                if classification not in {"Primary Equipment", "Secondary Equipment"}:
                    raise ValueError("Select Primary Equipment or Secondary Equipment.")
                source_name = request.form.get("source_name", "").strip()
                sensitivity = request.form.get("sensitivity_coefficient", type=float)
                degrees_of_freedom = request.form.get("degrees_of_freedom", type=float)
                if classification == "Primary Equipment":
                    missing_type_b = []
                    if not source_name:
                        missing_type_b.append("Source Name")
                    if sensitivity is None:
                        missing_type_b.append("Sensitivity Coefficient")
                    if missing_type_b:
                        raise ValueError("Primary Equipment Type B fields required: " + ", ".join(missing_type_b) + ".")
                    if degrees_of_freedom is not None and degrees_of_freedom <= 0:
                        raise ValueError("Degrees of freedom must be positive or left blank for infinite.")
                notes = request.form.get("notes", "").strip()
                status = controlled_status()
                if status != "Approved" and request.form.get("intent") == "submit":
                    status = "Submitted for Review"
                reviewer_id = request.form.get("reviewer_user_id", type=int) if status == "Submitted for Review" else None
                reviewer = connection.execute("""SELECT u.full_name FROM users u JOIN roles r ON r.id=u.role_id
                    WHERE u.id=? AND u.is_active=1 AND r.name IN
                    ('Chief Meteorologist','Supervisor','Administrator','HOD')""",
                    (reviewer_id,)).fetchone() if reviewer_id else None
                if status == "Submitted for Review" and not reviewer:
                    raise ValueError("Select the specific user who will review this calibration.")
                cursor = connection.execute("""INSERT INTO calibration_history
                    (equipment_id,calibration_date,next_due_date,calibrated_by,certificate_path,
                     remarks,uncertainty_value,uncertainty_unit,coverage_factor,
                     calibration_range,range_unit,result,uploaded_by,notes,certificate_number,
                     uncertainty_input_type,status,submitted_by,validity_value,validity_unit,
                     classification,source_name,sensitivity_coefficient,degrees_of_freedom,assigned_reviewer_id)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
                    equipment_id, calibration_date, next_due_date,
                    calibrated_by, certificate_path, notes, uncertainty_value,
                    request.form.get("uncertainty_unit", "").strip(), coverage_factor,
                    calibration_range, request.form["range_unit"].strip(), request.form["result"],
                    current_actor(), notes,certificate_number,request.form["uncertainty_input_type"],
                    status,current_actor(),validity_value,validity_unit,classification,
                    source_name or None,sensitivity,degrees_of_freedom,reviewer_id))
                if status == "Submitted for Review":
                    connection.execute("""INSERT INTO workflow_tasks
                        (entity_type,entity_id,task_type,status,submitted_by,submitted_user_id,
                         assigned_to,assigned_user_id) VALUES (?,?,'Technical Review','Pending',?,?,?,?)""",
                        ("calibration",cursor.lastrowid,current_actor(),session.get("user_id"),reviewer[0],reviewer_id))
                    add_notification(connection, reviewer_id, "Calibration assigned for review",
                        f"{item[2]} · {item[1]} was submitted by {current_actor()}.",
                        url_for("equipment_detail",equipment_id=equipment_id,tab="calibration"),
                        "calibration",cursor.lastrowid)
                add_note(connection, "calibration", cursor.lastrowid, notes)
                audit_change(connection, "calibration", cursor.lastrowid, "create", after={
                    "equipment_id": equipment_id, "calibration_date": calibration_date,
                    "next_due_date": next_due_date, "report": bool(certificate_path)})
                if status == "Approved":
                    _activate_calibration(connection,equipment_id,cursor.lastrowid,current_actor())
                flash("Calibration approved and made current." if status=="Approved" else
                    f"Calibration saved with status {status}.", "success")
            elif action == "maintenance":
                maintenance_date = request.form.get("maintenance_date")
                if not maintenance_date:
                    raise ValueError("Maintenance date is required.")
                next_maintenance_date = request.form.get("next_maintenance_date") or None
                if next_maintenance_date and next_maintenance_date < maintenance_date:
                    raise ValueError("Next maintenance date cannot be earlier than maintenance date.")
                maintenance_type = request.form.get("maintenance_type")
                if maintenance_type not in {"Repair", "Corrective", "Preventive"}:
                    raise ValueError("Select Repair, Corrective, or Preventive maintenance.")
                value = float(request.form.get("cost") or 0)
                if value < 0:
                    raise ValueError("Maintenance value cannot be negative.")
                maintenance_report = request.files.get("maintenance_document")
                validate_report(maintenance_report)
                document_path = _save_equipment_upload(
                    maintenance_report, item[2], "maintenance")
                notes = request.form.get("notes", "").strip()
                status = controlled_status()
                if status != "Approved" and request.form.get("intent") == "submit":
                    status = "Submitted for Review"
                reviewer_id = request.form.get("reviewer_user_id", type=int) if status == "Submitted for Review" else None
                reviewer = connection.execute("""SELECT u.full_name FROM users u JOIN roles r ON r.id=u.role_id
                    WHERE u.id=? AND u.is_active=1 AND r.name IN
                    ('Chief Meteorologist','Supervisor','Administrator','HOD')""",
                    (reviewer_id,)).fetchone() if reviewer_id else None
                if status == "Submitted for Review" and not reviewer:
                    raise ValueError("Select the specific user who will review this maintenance record.")
                cursor = connection.execute("""INSERT INTO maintenance_history
                    (equipment_id,maintenance_date,maintenance_type,performed_by,cost,
                     document_path,remarks,next_maintenance_date,currency,supplier,
                     work_performed,findings,parts_replaced,metrological_impact,recalibration_required,
                     uploaded_by,notes,status,submitted_by,assigned_reviewer_id)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
                    equipment_id, maintenance_date, maintenance_type,
                    request.form.get("performed_by", ""), value,
                    document_path, notes, next_maintenance_date,
                    request.form.get("currency", "NGN"), request.form.get("supplier", ""),
                    request.form.get("work_performed", ""), request.form.get("findings", ""),
                    request.form.get("parts_replaced", ""), request.form.get("metrological_impact", ""),
                    int(bool(request.form.get("recalibration_required"))),
                    current_actor() if document_path else None, notes,status,current_actor(),reviewer_id))
                if status == "Submitted for Review":
                    connection.execute("""INSERT INTO workflow_tasks
                        (entity_type,entity_id,task_type,status,submitted_by,submitted_user_id,
                         assigned_to,assigned_user_id) VALUES (?,?,'Technical Review','Pending',?,?,?,?)""",
                        ("maintenance",cursor.lastrowid,current_actor(),session.get("user_id"),reviewer[0],reviewer_id))
                    add_notification(connection, reviewer_id, "Maintenance record assigned for review",
                        f"{item[2]} · {item[1]} was submitted by {current_actor()}.",
                        url_for("equipment_detail",equipment_id=equipment_id,tab="maintenance"),
                        "maintenance",cursor.lastrowid)
                add_note(connection, "maintenance", cursor.lastrowid, notes)
                audit_change(connection, "maintenance", cursor.lastrowid, "create", after={
                    "equipment_id": equipment_id, "maintenance_date": maintenance_date,
                    "maintenance_type": maintenance_type, "value": value,
                    "report": bool(document_path)})
                flash(f"Maintenance saved with status {status}.", "success")
            else:
                raise ValueError("Unknown equipment action.")
            connection.commit()
        except Exception as error:
            connection.rollback()
            flash(str(error), "error")
        finally:
            connection.close()
        return redirect(url_for("equipment_detail", equipment_id=equipment_id,
            tab="maintenance" if action == "maintenance" else "calibration"))
    calibrations = query("""SELECT id,calibration_date,next_due_date,COALESCE(calibrated_by,''),
        COALESCE(certificate_path,''),COALESCE(remarks,''),uncertainty_value,
        COALESCE(uncertainty_unit,''),coverage_factor,COALESCE(calibration_range,''),
        COALESCE(range_unit,''),COALESCE(result,''),COALESCE(certificate_number,''),status,
        validity_value,COALESCE(validity_unit,''),COALESCE(submitted_by,''),COALESCE(approved_by,''),
        assigned_reviewer_id,COALESCE(review_comment,'')
        FROM calibration_history
        WHERE equipment_id=? AND is_deleted=0 ORDER BY calibration_date DESC,id DESC""", (equipment_id,))
    maintenance = query("""SELECT id,maintenance_date,COALESCE(maintenance_type,''),
        COALESCE(performed_by,''),cost,COALESCE(document_path,''),COALESCE(remarks,''),
        next_maintenance_date,COALESCE(currency,''),COALESCE(supplier,''),
        COALESCE(work_performed,''),COALESCE(findings,''),COALESCE(parts_replaced,''),
        COALESCE(metrological_impact,''),recalibration_required,status,
        COALESCE(submitted_by,''),COALESCE(approved_by,''),assigned_reviewer_id,
        COALESCE(review_comment,'') FROM maintenance_history
        WHERE equipment_id=? AND is_deleted=0 ORDER BY maintenance_date DESC,id DESC""", (equipment_id,))
    profiles = query("""SELECT version,COALESCE(certificate_number,''),calibration_date,
        next_due_date,coverage_factor,expanded_uncertainty,standard_uncertainty,
        COALESCE(certificate_path,''),is_active FROM uncertainty_profiles
        WHERE equipment_id=? ORDER BY id DESC""", (equipment_id,))
    notes = query("""SELECT note,author,created_at,edited_at,visibility FROM record_notes
        WHERE entity_type='equipment' AND entity_id=? ORDER BY created_at DESC,id DESC""", (equipment_id,))
    return render_template("equipment_detail.html", item=item, calibrations=calibrations,
        maintenance=maintenance, profiles=profiles, notes=notes,
        reviewers=query("""SELECT u.id,u.full_name,r.name FROM users u JOIN roles r ON r.id=u.role_id
            WHERE u.is_active=1 AND u.id<>? AND r.name IN
            ('Chief Meteorologist','Supervisor','Administrator','HOD') ORDER BY u.full_name""",
            (session.get("user_id") or -1,)),
        tab=request.args.get("tab", "overview"), show_form=request.args.get("add") == "1",
        page_title="Equipment", page_subtitle="", active_nav="iso17025")


@app.post("/equipment/<int:equipment_id>/<kind>/<int:record_id>/workflow")
def equipment_record_workflow(equipment_id,kind,record_id):
    action = request.form.get("action", "")
    if action == "approve":
        require_permission("record_approve")
    elif action == "revert":
        require_permission("record_review")
    else:
        require_permission("calibration" if kind=="calibration" else "maintenance")
    table = {"calibration":"calibration_history","maintenance":"maintenance_history"}.get(kind)
    if not table:
        abort(404)
    comment = request.form.get("comment", "").strip()
    connection = get_connection()
    row = connection.execute(f"""SELECT status,submitted_by,assigned_reviewer_id
        FROM {table} WHERE id=? AND equipment_id=? AND is_deleted=0""",
        (record_id,equipment_id)).fetchone()
    if not row:
        connection.close(); abort(404)
    try:
        if action == "submit":
            if row[0] not in {"Draft","Reverted"}:
                raise ValueError("Only a Draft or Reverted record can be submitted.")
            reviewer_id = request.form.get("reviewer_user_id", type=int)
            reviewer = connection.execute("""SELECT u.full_name FROM users u JOIN roles r ON r.id=u.role_id
                WHERE u.id=? AND u.is_active=1 AND r.name IN
                ('Chief Meteorologist','Supervisor','Administrator','HOD')""",
                (reviewer_id,)).fetchone() if reviewer_id else None
            if not reviewer or reviewer_id == session.get("user_id"):
                raise ValueError("Select a different active user to review this record.")
            status = "Submitted for Review"
            connection.execute(f"""UPDATE {table} SET status=?,submitted_by=?,assigned_reviewer_id=?,
                review_comment=NULL WHERE id=?""", (status,current_actor(),reviewer_id,record_id))
            connection.execute("""INSERT INTO workflow_tasks
                (entity_type,entity_id,task_type,status,submitted_by,submitted_user_id,assigned_to,assigned_user_id)
                VALUES (?,?,'Technical Review','Pending',?,?,?,?)""",
                (kind,record_id,current_actor(),session.get("user_id"),reviewer[0],reviewer_id))
            add_notification(connection,reviewer_id,f"{kind.title()} assigned for review",
                f"Equipment record {record_id} was submitted by {current_actor()}.",
                url_for("equipment_detail",equipment_id=equipment_id,tab=kind),kind,record_id)
        elif action == "revert":
            if row[0] != "Submitted for Review" or not comment:
                raise ValueError("A submitted record and a review comment are required to revert.")
            if row[2] and row[2] != session.get("user_id") and not workflow_authority():
                raise ValueError("This review is assigned to another user.")
            status = "Reverted"
            connection.execute(f"""UPDATE {table} SET status=?,reviewed_by=?,reviewed_at=CURRENT_TIMESTAMP,
                review_comment=? WHERE id=?""",(status,current_actor(),comment,record_id))
            sender_id = user_id_for_actor(connection,row[1])
            add_notification(connection,sender_id,f"{kind.title()} returned with comment",
                f"{current_actor()}: {comment}",url_for("equipment_detail",equipment_id=equipment_id,tab=kind),kind,record_id)
        elif action == "approve":
            if row[0] != "Submitted for Review":
                raise ValueError("Only a submitted record can be approved.")
            if row[1] == current_actor() and not workflow_authority():
                raise ValueError("The submitter cannot approve their own controlled record.")
            if row[2] and row[2] != session.get("user_id") and not workflow_authority():
                raise ValueError("This review is assigned to another user.")
            if kind == "calibration":
                candidate = connection.execute("SELECT calibration_date,certificate_number FROM calibration_history WHERE id=?",
                    (record_id,)).fetchone()
                if connection.execute("""SELECT 1 FROM calibration_history WHERE equipment_id=? AND id<>?
                    AND calibration_date=? AND status IN ('Approved','Superseded') AND is_deleted=0""",
                    (equipment_id,record_id,candidate[0])).fetchone():
                    raise ValueError("A calibration for this equipment already exists with this calibration date.")
                if connection.execute("""SELECT 1 FROM calibration_history WHERE equipment_id=? AND id<>?
                    AND LOWER(certificate_number)=LOWER(?) AND status IN ('Approved','Superseded') AND is_deleted=0""",
                    (equipment_id,record_id,candidate[1])).fetchone():
                    raise ValueError("This certificate number already exists for this equipment.")
                _activate_calibration(connection,equipment_id,record_id,current_actor())
            else:
                connection.execute("""UPDATE maintenance_history SET status='Approved',approved_by=?,
                    approved_at=CURRENT_TIMESTAMP WHERE id=?""",(current_actor(),record_id))
            status = "Approved"
            add_notification(connection,user_id_for_actor(connection,row[1]),f"{kind.title()} approved",
                f"Record {record_id} was approved by {current_actor()}.",
                url_for("equipment_detail",equipment_id=equipment_id,tab=kind),kind,record_id)
        else:
            raise ValueError("Unknown workflow action.")
        connection.execute("INSERT INTO approval_history(entity_type,entity_id,action,actor,notes) VALUES (?,?,?,?,?)",
            (kind,record_id,action,current_actor(),comment or None))
        if action in {"approve","revert"}:
            connection.execute("""UPDATE workflow_tasks SET status='Completed',reviewed_at=CURRENT_TIMESTAMP,
                decision=?,comment=? WHERE entity_type=? AND entity_id=? AND status='Pending'""",
                (status,comment or None,kind,record_id))
        audit_change(connection,kind,record_id,action,before={"status":row[0]},
            after={"status":status,"comment":comment or None})
        connection.commit(); flash(f"{kind.title()} status changed to {status}.","success")
    except ValueError as error:
        connection.rollback(); flash(str(error),"error")
    finally:
        connection.close()
    return redirect(url_for("equipment_detail",equipment_id=equipment_id,tab=kind))


@app.post("/equipment/<int:equipment_id>/workflow")
def equipment_workflow(equipment_id):
    action = request.form.get("action", "")
    require_permission("record_approve" if action == "approve" else
        "record_review" if action == "revert" else "equipment_edit")
    comment = request.form.get("comment", "").strip()
    connection = get_connection()
    row = connection.execute("""SELECT record_status,created_by,submitted_by
        FROM laboratory_equipment WHERE id=? AND is_active=1""", (equipment_id,)).fetchone()
    if not row:
        connection.close(); abort(404)
    try:
        if action == "submit" and row[0] in {"Draft", "Reverted"}:
            status = "Submitted for Review"
            connection.execute("""UPDATE laboratory_equipment SET record_status=?,submitted_by=?,
                submitted_at=CURRENT_TIMESTAMP,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                (status,current_actor(),current_actor(),equipment_id))
        elif action == "revert" and row[0] == "Submitted for Review" and comment:
            status = "Reverted"
            connection.execute("""UPDATE laboratory_equipment SET record_status=?,reviewed_by=?,
                reviewed_at=CURRENT_TIMESTAMP,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                (status,current_actor(),current_actor(),equipment_id))
        elif action == "approve" and row[0] == "Submitted for Review":
            if row[2] == current_actor():
                raise ValueError("The submitter cannot approve their own equipment record.")
            status = "Approved"
            connection.execute("""UPDATE laboratory_equipment SET record_status=?,approved_by=?,
                approved_at=CURRENT_TIMESTAMP,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                (status,current_actor(),current_actor(),equipment_id))
        else:
            raise ValueError("This equipment workflow action is not available for its current status.")
        connection.execute("""INSERT INTO approval_history(entity_type,entity_id,action,actor,notes)
            VALUES ('equipment',?,?,?,?)""", (equipment_id,action,current_actor(),comment or None))
        audit_change(connection,"equipment",equipment_id,action,before={"status":row[0]},
            after={"status":status,"comment":comment or None})
        connection.commit(); flash(f"Equipment status changed to {status}.", "success")
    except ValueError as error:
        connection.rollback(); flash(str(error), "error")
    finally:
        connection.close()
    return redirect(url_for("equipment"))


@app.post("/equipment/<int:equipment_id>/delete")
def equipment_delete(equipment_id):
    require_permission("equipment_delete")
    connection = get_connection()
    item = connection.execute("""SELECT asset_number,equipment_name,is_active,record_status
        FROM laboratory_equipment WHERE id=?""", (equipment_id,)).fetchone()
    if not item:
        connection.close(); abort(404)
    reason = request.form.get("reason", "").strip()
    if item[3] == "Approved" and not reason:
        connection.close(); flash("A reason is required to remove approved equipment.", "error")
        return redirect(url_for("equipment"))
    connection.execute("""UPDATE laboratory_equipment SET is_active=0,
        deactivated_at=CURRENT_TIMESTAMP,deactivated_by=?,updated_by=?,updated_at=CURRENT_TIMESTAMP
        WHERE id=?""", (current_actor(), current_actor(), equipment_id))
    audit_change(connection, "equipment", equipment_id, "deactivate",
        before={"asset_number": item[0], "equipment_name": item[1], "is_active": item[2]},
        after={"is_active": 0, "reason": reason or None})
    connection.commit(); connection.close()
    flash("Equipment removed from the active register. Its history and files were preserved.", "success")
    return redirect(url_for("equipment"))


@app.post("/equipment/<int:equipment_id>/<kind>/<int:record_id>/delete")
def equipment_history_delete(equipment_id, kind, record_id):
    require_permission("equipment_delete")
    table = {"calibration": "calibration_history", "maintenance": "maintenance_history"}.get(kind)
    if not table:
        abort(404)
    connection = get_connection()
    record = connection.execute(f"SELECT id,status FROM {table} WHERE id=? AND equipment_id=? AND is_deleted=0",
        (record_id, equipment_id)).fetchone()
    if not record:
        connection.close(); abort(404)
    reason = request.form.get("reason", "").strip()
    if record[1] in {"Approved", "Superseded"} and not reason:
        connection.close(); flash("A reason is required to remove an approved controlled record.", "error")
        return redirect(url_for("equipment_detail", equipment_id=equipment_id, tab=kind))
    connection.execute(f"""UPDATE {table} SET is_deleted=1,deleted_at=CURRENT_TIMESTAMP,
        deleted_by=? WHERE id=?""", (current_actor(), record_id))
    audit_change(connection, kind, record_id, "soft_delete",
        before={"equipment_id": equipment_id, "is_deleted": 0, "status": record[1]},
        after={"is_deleted": 1, "reason": reason or None})
    if kind == "calibration":
        latest = connection.execute("""SELECT calibration_date,next_due_date,certificate_path
            FROM calibration_history WHERE equipment_id=? AND is_deleted=0
            ORDER BY calibration_date DESC,id DESC LIMIT 1""", (equipment_id,)).fetchone()
        connection.execute("""UPDATE laboratory_equipment SET last_calibration_date=?,
            next_calibration_date=?,certificate_path=?,updated_by=?,updated_at=CURRENT_TIMESTAMP
            WHERE id=?""", ((latest[0] if latest else None), (latest[1] if latest else None),
            (latest[2] if latest else None), current_actor(), equipment_id))
    connection.commit(); connection.close()
    flash(f"{kind.title()} record removed. The audit record was preserved.", "success")
    return redirect(url_for("equipment_detail", equipment_id=equipment_id, tab=kind))


@app.get("/equipment/<int:equipment_id>/file/<kind>/<int:record_id>")
def equipment_file(equipment_id, kind, record_id):
    if kind == "calibration":
        result = query("SELECT certificate_path FROM calibration_history WHERE id=? AND equipment_id=? AND is_deleted=0",
            (record_id, equipment_id))
    elif kind == "maintenance":
        result = query("SELECT document_path FROM maintenance_history WHERE id=? AND equipment_id=? AND is_deleted=0",
            (record_id, equipment_id))
    else:
        abort(404)
    if not result or not result[0][0] or not Path(result[0][0]).is_file():
        abort(404)
    return send_file(Path(result[0][0]).resolve(), as_attachment=False)


@app.get("/programme")
def programme():
    due_soon_days = setting_integer("equipment_due_soon_days", 30)
    rows = query("""SELECT e.asset_number,e.equipment_name,COALESCE(e.serial_number,''),
        COALESCE(e.equipment_type,''),e.validity_value,COALESCE(e.validity_unit,''),
        e.last_calibration_date,e.next_calibration_date,
        CASE WHEN next_calibration_date IS NOT NULL AND next_calibration_date < date('now','localtime') THEN 'Expired'
             WHEN next_calibration_date IS NOT NULL AND next_calibration_date <= date('now','localtime',?) THEN 'Due Soon'
             ELSE 'Active' END,
        COALESCE((SELECT c.certificate_number FROM calibration_history c WHERE c.equipment_id=e.id
            AND c.status='Approved' AND c.is_deleted=0 ORDER BY c.calibration_date DESC,c.id DESC LIMIT 1),''),
        COALESCE((SELECT c.status FROM calibration_history c WHERE c.equipment_id=e.id
            AND c.status='Approved' AND c.is_deleted=0 ORDER BY c.calibration_date DESC,c.id DESC LIMIT 1),'Not Calibrated')
        FROM laboratory_equipment e WHERE e.include_in_calibration_programme=1 AND e.is_active=1
        ORDER BY e.next_calibration_date,e.equipment_name""", (f"+{due_soon_days} days",))
    return render_template("programme.html", rows=rows,
        page_title="Calibration Programme", page_subtitle="",
        active_nav="iso17025")


@app.route("/documents", methods=["GET", "POST"])
def documents():
    if request.method == "POST":
        require_permission("reports")
        document = request.files.get("document")
        title = request.form.get("title", "").strip()
        if not document or not document.filename or not title:
            flash("Title and document file are required.", "error")
        else:
            try:
                validate_report(document, required=True)
                connection = get_connection()
                try: ensure_unique_document_number(connection,request.form.get("number"))
                finally: connection.close()
            except ValueError as error:
                flash(str(error), "error")
                return redirect(url_for("documents"))
            destination = Path("data/documents/controlled")
            destination.mkdir(parents=True, exist_ok=True)
            target = destination / f"{uuid4().hex}_{Path(document.filename).name}"
            document.save(target)
            connection = get_connection()
            cursor = connection.execute("""INSERT INTO controlled_documents
                (document_type, document_number, title, revision, original_filename, file_path,
                 status,notes,uploaded_by)
                VALUES (?, ?, ?, ?, ?, ?, 'Draft',?,?)""",
                (request.form.get("document_type", "Procedure"), request.form.get("number") or None,
                 title, request.form.get("revision") or None, document.filename, str(target),
                 request.form.get("notes", "").strip() or None, current_actor()))
            add_note(connection, "controlled_document", cursor.lastrowid, request.form.get("notes"))
            audit_change(connection, "controlled_document", cursor.lastrowid, "create", after={
                "title": title, "filename": document.filename, "status": "Draft"})
            connection.commit(); connection.close(); flash("Controlled document uploaded.", "success")
            return redirect(url_for("documents"))
    rows = query("""SELECT id,document_type,COALESCE(document_number,''),title,
        COALESCE(revision,''),original_filename,status FROM controlled_documents
        WHERE is_deleted=0 ORDER BY created_at DESC""")
    return render_template("documents.html", rows=rows,
        page_title="Standards & Procedures",
        page_subtitle="Controlled laboratory documents and reference materials",
        active_nav="standards")


@app.get("/documents/<int:document_id>/file")
def document_file(document_id):
    row = query("SELECT file_path FROM controlled_documents WHERE id=? AND is_deleted=0", (document_id,))
    if not row or not row[0][0] or not Path(row[0][0]).is_file():
        abort(404)
    return send_file(Path(row[0][0]).resolve(), as_attachment=False)


@app.post("/documents/<int:document_id>/workflow")
def document_workflow(document_id):
    action = request.form.get("action","")
    require_permission("record_approve" if action=="approve" else "record_review" if action=="revert" else "reports")
    comment = request.form.get("comment","").strip()
    connection = get_connection()
    row = connection.execute("SELECT status,uploaded_by FROM controlled_documents WHERE id=? AND is_deleted=0",
        (document_id,)).fetchone()
    if not row:
        connection.close(); abort(404)
    try:
        if action=="submit" and row[0] in {"Draft","Reverted"}:
            status="Approved" if chief_auto_approval() else "Submitted for Review"
            if status == "Approved":
                connection.execute("""UPDATE controlled_documents SET status=?,submitted_by=?,approved_by=?,
                    approved_at=CURRENT_TIMESTAMP,effective_date=date('now'),updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                    (status,current_actor(),current_actor(),document_id))
            else:
                connection.execute("UPDATE controlled_documents SET status=?,submitted_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",
                    (status,current_actor(),document_id))
        elif action=="revert" and row[0]=="Submitted for Review" and comment:
            status="Reverted"
            connection.execute("""UPDATE controlled_documents SET status=?,reviewed_by=?,reviewed_at=CURRENT_TIMESTAMP,
                updated_at=CURRENT_TIMESTAMP WHERE id=?""",(status,current_actor(),document_id))
        elif action=="approve" and row[0]=="Submitted for Review":
            if row[1]==current_actor() and not workflow_authority():
                raise ValueError("The uploader cannot approve their own controlled document.")
            status="Approved"
            connection.execute("""UPDATE controlled_documents SET status=?,approved_by=?,approved_at=CURRENT_TIMESTAMP,
                effective_date=date('now'),updated_at=CURRENT_TIMESTAMP WHERE id=?""",(status,current_actor(),document_id))
        else:
            raise ValueError("This document workflow action is not available for its current status.")
        connection.execute("INSERT INTO approval_history(entity_type,entity_id,action,actor,notes) VALUES ('controlled_document',?,?,?,?)",
            (document_id,action,current_actor(),comment or None))
        if action == "submit" and status == "Approved":
            connection.execute("""INSERT INTO approval_history(entity_type,entity_id,action,actor,notes)
                VALUES ('controlled_document',?,'auto_approve',?,?)""", (document_id,current_actor(),
                "Automatically approved under Chief Meteorologist authority."))
        audit_change(connection,"controlled_document",document_id,action,before={"status":row[0]},after={"status":status,"comment":comment or None})
        connection.commit(); flash(f"Document status changed to {status}.","success")
    except ValueError as error:
        connection.rollback(); flash(str(error),"error")
    finally:
        connection.close()
    return redirect(url_for("documents"))


@app.post("/documents/<int:document_id>/delete")
def document_delete(document_id):
    require_permission("equipment_delete")
    connection = get_connection()
    row = connection.execute("SELECT title,status,file_path FROM controlled_documents WHERE id=? AND is_deleted=0",
        (document_id,)).fetchone()
    if not row:
        connection.close(); abort(404)
    if row[1] == "Approved":
        connection.close(); flash("Approved documents must be superseded through review and approval.", "error")
        return redirect(url_for("documents"))
    connection.execute("""UPDATE controlled_documents SET is_deleted=1,deleted_at=CURRENT_TIMESTAMP,
        deleted_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""", (current_actor(), document_id))
    recycle_record(connection, "controlled_document", document_id, row[0], row[2])
    audit_change(connection, "controlled_document", document_id, "soft_delete",
        before={"title": row[0], "status": row[1]}, after={"is_deleted": 1})
    connection.commit(); connection.close()
    flash("Document removed. The stored file and audit history were preserved.", "success")
    return redirect(url_for("documents"))


@app.post("/uncertainty/api/calculate")
def uncertainty_api_calculate():
    require_permission("uncertainty_calculate")
    connection = get_connection()
    try:
        result = calculate_payload(connection, request.get_json(force=True) or {})
        return jsonify(ok=True, result=result)
    except ValueError as error:
        return jsonify(ok=False, error=str(error)), 422
    finally:
        connection.close()


@app.post("/uncertainty/api/save")
def uncertainty_api_save():
    require_permission("uncertainty_save")
    connection = get_connection()
    try:
        payload = request.get_json(force=True) or {}
        if payload.get("standalone"):
            raise ValueError("Standalone Tool calculations are not retained. Start from a Project Job to create a controlled record.")
        result = calculate_payload(connection, payload)
        saved = save_cmc(connection, payload, result, current_actor()) if result["isCmc"] else \
            save_calculation(connection, payload, result, current_actor())
        entity = "cmc_revision" if result["isCmc"] else "uncertainty_calculation"
        audit_change(connection, entity, saved["id"], "save_draft",
            after={"revision": saved["revision"], "status": saved["status"]})
        connection.commit()
        return jsonify(ok=True, record=saved, result=result)
    except ValueError as error:
        connection.rollback()
        return jsonify(ok=False, error=str(error)), 422
    finally:
        connection.close()


@app.post("/uncertainty/api/autosave")
def uncertainty_api_autosave():
    require_permission("uncertainty_save")
    connection = get_connection()
    try:
        payload=request.get_json(silent=True) or {}
        if payload.get("standalone"):
            raise ValueError("Standalone Tool calculations are not retained.")
        saved = save_partial_draft(connection,payload,current_actor())
        connection.commit()
        return jsonify(ok=True, record=saved)
    except ValueError as error:
        connection.rollback(); return jsonify(ok=False, error=str(error)), 422
    finally:
        connection.close()


@app.post("/uncertainty/api/records/<record_type>/<int:record_id>/discard")
def uncertainty_api_discard(record_type, record_id):
    require_permission("uncertainty_save")
    connection = get_connection()
    try:
        is_cmc = record_type == "cmc"
        if record_type not in {"calculation", "cmc"}:
            abort(404)
        table = "cmc_revisions" if is_cmc else "uncertainty_calculations"
        creator = "created_by" if is_cmc else "calculated_by"
        reference = "method_name" if is_cmc else "calculation_uid"
        row = connection.execute(f"SELECT status,{creator},{reference} FROM {table} WHERE id=?", (record_id,)).fetchone()
        if not row or row[0] not in {"Draft", "Reverted"}:
            raise ValueError("Only a Draft or Reverted uncertainty record can be discarded.")
        if row[1] != current_actor() and not Permissions.can(session.get("role_name", ""), "uncertainty_approve"):
            raise ValueError("Only the draft creator or an approving officer can discard this draft.")
        entity = "cmc_revision" if is_cmc else "uncertainty_calculation"
        connection.execute("DELETE FROM workflow_tasks WHERE entity_type=? AND entity_id=?", (entity, record_id))
        connection.execute("DELETE FROM approval_history WHERE entity_type=? AND entity_id=?", (entity, record_id))
        if not is_cmc:
            connection.execute("DELETE FROM uncertainty_flow_points WHERE calculation_id=?", (record_id,))
            connection.execute("DELETE FROM uncertainty_sources WHERE calculation_id=?", (record_id,))
        connection.execute(f"DELETE FROM {table} WHERE id=?", (record_id,))
        audit_change(connection, entity, record_id, "discard_draft",
            before={"status": row[0], "calculation_id": row[2]}, after={"discarded": True})
        connection.commit(); return jsonify(ok=True)
    except ValueError as error:
        connection.rollback(); return jsonify(ok=False, error=str(error)), 422
    finally:
        connection.close()


@app.get("/uncertainty/api/records/<record_type>/<int:record_id>")
def uncertainty_api_record(record_type, record_id):
    connection = get_connection()
    try:
        table = "cmc_revisions" if record_type == "cmc" else "uncertainty_calculations"
        creator = "created_by" if record_type == "cmc" else "calculated_by"
        access = connection.execute(f"SELECT status,{creator} FROM {table} WHERE id=?", (record_id,)).fetchone()
        if not access:
            raise ValueError("Uncertainty record was not found.")
        if access[0] in {"Draft","Reverted"} and access[1] != current_actor():
            return jsonify(ok=False,error="This unfinished draft is private to its creator."), 403
        return jsonify(ok=True, record=record_payload(connection, record_id,
            "cmc" if record_type == "cmc" else "calculation"))
    except ValueError as error:
        return jsonify(ok=False, error=str(error)), 404
    finally:
        connection.close()


@app.get("/uncertainty/api/records/calculation/<int:record_id>/details/<kind>")
def uncertainty_api_record_details(record_id, kind):
    connection = get_connection()
    try:
        payload = record_payload(connection, record_id)
        backend = payload.get("backendResult", {})
        if kind == "sources":
            data = [{"flowPoint": point.get("nominalFlow"), "sources": [component.get("input") |
                {"standardUncertainty":component.get("standard_uncertainty"),
                 "contribution":component.get("contribution"),"contributionPercent":component.get("contribution_percent")}
                for component in point.get("result", {}).get("components", [])]}
                for point in backend.get("points", [])]
        elif kind == "observations":
            data = [{"flowPoint":point.get("nominalFlow"),"observations":point.get("observations",[]),
                     "statistics":point.get("statistics",{})} for point in backend.get("points", [])]
        elif kind == "review":
            data = [dict(zip(("action","actor","comment","occurredAt"), row)) for row in connection.execute("""SELECT action,actor,notes,occurred_at
                FROM approval_history WHERE entity_type='uncertainty_calculation' AND entity_id=?
                ORDER BY id""", (record_id,)).fetchall()]
        elif kind == "audit":
            data = [dict(zip(("action","actor","before","after","occurredAt"), row)) for row in connection.execute("""SELECT action,actor,before_state,after_state,occurred_at
                FROM audit_trail WHERE entity_type='uncertainty_calculation' AND entity_id=?
                ORDER BY id""", (record_id,)).fetchall()]
        elif kind == "calculation":
            data = backend
        else:
            return jsonify(ok=False,error="Unknown uncertainty record detail."), 404
        return jsonify(ok=True,kind=kind,data=data)
    finally:
        connection.close()


@app.post("/uncertainty/api/records/<record_type>/<int:record_id>/workflow")
def uncertainty_api_workflow(record_type, record_id):
    body = request.get_json(force=True) or {}
    action = body.get("action", "")
    permission = ({"submit": "uncertainty_save", "review": "record_review",
        "revert": "record_review", "approve": "uncertainty_approve"}).get(action)
    if not permission:
        return jsonify(ok=False, error="Unknown uncertainty workflow action."), 422
    require_permission(permission)
    connection = get_connection()
    try:
        is_cmc = record_type == "cmc"
        table = "cmc_revisions" if is_cmc else "uncertainty_calculations"
        creator_column = "created_by" if is_cmc else "calculated_by"
        entity = "cmc_revision" if is_cmc else "uncertainty_calculation"
        record_row = connection.execute(f"""SELECT status,{creator_column},assigned_reviewer,
            assigned_approver FROM {table} WHERE id=?""", (record_id,)).fetchone()
        if not record_row:
            raise ValueError("Calculation record was not found.")
        auto_approve = action == "submit" and chief_auto_approval()
        reviewer = approver = None
        reviewer_id = approver_id = None
        if action == "submit" and not auto_approve:
            reviewer_id = body.get("reviewerUserId"); approver_id = body.get("approverUserId")
            reviewer = connection.execute("""SELECT u.full_name FROM users u JOIN roles r ON r.id=u.role_id
                WHERE u.id=? AND u.is_active=1 AND r.name IN
                ('Engineer','Supervisor','Chief Meteorologist','Administrator','HOD')""",
                (reviewer_id,)).fetchone()
            approver = connection.execute("""SELECT u.full_name FROM users u JOIN roles r ON r.id=u.role_id
                WHERE u.id=? AND u.is_active=1 AND r.name IN
                ('Chief Meteorologist','Supervisor','Administrator','HOD')""",
                (approver_id,)).fetchone()
            if not reviewer or not approver:
                raise ValueError("Select an active Technical Reviewer and Chief/Supervisor Approving Officer.")
            reviewer, approver = reviewer[0], approver[0]
            if current_actor() in {reviewer, approver}:
                raise ValueError("The calculation creator cannot be assigned to review or approve their own record.")
            if reviewer == approver:
                raise ValueError("Technical Review and final approval must be assigned to different officers.")
        if action in {"submit", "review", "approve"}:
            saved_payload = record_payload(connection, record_id,
                "cmc" if record_type == "cmc" else "calculation")
            saved_payload.pop("backendResult", None)
            calculate_payload(connection, saved_payload)
        status = workflow(connection, "cmc" if record_type == "cmc" else "calculation",
            record_id, action, current_actor(), body.get("comment", ""), reviewer, approver,
            auto_approve=auto_approve, authority_override=workflow_authority())
        creator_id = user_id_for_actor(connection,record_row[1])
        link = url_for("uncertainty",record=record_id)
        if action == "submit" and not auto_approve:
            connection.execute("""UPDATE workflow_tasks SET submitted_user_id=?,assigned_user_id=?
                WHERE entity_type=? AND entity_id=? AND status='Pending'""",
                (session.get("user_id"),reviewer_id,entity,record_id))
            add_notification(connection,reviewer_id,"Uncertainty assigned for Technical Review",
                f"Record {record_id} was submitted by {current_actor()}.",link,entity,record_id)
        elif action == "review":
            approver_id = user_id_for_actor(connection,record_row[3])
            connection.execute("""UPDATE workflow_tasks SET assigned_user_id=?
                WHERE entity_type=? AND entity_id=? AND status='Pending'""",
                (approver_id,entity,record_id))
            add_notification(connection,approver_id,"Uncertainty assigned for final approval",
                f"Technical Review was completed by {current_actor()}.",link,entity,record_id)
        elif action == "revert":
            add_notification(connection,creator_id,"Uncertainty returned with comment",
                f"{current_actor()}: {body.get('comment','')}",link,entity,record_id)
        elif action == "approve":
            add_notification(connection,creator_id,"Uncertainty calculation approved",
                f"Record {record_id} was approved by {current_actor()}.",link,entity,record_id)
        audit_change(connection, "cmc_revision" if record_type == "cmc" else "uncertainty_calculation",
            record_id, "auto_approve" if auto_approve else action, before=None,
            after={"status": status, "comment": body.get("comment")})
        connection.commit()
        return jsonify(ok=True, status=status)
    except ValueError as error:
        connection.rollback(); return jsonify(ok=False, error=str(error)), 422
    finally:
        connection.close()


@app.post("/uncertainty/api/records/calculation/<int:record_id>/revision")
def uncertainty_api_revision(record_id):
    require_permission("uncertainty_save")
    connection = get_connection()
    try:
        body = request.get_json(silent=True) or {}
        saved = create_revision(connection, record_id, current_actor(), body.get("reason"))
        audit_change(connection, "uncertainty_calculation", saved["id"], "create_revision",
            after={"parent_id": record_id, "revision": saved["revision"]})
        connection.commit(); return jsonify(ok=True, record=saved)
    except ValueError as error:
        connection.rollback(); return jsonify(ok=False, error=str(error)), 422
    finally:
        connection.close()


@app.post("/uncertainty/api/records/cmc/<int:record_id>/revision")
def uncertainty_api_cmc_revision(record_id):
    require_permission("uncertainty_save")
    connection = get_connection()
    try:
        body = request.get_json(silent=True) or {}
        saved = create_cmc_revision(connection, record_id, current_actor(), body.get("reason"))
        audit_change(connection, "cmc_revision", saved["id"], "create_revision",
            after={"previous_id": record_id, "revision": saved["revision"]})
        connection.commit(); return jsonify(ok=True, record=saved)
    except ValueError as error:
        connection.rollback(); return jsonify(ok=False, error=str(error)), 422
    finally:
        connection.close()


@app.get("/uncertainty/records/<int:record_id>/pdf")
def uncertainty_pdf(record_id):
    require_permission("reports")
    connection = get_connection()
    try:
        payload = record_payload(connection, record_id)
        return send_file(pdf_report(payload), mimetype="application/pdf", as_attachment=True,
            download_name=f"{payload.get('calculationId') or 'uncertainty-budget'}.pdf")
    finally:
        connection.close()


@app.get("/uncertainty/records/<int:record_id>/report")
def uncertainty_report(record_id):
    require_permission("reports")
    connection = get_connection()
    try:
        payload = record_payload(connection, record_id)
        if payload.get("status") != "Approved":
            abort(404)
        return render_template("uncertainty_report.html", report=payload,
            result=payload.get("backendResult", {}),
            standards=("ISO/IEC 17025:2017", "JCGM 100:2008", "ILAC P14:09/2020"),
            page_title=f"Approved Uncertainty · {payload.get('calculationId', '')}",
            page_subtitle="Locked approved uncertainty report", active_nav="projects")
    finally:
        connection.close()


@app.get("/uncertainty/records/<int:record_id>/csv")
def uncertainty_csv(record_id):
    require_permission("reports")
    connection = get_connection()
    try:
        payload = record_payload(connection, record_id)
        if payload.get("status") != "Approved":
            abort(404)
        output = io.StringIO(); writer = csv.writer(output)
        writer.writerow(["FlowLab Pro - Approved Measurement Uncertainty Report"])
        for label, key in (("Calculation ID", "calculationId"), ("Revision", "revision"),
                ("Status", "status"), ("Job", "jobNumber"), ("Project", "projectNumber"),
                ("Project name", "projectName"), ("Customer", "customer"), ("MUT", "mut"),
                ("MUT asset ID", "mutAssetId"), ("Serial number", "serialNumber"),
                ("Manufacturer", "manufacturer"), ("Model", "model"), ("Meter type", "meterType"),
                ("Analyst", "analyst"), ("Technical reviewer", "reviewedBy"),
                ("Approving officer", "approvedBy"), ("Approved at", "approvedAt")):
            writer.writerow([label, payload.get(key, "")])
        writer.writerow(["Applicable standards", "ISO/IEC 17025:2017; JCGM 100:2008; ILAC P14:09/2020"])
        for point in payload.get("backendResult", {}).get("points", []):
            outcome = point.get("result", {}); stats = point.get("statistics", {})
            writer.writerow([]); writer.writerow(["Flow point", point.get("label"), point.get("nominalFlow"), point.get("flowUnit")])
            writer.writerow(["Type A statistics", "n", stats.get("n"), "Mean", stats.get("mean"),
                "Standard deviation", stats.get("standard_deviation"), "Standard uncertainty", stats.get("standard_uncertainty")])
            writer.writerow(["Repeatability diagnostics", "Informational only; component effects are captured in calibration-result repeatability and are not separately added to RSS"])
            writer.writerow(["Observed quantity", "Mean", "Unit", "Sample s", "Standard uncertainty", "Relative u (%)", "Treatment"])
            for diagnostic in point.get("repeatabilityDiagnostics", []):
                writer.writerow([diagnostic.get("quantity"), diagnostic.get("mean"), diagnostic.get("unit"),
                    diagnostic.get("standard_deviation"), diagnostic.get("standard_uncertainty"),
                    diagnostic.get("relative_standard_uncertainty"), diagnostic.get("treatment")])
            writer.writerow(["Run", "Reference flow", "MUT indication", "Error (%)"])
            for index, observation in enumerate(point.get("observations", []), 1):
                writer.writerow([index, observation.get("reference_flow"), observation.get("mut"), observation.get("error_percent")])
            writer.writerow(["Source", "Type", "Origin", "Input value", "Unit", "Basis", "Distribution",
                "Divisor", "Standard uncertainty", "Sensitivity", "Contribution", "Contribution (%)", "DOF", "Evidence"])
            for component in outcome.get("components", []):
                source = component.get("input", {})
                writer.writerow([source.get("source"), source.get("source_type"), source.get("source_origin"),
                    source.get("input_value"), source.get("unit"), source.get("uncertainty_input_type"),
                    source.get("distribution"), source.get("divisor"), component.get("standard_uncertainty"),
                    source.get("sensitivity_coefficient"), component.get("contribution"),
                    component.get("contribution_percent"), source.get("degrees_of_freedom"), source.get("evidence")])
            writer.writerow(["Combined standard uncertainty", outcome.get("combined_standard_uncertainty")])
            writer.writerow(["Effective degrees of freedom", outcome.get("effective_degrees_of_freedom")])
            writer.writerow(["Coverage factor", outcome.get("coverage_factor")])
            writer.writerow(["Coverage probability (%)", outcome.get("coverage_probability")])
            writer.writerow(["Expanded uncertainty (%)", outcome.get("expanded_uncertainty")])
        data = io.BytesIO(output.getvalue().encode("utf-8-sig"))
        return send_file(data, mimetype="text/csv", as_attachment=True,
            download_name=f"{payload.get('calculationId') or 'uncertainty-report'}.csv")
    finally:
        connection.close()


@app.get("/uncertainty")
def uncertainty():
    standalone = request.args.get("standalone") == "1"
    requested_job = request.args.get("job", type=int)
    requested_record = request.args.get("record", type=int)
    connection = get_connection()
    try:
        if requested_job and not requested_record:
            unfinished = connection.execute("""SELECT id FROM uncertainty_calculations
                WHERE job_id=? AND status IN ('Draft','Reverted') AND calculated_by=?
                ORDER BY updated_at DESC,id DESC LIMIT 1""",
                (requested_job,current_actor())).fetchone()
            requested_record = unfinished[0] if unfinished else None
        jobs = [] if standalone else connection.execute("""SELECT j.id,j.job_number,c.customer_name,
            COALESCE(j.mut_description,e.equipment_name,''),COALESCE(j.mut_serial_number,e.serial_number,''),
            COALESCE(j.model_number,e.model,''),COALESCE(m.method_name,''),COALESCE(j.meter_type,''),j.status,
            COALESCE(p.project_number,''),COALESCE(p.project_name,''),COALESCE(j.mut_asset_id,e.asset_number,''),
            COALESCE(j.manufacturer,e.manufacturer,''),j.flow_min,j.flow_max,COALESCE(j.flow_unit,''),
            COALESCE(j.fluid_medium,''),COALESCE(j.calibration_quantity,''),COALESCE(j.engineer_operator,''),
            COALESCE(j.job_title,''),COALESCE(j.job_type,''),COALESCE(j.flow_point_count,1)
            FROM calibration_jobs j LEFT JOIN customers c ON c.id=j.customer_id
            LEFT JOIN laboratory_equipment e ON e.id=j.equipment_id
            LEFT JOIN calibration_methods m ON m.id=j.method_id
            LEFT JOIN project_jobs pj ON pj.job_id=j.id LEFT JOIN projects p ON p.id=pj.project_id
            WHERE j.is_deleted=0 ORDER BY j.id DESC""").fetchall()
        equipment_data = [equipment_dict(row) for row in equipment_rows(connection)]
        uncertainty_reviewers = connection.execute("""SELECT u.id,u.full_name,r.name FROM users u
            JOIN roles r ON r.id=u.role_id WHERE u.is_active=1 AND r.name IN
            ('Engineer','Supervisor','Chief Meteorologist','Administrator','HOD') ORDER BY u.full_name""").fetchall()
        uncertainty_approvers = connection.execute("""SELECT u.id,u.full_name,r.name FROM users u
            JOIN roles r ON r.id=u.role_id WHERE u.is_active=1 AND r.name IN
            ('Chief Meteorologist','Supervisor','Administrator','HOD') ORDER BY u.full_name""").fetchall()
        records = [] if standalone else connection.execute("""SELECT u.id,u.calculation_uid,j.job_number,c.customer_name,
            COALESCE(j.mut_description,e.equipment_name,''),u.method_name,u.calculation_type,
            u.calculated_at,u.analyst,u.status,u.revision,COALESCE(p.project_number,''),
            COALESCE(p.project_name,''),j.flow_range,u.calculated_by,u.expanded_uncertainty,
            u.final_job_uncertainty,COALESCE(u.reviewed_by,u.assigned_reviewer),
            COALESCE(u.approved_by,u.assigned_approver),u.approved_at,
            json_extract(u.snapshot_json,'$.backendResult.cmcComparison[0].applicableCmc'),j.id
            FROM uncertainty_calculations u JOIN calibration_jobs j ON j.id=u.job_id
            LEFT JOIN customers c ON c.id=j.customer_id LEFT JOIN laboratory_equipment e ON e.id=j.equipment_id
            LEFT JOIN project_jobs pj ON pj.job_id=j.id LEFT JOIN projects p ON p.id=pj.project_id
            WHERE u.status NOT IN ('Draft','Reverted') OR u.calculated_by=?
            ORDER BY u.id DESC""", (current_actor(),)).fetchall()
        cmc_records = connection.execute("""SELECT id,method_name,revision,status,proposed_cmc,
            coverage_factor,effective_at,created_by,created_at,snapshot_json
            FROM cmc_revisions WHERE status NOT IN ('Draft','Reverted') OR created_by=?
            ORDER BY id DESC""", (current_actor(),)).fetchall()
        audit_records = connection.execute("""SELECT entity_type,entity_id,action,actor,notes,occurred_at
            FROM approval_history WHERE entity_type IN ('uncertainty_calculation','cmc_revision')
            ORDER BY id DESC LIMIT 100""").fetchall()
        return render_template("uncertainty.html", jobs=jobs, equipment=equipment_data,
            records=records, cmc_records=cmc_records, audit_records=audit_records,
            uncertainty_reviewers=uncertainty_reviewers, uncertainty_approvers=uncertainty_approvers,
            source_templates=SOURCE_TEMPLATES, preselected_job_id=None if requested_record else requested_job,
            preselected_record_id=requested_record,standalone=standalone,
            page_title="Uncertainty Calculator" if standalone else "Uncertainty Management",
            page_subtitle="Standalone calculation tool — results are not retained" if standalone else "Measurement uncertainty budgets and laboratory CMC",
            active_nav="tools" if standalone else "iso17025")
    finally:
        connection.close()


@app.get("/tools")
def tools_uncertainty():
    modules=(
        {"title":"Uncertainty Calculator","description":"Run standalone flow measurement uncertainty calculations without creating a Job record.","url":url_for("uncertainty",standalone=1),"icon":"01","action":"Open calculator"},
        {"title":"Coveter","description":"Convert volumetric flow rates between SI, litre, US gallon, Imperial gallon, cubic-foot and petroleum barrel units.","url":url_for("flow_converter"),"icon":"02","action":"Open converter"},
    )
    return render_template("workspace.html",modules=modules,eyebrow="Engineering utilities",
        description="Independent calculation utilities that do not create controlled laboratory records.",
        page_title="Tools",page_subtitle="Standalone laboratory utilities",active_nav="tools")


@app.get("/tools/coveter")
def flow_converter():
    units=(
        ("m3_s","m³/s"),("m3_min","m³/min"),("m3_h","m³/h"),("m3_d","m³/day"),
        ("l_s","L/s"),("l_min","L/min"),("l_h","L/h"),("l_d","L/day"),
        ("ml_s","mL/s"),("cm3_s","cm³/s"),("ft3_s","ft³/s"),("ft3_min","ft³/min (CFM)"),
        ("ft3_h","ft³/h"),("usgal_s","US gal/s"),("usgal_min","US gal/min (GPM)"),
        ("usgal_h","US gal/h"),("impgal_s","Imp gal/s"),("impgal_min","Imp gal/min"),
        ("impgal_h","Imp gal/h"),("bbl_s","bbl/s"),("bbl_h","bbl/h"),
        ("bbl_d","bbl/day (BPD)"),("ml_d","ML/day"),
    )
    return render_template("converter.html",units=units,page_title="Coveter",
        page_subtitle="Volumetric flow-rate unit conversion",active_nav="tools")
    uncertainty_repository = UncertaintyRepository()
    uncertainty_nav = (
        ("Overview", (("dashboard", "Dashboard"),)),
        ("Jobs", (("jobs", "Calibration Jobs"), ("job-mut", "Job & MUT Information"),
            ("range", "Calibration Range"), ("flow-points", "Multiple Flow Points"),
            ("readiness", "Readiness"))),
        ("Calculation", (("type-a", "Type A"), ("type-b", "Type B"),
            ("budget", "Full Budget"), ("mut-result", "MUT Result"))),
        ("Capability", (("cmc", "CMC"), ("cmc-evaluation", "CMC Evaluation"),
            ("matrix", "Primary Instrument Matrix"))),
        ("Control & Evidence", (("equipment-sources", "Equipment / Certificates"),
            ("approvals", "Approvals"), ("analytics", "Analytics"),
            ("audit", "Audit"), ("history", "History"))),
    )
    view = request.args.get("view", "dashboard")
    valid_views = {key for _, items in uncertainty_nav for key, _ in items}
    if view not in valid_views:
        return redirect(url_for("uncertainty"))
    matrix_version, matrix_bands = uncertainty_repository.active_matrix()
    matrix_matches, matrix_error = (), None
    matrix_min = request.args.get("minimum", "")
    matrix_max = request.args.get("maximum", "")
    matrix_unit = request.args.get("unit", "L/min")
    if view == "matrix" and matrix_min != "" and matrix_max != "":
        try:
            matrix_matches = MatrixResolver.resolve_range(
                matrix_bands, float(matrix_min), float(matrix_max), matrix_unit)
        except ValueError as error:
            matrix_error = str(error)
    view_meta = {
        "jobs": ("Jobs", "Calibration Jobs", "Create and control job-level uncertainty work.", ("Create calibration job", "Capture MUT identity", "Assign approved method", "Define flow range", "Submit for review")),
        "job-mut": ("Jobs", "Job & MUT Information", "Controlled job and meter-under-test identification.", ("Select project and job", "Record MUT details", "Select calibration method", "Link calibration report", "Save controlled revision")),
        "range": ("Jobs", "Calibration Range", "Declare the working range and resolve the applicable approved matrix revision.", ("Enter minimum and maximum", "Select engineering unit", "Resolve matrix range", "Validate equipment coverage")),
        "flow-points": ("Jobs", "Multiple Flow Points", "Manage each calibration point and its retained observations.", ("Add flow point", "Resolve matrix equipment", "Link observations", "Calculate point result", "Review all points")),
        "readiness": ("Jobs", "Readiness", "Calculate readiness from controlled sources, certificates and observations.", ("Validate method", "Validate equipment certificates", "Check range coverage", "Check Type A observations", "Resolve required actions")),
        "type-a": ("Calculation", "Type A", "Derive statistics from retained raw repeated observations.", ("Upload calibration report", "Preserve original evidence", "Parse repeated observations", "Calculate mean and standard deviation", "Link result to flow point")),
        "budget": ("Calculation", "Full Budget", "Combine traceable Type A and Type B contributors without manual source substitution.", ("Load Type A result", "Resolve Type B sources", "Apply distributions and sensitivity", "Calculate combined uncertainty", "Retain contributor lineage")),
        "mut-result": ("Calculation", "MUT Result", "Present the customer-facing result while retaining the complete calculation.", ("Calculate effective degrees of freedom", "Resolve coverage factor", "Calculate expanded uncertainty", "Compare against CMC", "Submit result for approval")),
        "cmc": ("Capability", "Laboratory CMC", "Maintain laboratory capability separately from customer-job uncertainty.", ("Select approved method", "Load master budget", "Resolve matrix and sources", "Calculate flow-dependent profile", "Approve controlled CMC")),
        "cmc-evaluation": ("Capability", "CMC Evaluation", "Calculate a controlled CMC profile for the declared scope.", ("Define evaluation inputs", "Generate profile points", "Review capability margin", "Submit to authorized metrology", "Lock approved revision")),
        "equipment-sources": ("Control", "Equipment / Certificates", "Open authoritative equipment and certificate sources from Calibration Programme.", ("Select equipment", "Verify current certificate", "Check source basis and value", "Verify range coverage", "Expose read-only Type B source")),
        "approvals": ("Control", "Approvals", "Use the application-wide technical review and approval workflow.", WORKFLOWS["approvals"]),
        "analytics": ("Evidence", "Uncertainty Analytics", "Analyze persisted results, capability margins and dominant contributors.", ("Read approved results", "Compare CMC vs actual", "Trend uncertainty", "Identify dominant contributors")),
        "audit": ("Evidence", "Audit / Supporting Records", "Trace every result through calculations, contributors and source revisions.", ("Open calculation", "Trace contributor", "Open source record", "Verify certificate or report", "Review approval history")),
    }
    meta = view_meta.get(view, ("Uncertainty", view.replace("-", " ").title(), "Controlled uncertainty workflow.", DEFAULT_WORKFLOW))
    profiles = query("""SELECT e.asset_number, e.equipment_name, COALESCE(e.equipment_type,''),
        p.standard_uncertainty, p.coverage_factor, p.certificate_number,
        CASE WHEN e.next_calibration_date < date('now') THEN 'Expired' ELSE 'Valid' END
        FROM uncertainty_profiles p JOIN laboratory_equipment e ON e.id=p.equipment_id
        WHERE p.is_active=1 ORDER BY e.equipment_name""")
    rules = query("""SELECT r.meter_type, r.method_name, r.flow_min, r.flow_max,
        r.equipment_role, e.asset_number || ' - ' || e.equipment_name
        FROM primary_selection_rules r JOIN laboratory_equipment e ON e.id=r.equipment_id
        JOIN primary_selection_configurations c ON c.id=r.configuration_id WHERE c.is_active=1""")
    sessions = query("SELECT id, meter_type, method_name, flow_range, combined_standard_uncertainty, expanded_uncertainty, created_at FROM measurement_sessions ORDER BY id DESC LIMIT 10")
    source_records = query("""SELECT e.equipment_name, e.asset_number,
        p.standard_uncertainty, p.sensitivity, p.coverage_factor,
        COALESCE(p.certificate_number, e.certificate_path),
        CASE WHEN e.next_calibration_date IS NULL THEN 'Missing'
             WHEN e.next_calibration_date < date('now') THEN 'Expired' ELSE 'Valid' END,
        p.version, e.next_calibration_date
        FROM uncertainty_profiles p JOIN laboratory_equipment e ON e.id=p.equipment_id
        WHERE p.is_active=1 ORDER BY e.equipment_name""")
    engine_inputs, source_rows = [], []
    for record in source_records:
        name, equipment_id, value, sensitivity, _, evidence, status, version, due = record
        row = {"source": name, "type": "B", "origin": "Equipment Register",
            "equipment": equipment_id, "flow": None, "value": value, "unit": "",
            "input_type": "Standard Uncertainty", "distribution": "Controlled",
            "divisor": 1.0, "standard": value, "sensitivity": sensitivity or 1.0,
            "contribution": None, "percent": None, "dof": None, "status": status,
            "evidence": evidence, "notes": f"Profile {version}; due {due or 'not configured'}"}
        source_rows.append(row)
        if value is not None and status == "Valid":
            engine_inputs.append(UncertaintyInput(name, "B", "Equipment Register",
                value, "", "Standard Uncertainty", "Controlled", 1.0,
                sensitivity or 1.0, equipment_id=equipment_id, source_status=status,
                evidence=evidence, source_version=version))
    latest_raw = query("""SELECT flow_range, imported_observations, type_a_report_path
        FROM measurement_sessions ORDER BY id DESC LIMIT 1""")
    if latest_raw:
        try:
            observations = json.loads(latest_raw[0][1])
            type_a_input, stats = TypeAProcessor.from_observations(
                "Repeatability", observations, "", latest_raw[0][0],
                "Uploaded Report", latest_raw[0][2])
            engine_inputs.insert(0, type_a_input)
            source_rows.insert(0, {"source": "Repeatability", "type": "A",
                "origin": "Uploaded Report", "equipment": None, "flow": latest_raw[0][0],
                "value": type_a_input.input_value, "unit": "", "input_type": "Standard Uncertainty",
                "distribution": "Normal", "divisor": 1.0, "standard": type_a_input.input_value,
                "sensitivity": 1.0, "contribution": None, "percent": None,
                "dof": stats["n"] - 1, "status": "Valid", "evidence": latest_raw[0][2],
                "notes": f'n={stats["n"]}; mean={stats["mean"]:.6g}; raw observations retained'})
        except (TypeError, ValueError, json.JSONDecodeError):
            pass
    calculation = RSSEngine.calculate(engine_inputs) if engine_inputs else None
    if calculation:
        calculated = {(item.input.source, item.input.equipment_id): item for item in calculation.components}
        for row in source_rows:
            component = calculated.get((row["source"], row["equipment"]))
            if component:
                row["standard"] = component.standard_uncertainty
                row["contribution"] = component.contribution
                row["percent"] = component.contribution_percent
    return render_template("uncertainty.html", profiles=profiles, rules=rules,
        sessions=sessions, source_rows=source_rows, calculation=calculation,
        budget2_rules=uncertainty_repository.approved_budget2_rules(),
        matrix_version=matrix_version, matrix_bands=matrix_bands,
        matrix_matches=matrix_matches, matrix_error=matrix_error,
        matrix_min=matrix_min, matrix_max=matrix_max, matrix_unit=matrix_unit,
        view=view, uncertainty_nav=uncertainty_nav,
        view_meta={"stage": meta[0], "title": meta[1], "description": meta[2], "steps": meta[3]},
        page_title="Uncertainty Management",
        page_subtitle="Controlled measurement uncertainty and CMC engine",
        active_nav="iso17025")


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
