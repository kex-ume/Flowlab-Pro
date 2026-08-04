from __future__ import annotations

from pathlib import Path
import shutil
from uuid import uuid4

from PySide6.QtCore import QDate, QTime, Qt
from PySide6.QtCore import QUrl
from PySide6.QtGui import QDesktopServices
from PySide6.QtWidgets import (
    QComboBox, QDialog, QDialogButtonBox, QDoubleSpinBox, QFormLayout,
    QDateEdit, QFileDialog, QGridLayout, QHBoxLayout, QLabel, QLineEdit, QMessageBox, QPushButton, QTableWidget,
    QTableWidgetItem, QTabWidget, QVBoxLayout, QWidget,
    QTimeEdit,
)

from app.modules.uncertainty.service import UncertaintyService
from app.ui.base_page import BasePage
from app.ui.theme import ThemeManager


class _EntryDialog(QDialog):
    def __init__(self, title, fields, parent=None):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.widgets = {}
        form = QFormLayout(self)
        for key, label, kind, default in fields:
            widget = QDoubleSpinBox() if kind == "number" else QLineEdit()
            if kind == "number":
                widget.setRange(-1e12, 1e12); widget.setDecimals(10); widget.setValue(default)
            else:
                widget.setText(default)
            self.widgets[key] = widget; form.addRow(label, widget)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject)
        form.addRow(buttons)

    def value(self, key):
        widget = self.widgets[key]
        return widget.value() if isinstance(widget, QDoubleSpinBox) else widget.text().strip()


class _ProfileDialog(QDialog):
    """Controlled profile entry using equipment records rather than typed IDs."""
    def __init__(self, equipment, parent=None):
        super().__init__(parent); self.setWindowTitle("New Uncertainty Profile"); self.resize(480, 300)
        form = QFormLayout(self); self.equipment = QComboBox()
        for equipment_id, label in equipment: self.equipment.addItem(label, equipment_id)
        self.version = QLineEdit("1.0"); self.certificate = QLineEdit(); self.standard_u = QDoubleSpinBox(); self.standard_u.setRange(0, 1e12); self.standard_u.setDecimals(10)
        self.coverage = QDoubleSpinBox(); self.coverage.setRange(0.1, 10); self.coverage.setValue(2); self.coverage.setDecimals(2)
        form.addRow("Equipment *", self.equipment); form.addRow("Profile version *", self.version); form.addRow("Certificate number", self.certificate); form.addRow("Standard uncertainty *", self.standard_u); form.addRow("Coverage factor, k", self.coverage)
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel); buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject); form.addRow(buttons)


class _PrimaryRuleDialog(QDialog):
    """Maps a range to approved primary equipment without exposing database IDs."""
    def __init__(self, meter_types, methods, equipment, parent=None):
        super().__init__(parent); self.setWindowTitle("Add Primary Flow-Range Rule"); self.resize(500, 340)
        form = QFormLayout(self); self.meter = QComboBox(); self.method = QComboBox(); self.role = QComboBox(); self.equipment = QComboBox()
        self.meter.addItems(meter_types); self.method.addItems(methods); self.role.addItems(["Pump", "Tank", "Diverter", "Master Meter", "Balance", "Temperature Probe", "Pressure Gauge"])
        for equipment_id, label in equipment: self.equipment.addItem(label, equipment_id)
        self.minimum = QDoubleSpinBox(); self.maximum = QDoubleSpinBox()
        for field in (self.minimum, self.maximum): field.setRange(0, 1e12); field.setDecimals(6)
        form.addRow("Meter type *", self.meter); form.addRow("Calibration method *", self.method); form.addRow("Flow range from *", self.minimum); form.addRow("Flow range to *", self.maximum); form.addRow("Equipment role *", self.role); form.addRow("Primary equipment *", self.equipment)
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel); buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject); form.addRow(buttons)


class UncertaintyPage(BasePage):
    """Measurement-session workspace backed by the controlled uncertainty engine."""
    def __init__(self):
        super().__init__("Measurement & Uncertainty", "Measurement sessions use the approved uncertainty engine; operators do not construct uncertainty budgets.")
        self.service = UncertaintyService()
        self.tabs = QTabWidget()
        self.content_layout.addWidget(self.tabs, 1)
        self._build_calculator(); self._build_profiles(); self._build_library(); self._build_configuration(); self._build_environment_log()
        ThemeManager.register(self._apply_uncertainty_theme)
        self._apply_uncertainty_theme(ThemeManager.current())
        self.refresh()

    def _table(self, headers):
        table = QTableWidget(0, len(headers)); table.setHorizontalHeaderLabels(headers)
        table.setEditTriggers(QTableWidget.NoEditTriggers); table.setSelectionBehavior(QTableWidget.SelectRows)
        table.verticalHeader().setVisible(False); table.setAlternatingRowColors(False)
        table.horizontalHeader().setStretchLastSection(True)
        return table

    def _build_calculator(self):
        page = QWidget(); layout = QVBoxLayout(page); layout.setContentsMargins(18, 18, 18, 18); layout.setSpacing(14)
        guidance = QLabel("Start a measurement session by selecting the meter type, approved method, flow range, and Type A measurement report. The engine resolves primary equipment, retrieves approved Type B data and today's environmental record, then assembles the uncertainty budget.")
        guidance.setObjectName("CalculatorGuidance"); guidance.setWordWrap(True); layout.addWidget(guidance)
        form = QGridLayout(); form.setHorizontalSpacing(14); form.setVerticalSpacing(10)
        self.meter_type = QComboBox(); self.meter_type.setObjectName("CalculatorField")
        self.method_name = QComboBox(); self.method_name.setObjectName("CalculatorField")
        self.flow_range = QDoubleSpinBox(); self.flow_range.setObjectName("CalculatorField"); self.flow_range.setRange(0, 1e12); self.flow_range.setDecimals(6); self.flow_range.setSuffix(" flow units")
        for row, label, widget in ((0, "Meter type", self.meter_type), (1, "Calibration method", self.method_name), (2, "Flow range", self.flow_range)):
            form.addWidget(QLabel(label), row, 0); form.addWidget(widget, row, 1)
        self.upload_button = QPushButton("Upload Type A Report"); self.upload_button.clicked.connect(self._upload_type_a)
        self.report_label = QLabel("No report uploaded"); self.report_label.setObjectName("ReportLabel"); self.report_label.setWordWrap(True)
        form.addWidget(self.upload_button, 0, 2); form.addWidget(self.report_label, 0, 3, 1, 2)
        self.execute_button = QPushButton("Execute Measurement Session"); self.execute_button.setObjectName("ControlledAction"); self.execute_button.clicked.connect(self._execute_measurement)
        form.addWidget(self.execute_button, 2, 3, 1, 2); form.setColumnStretch(1, 2); form.setColumnStretch(3, 2); layout.addLayout(form)
        self.result = QLabel("Awaiting an uploaded Type A measurement report.")
        self.result.setObjectName("CalculationResult"); self.result.setWordWrap(True); layout.addWidget(self.result)
        self.budget_table = self._table(["COMPONENT", "STANDARD CONTRIBUTION", "CONTROLLED SOURCE"]); layout.addWidget(self.budget_table, 1)
        self.tabs.addTab(page, "Measurement Session")

    def _build_profiles(self):
        page = QWidget(); layout = QVBoxLayout(page); actions = QHBoxLayout()
        self.add_profile_button = QPushButton("New Active Profile"); self.add_profile_button.setObjectName("ControlledAction"); self.add_profile_button.clicked.connect(self._add_profile)
        actions.addWidget(self.add_profile_button); actions.addStretch(); layout.addLayout(actions)
        self.profile_table = self._table(["ASSET", "EQUIPMENT", "TYPE", "CLASS", "RANGE / UNIT", "STANDARD U", "K", "STATUS", "CERTIFICATE", "VIEW"]); layout.addWidget(self.profile_table, 1)
        self.tabs.addTab(page, "Uncertainty Profiles")

    def _build_library(self):
        page = QWidget(); layout = QVBoxLayout(page); actions = QHBoxLayout()
        self.add_type_b_button = QPushButton("Add Approved Type B Component"); self.add_type_b_button.setObjectName("ControlledAction"); self.add_type_b_button.clicked.connect(self._add_type_b)
        actions.addWidget(self.add_type_b_button); actions.addStretch(); layout.addLayout(actions)
        self.library_table = self._table(["COMPONENT", "VALUE", "UNIT", "DISTRIBUTION", "DIVISOR", "SENSITIVITY", "SOURCE", "REVISION", "STATUS", "METHOD"]); layout.addWidget(self.library_table, 1)
        self.tabs.addTab(page, "Type B Library")

    def _build_configuration(self):
        page = QWidget(); layout = QVBoxLayout(page)
        label = QLabel("Define each flow range and its required primary equipment. During calculation the system uses these active rules; the operator does not select primary equipment.")
        label.setObjectName("CalculatorGuidance"); label.setWordWrap(True); layout.addWidget(label)
        actions = QHBoxLayout(); self.upload_rule_button = QPushButton("Upload Decision Rule"); self.upload_rule_button.setObjectName("SecondaryAction"); self.upload_rule_button.clicked.connect(self._upload_decision_rule); self.add_rule_button = QPushButton("Add Flow-Range Rule"); self.add_rule_button.setObjectName("ControlledAction"); self.add_rule_button.clicked.connect(self._add_primary_rule); actions.addWidget(self.upload_rule_button); actions.addWidget(self.add_rule_button); actions.addStretch(); layout.addLayout(actions)
        self.rule_table = self._table(["METER TYPE", "METHOD", "FLOW FROM", "FLOW TO", "ROLE", "PRIMARY EQUIPMENT", "CONFIGURATION"]); layout.addWidget(self.rule_table, 1)
        self.tabs.addTab(page, "Primary Selection")

    def _build_environment_log(self):
        page = QWidget(); layout = QVBoxLayout(page); layout.setContentsMargins(18, 18, 18, 18); layout.setSpacing(14)
        guidance = QLabel("Daily environmental monitoring is retained as controlled compliance evidence. Measurement-session temperature is taken from the uploaded Type A report.")
        guidance.setObjectName("CalculatorGuidance"); guidance.setWordWrap(True); layout.addWidget(guidance)
        form = QGridLayout(); self.environment_date = QDateEdit(QDate.currentDate()); self.environment_date.setCalendarPopup(True)
        self.environment_time = QTimeEdit(QTime.currentTime()); self.temperature = QDoubleSpinBox(); self.humidity = QDoubleSpinBox(); self.pressure = QDoubleSpinBox(); self.environment_operator = QLineEdit(); self.environment_instrument = QLineEdit()
        self.temperature.setSuffix(" °C"); self.temperature.setRange(-100, 200); self.temperature.setDecimals(3)
        self.humidity.setSuffix(" %RH"); self.humidity.setRange(0, 100); self.humidity.setDecimals(2)
        self.pressure.setSuffix(" hPa"); self.pressure.setRange(0, 2000); self.pressure.setDecimals(2)
        fields = (("Date", self.environment_date), ("Time", self.environment_time), ("Temperature", self.temperature), ("Relative humidity", self.humidity), ("Pressure", self.pressure), ("Operator", self.environment_operator), ("Instrument used", self.environment_instrument))
        for row, (label, widget) in enumerate(fields): form.addWidget(QLabel(label), row, 0); form.addWidget(widget, row, 1)
        save = QPushButton("Save Daily Environmental Log"); save.setObjectName("ControlledAction"); save.clicked.connect(self._save_environment); form.addWidget(save, len(fields), 1)
        layout.addLayout(form); self.environment_table = self._table(["DATE", "TIME", "TEMPERATURE", "HUMIDITY", "PRESSURE", "OPERATOR", "INSTRUMENT"]); layout.addWidget(self.environment_table, 1)
        self.tabs.addTab(page, "Environmental Log")

    def refresh(self):
        current_type = self.meter_type.currentText(); current_method = self.method_name.currentText()
        self.meter_type.clear(); self.meter_type.addItem("Select meter type")
        self.meter_type.addItems([row[0] for row in self.service.repository.meter_types()])
        self.method_name.clear(); self.method_name.addItem("Select approved method")
        self.method_name.addItems([row[0] for row in self.service.repository.approved_methods()])
        self.meter_type.setCurrentText(current_type); self.method_name.setCurrentText(current_method)
        self._fill_profiles(self.service.repository.profiles())
        self._fill(self.library_table, self.service.repository.type_b_components(), lambda row: tuple(str(value or "-") for value in row[1:]))
        self._fill(self.rule_table, self.service.repository.primary_rules(), lambda row: (row[1], row[2], str(row[3]), str(row[4]), row[5], row[6], f"{row[7]} v{row[8]}"))
        self._fill(self.environment_table, self.service.repository.environmental_records(), lambda row: tuple(str(value if value is not None else "-") for value in row))

    def set_role(self, role_name: str):
        """Technicians can calculate only; controlled Type B data stays protected."""
        can_manage = role_name in {"Administrator", "Calibration Engineer"}
        self.add_profile_button.setEnabled(can_manage)
        self.add_type_b_button.setEnabled(can_manage)
        self.add_rule_button.setEnabled(can_manage)
        self.upload_rule_button.setEnabled(can_manage)

    def _upload_decision_rule(self):
        filename, _ = QFileDialog.getOpenFileName(self, "Upload Decision Rule", "", "Supported documents (*.xlsx *.xlsm *.csv *.pdf *.docx *.doc *.json *.yaml *.yml);;All files (*.*)")
        if not filename:
            return
        source = Path(filename)
        try:
            destination = Path("data/documents/uncertainty-rules")
            destination.mkdir(parents=True, exist_ok=True)
            target = destination / f"{uuid4().hex}_{source.name}"
            shutil.copy2(source, target)
            self.service.repository.add_configuration_document(source.name, str(target))
            QMessageBox.information(self, "Decision Rule", "The rule document is controlled evidence. Add or import its validated flow-range mappings before using it for measurement sessions.")
        except Exception as error:
            QMessageBox.warning(self, "Decision Rule", str(error))

    def _add_primary_rule(self):
        equipment = self.service.repository.equipment()
        if not equipment:
            QMessageBox.information(self, "Primary Selection", "Add the primary equipment to the Equipment Register first."); return
        meter_types = [row[0] for row in self.service.repository.meter_types()]
        methods = [row[0] for row in self.service.repository.approved_methods()]
        if not meter_types or not methods:
            QMessageBox.information(self, "Primary Selection", "Add approved calibration methods before configuring primary-selection rules."); return
        dialog = _PrimaryRuleDialog(meter_types, methods, equipment, self)
        if dialog.exec():
            try:
                if dialog.minimum.value() > dialog.maximum.value():
                    raise ValueError("Flow range from cannot exceed flow range to.")
                self.service.repository.add_primary_rule(dialog.meter.currentText(), dialog.method.currentText(), dialog.minimum.value(), dialog.maximum.value(), dialog.role.currentText(), dialog.equipment.currentData()); self.refresh()
            except Exception as error: QMessageBox.warning(self, "Primary Selection", str(error))

    def _fill(self, table, rows, mapper):
        table.setRowCount(len(rows))
        for index, row in enumerate(rows):
            for column, value in enumerate(mapper(row)):
                item = QTableWidgetItem(value); item.setTextAlignment(Qt.AlignVCenter | (Qt.AlignRight if column in (1, 4, 5) else Qt.AlignLeft)); table.setItem(index, column, item)
            table.setRowHeight(index, 36)

    def _fill_profiles(self, rows):
        self.profile_table.setRowCount(len(rows))
        for index, row in enumerate(rows):
            _, asset, equipment, equipment_type, is_primary, due, certificate_path, version, certificate, standard_u, k, active, range_min, range_max, unit, status = row
            operating_range = "-" if range_min is None or range_max is None else f"{range_min:g}–{range_max:g} {unit or ''}".strip()
            values = (asset, equipment, equipment_type or "-", "Primary" if is_primary else "Secondary", operating_range, f"{standard_u:.8g}" if standard_u is not None else "-", f"{k:g}", status, certificate or f"Profile v{version}")
            for column, value in enumerate(values):
                item = QTableWidgetItem(value); item.setTextAlignment(Qt.AlignVCenter | (Qt.AlignCenter if column in (3, 5, 6, 7) else Qt.AlignLeft)); self.profile_table.setItem(index, column, item)
            button = QPushButton("View certificate")
            button.setObjectName("TableAction")
            button.setEnabled(bool(certificate_path))
            if certificate_path:
                button.clicked.connect(lambda _, path=certificate_path: QDesktopServices.openUrl(QUrl.fromLocalFile(str(Path(path).resolve()))))
            self.profile_table.setCellWidget(index, 9, button); self.profile_table.setRowHeight(index, 42)

    def _add_profile(self):
        equipment = self.service.repository.equipment()
        if not equipment:
            QMessageBox.information(self, "Uncertainty Profile", "Add the primary standard to the Equipment Register first."); return
        dialog = _ProfileDialog(equipment, self)
        if dialog.exec():
            try:
                if not dialog.version.text().strip(): raise ValueError("Profile version is required.")
                self.service.repository.add_profile(dialog.equipment.currentData(), dialog.version.text().strip(), dialog.certificate.text().strip(), dialog.standard_u.value(), dialog.coverage.value()); self.refresh()
            except Exception as error: QMessageBox.warning(self, "Uncertainty Profile", str(error))

    def _add_type_b(self):
        dialog = _EntryDialog("Approved Type B Component", (("name", "Component name", "line", ""), ("value", "Value", "number", 0.0), ("unit", "Unit", "line", ""), ("distribution", "Distribution", "line", "Rectangular"), ("divisor", "Divisor", "number", 1.7320508), ("sensitivity", "Sensitivity coefficient", "number", 1.0), ("source", "Source", "line", ""), ("revision", "Revision", "line", "1.0"), ("method", "Applies to (Both, Volumetric, Mass)", "line", "Both")), self)
        if dialog.exec():
            try:
                if not dialog.value("name"): raise ValueError("Component name is required.")
                method = dialog.value("method").title()
                if method not in {"Both", "Volumetric", "Mass"}: raise ValueError("Method applicability must be Both, Volumetric, or Mass.")
                self.service.repository.add_type_b(dialog.value("name"), dialog.value("value"), dialog.value("unit"), dialog.value("distribution"), dialog.value("divisor"), dialog.value("sensitivity"), dialog.value("source"), dialog.value("revision"), method); self.refresh()
            except Exception as error: QMessageBox.warning(self, "Type B Library", str(error))

    def _execute_measurement(self):
        try:
            if self.meter_type.currentIndex() == 0 or self.method_name.currentIndex() == 0:
                raise ValueError("Select a meter type and an approved calibration method.")
            if not hasattr(self, "observations"):
                raise ValueError("Upload the Type A measurement report before executing the session.")
            result = self.service.execute_measurement(self.meter_type.currentText(), self.method_name.currentText(), self.flow_range.value(), self.observations, self.report_path, self.report_temperature)
            self.budget_table.setRowCount(len(result["components"]))
            for index, (name, contribution, source) in enumerate(result["components"]):
                for column, value in enumerate((name, f"{contribution:.8g}", source)):
                    self.budget_table.setItem(index, column, QTableWidgetItem(value))
            environment = result["environment"]
            environmental_note = f" Daily environmental log: {environment[0]}." if environment else " No daily environmental log was attached to this session."
            self.result.setText(f"Measurement session {result['session_id']} complete. Combined standard uncertainty: {result['combined']:.8g} | Expanded uncertainty (k=2): {result['expanded']:.8g}. Report temperature: {self.report_temperature:.4g} °C." + environmental_note)
        except ValueError as error: QMessageBox.warning(self, "Measurement Session", str(error))

    def _upload_type_a(self):
        filename, _ = QFileDialog.getOpenFileName(self, "Select Type A Report", "", "Excel files (*.xlsx *.xlsm);;CSV files (*.csv);;All Files (*.*)")
        if not filename:
            return
        try:
            self.observations, self.report_temperature = self._report_measurement_data(filename)
            self.report_path = filename
            self.service.type_a_from_observations(self.observations)
            self.report_label.setText(f"{Path(filename).name}: {len(self.observations)} observations; report temperature {self.report_temperature:.4g} °C")
        except Exception as error:
            QMessageBox.warning(self, "Type A Report", f"Could not calculate Type A uncertainty from this report. {error}")

    @staticmethod
    def _report_measurement_data(filename):
        path = Path(filename)
        if path.suffix.lower() == ".csv":
            import csv
            with path.open(newline="", encoding="utf-8-sig") as file:
                rows = list(csv.reader(file))
        else:
            from openpyxl import load_workbook
            workbook = load_workbook(path, data_only=True, read_only=True)
            rows = [list(row) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)]
        for row_index, row in enumerate(rows):
            headers = [str(value).strip().lower() if value is not None else "" for value in row]
            observation_index = next((index for index, value in enumerate(headers) if any(key in value for key in ("measurement", "reading", "flow", "result", "value"))), None)
            temperature_index = next((index for index, value in enumerate(headers) if "temp" in value), None)
            if observation_index is None or temperature_index is None:
                continue
            observations, temperatures = [], []
            for data_row in rows[row_index + 1:]:
                if observation_index < len(data_row) and isinstance(data_row[observation_index], (int, float)) and not isinstance(data_row[observation_index], bool): observations.append(float(data_row[observation_index]))
                if temperature_index < len(data_row) and isinstance(data_row[temperature_index], (int, float)) and not isinstance(data_row[temperature_index], bool): temperatures.append(float(data_row[temperature_index]))
            if len(observations) >= 2 and temperatures:
                return observations, sum(temperatures) / len(temperatures)
        raise ValueError("The report must contain labeled measurement/reading and temperature columns with at least two observations.")

    def _save_environment(self):
        try:
            if not self.environment_operator.text().strip() or not self.environment_instrument.text().strip():
                raise ValueError("Operator and instrument used are required.")
            self.service.repository.save_environmental_record(self.environment_date.date().toString("yyyy-MM-dd"), self.environment_time.time().toString("HH:mm:ss"), self.temperature.value(), self.humidity.value(), self.pressure.value(), self.environment_operator.text().strip(), self.environment_instrument.text().strip())
            self.refresh()
        except ValueError as error:
            QMessageBox.warning(self, "Environmental Log", str(error))

    def _apply_uncertainty_theme(self, theme):
        c = theme.color
        self.setStyleSheet(f"QTabWidget::pane {{ border: 1px solid {c.border}; }} QTabBar::tab {{ padding: 10px 16px; font-size: 12px; }} QTableWidget {{ background: {c.card}; border: 1px solid {c.border}; border-radius: 8px; color: {c.text}; font-size: 12px; }} QTableWidget::item {{ padding: 8px; border-bottom: 1px solid {c.divider}; }} QHeaderView::section {{ background: {c.window}; color: {c.text}; border: none; border-bottom: 1px solid {c.border}; padding: 10px 8px; font-size: 11px; font-weight: 700; }} QLabel#CalculatorGuidance {{ background: {c.card}; color: {c.text_secondary}; border: 1px solid {c.border}; border-radius: 8px; padding: 11px; font-size: 12px; }} QDoubleSpinBox#CalculatorField, QComboBox#CalculatorField {{ min-height: 38px; padding: 0 8px; }} QLabel#ReportLabel {{ color: {c.text_secondary}; font-size: 11px; }} QLabel#CalculationResult {{ background: {c.card}; border: 1px solid {c.border}; border-radius: 8px; padding: 12px; font-size: 13px; font-weight: 600; }} QPushButton#ControlledAction, QPushButton#SecondaryAction, QPushButton#TableAction {{ min-height: 38px; border-radius: 7px; padding: 0 16px; font-size: 12px; font-weight: 600; }} QPushButton#ControlledAction {{ background: {c.primary}; color: white; border: none; }} QPushButton#ControlledAction:hover {{ background: {c.primary_hover}; }} QPushButton#SecondaryAction, QPushButton#TableAction {{ background: {c.card}; color: {c.text}; border: 1px solid {c.border}; }} QPushButton#SecondaryAction:hover, QPushButton#TableAction:hover {{ border-color: {c.primary}; color: {c.primary}; }} QPushButton#TableAction {{ min-height: 30px; padding: 0 10px; font-size: 11px; }}")

    def closeEvent(self, event):
        ThemeManager.unregister(self._apply_uncertainty_theme)
        super().closeEvent(event)
