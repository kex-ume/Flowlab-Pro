from datetime import timedelta
import json
from pathlib import Path
import tempfile
import uuid
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.signing import TimestampSigner
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, Client as TestClient, override_settings
from django.urls import reverse
from django.utils import timezone

from . import ei_workflow as service
from .models import (Client, Job, JobDocument, JobUUT, Role, Permission, Personnel,
    PersonnelAuthorization, PersonnelCompetency, CalibrationType,
    CalibrationLogTemplate, TechnicalRecord, Notification, WorkflowAudit)
from .models import Primary


class IntegratedEIWorkflowTests(TestCase):
    def setUp(self):
        self.files = tempfile.TemporaryDirectory()
        self.override = override_settings(EI_PRIVATE_STORAGE_ROOT=self.files.name)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.addCleanup(self.files.cleanup)
        self.scope = CalibrationType.objects.create(name='Workflow Test Scope', abbreviation='WTS')
        self.standard = Primary.objects.create(name='Test reference standard', serial_no='REF-TEST', certificate_no='TEST-CERT', scope=self.scope, last_calibrated=timezone.localdate())
        self.customer = Client.objects.create(name='E&I Test Client')
        self.user, self.person = self.make_person('analyst', ['create_jobs', 'view_jobs', 'edit_calibration', 'repair_equipment'])
        self.lead, self.lead_person = self.make_person('supervisor', ['manage_ei_jobs', 'create_jobs', 'view_jobs', 'edit_calibration', 'repair_equipment', 'review_calibration', 'approve_calibration', 'manage_personnel','verify_personnel_records','assess_personnel_competency','authorize_personnel'])
        # Final signing authority is explicitly Chief, not an ordinary Supervisor.
        self.lead_person.role=Role.objects.get(name='Chief Metrologist')
        self.lead_person.save(update_fields=['role'])
        self.job_reviewer,self.job_reviewer_person=self.make_person('job-reviewer',['manage_ei_jobs','view_jobs','review_calibration'])
        self.other, self.other_person = self.make_person('other', ['create_jobs', 'view_jobs', 'edit_calibration', 'repair_equipment'])
        self.stranger, self.stranger_person = self.make_person('stranger', ['view_jobs'])
        self.template = CalibrationLogTemplate.objects.create(scope=self.scope, title='Approved test method', version=1,
            schema={'method_reference':'TEST-METHOD', 'fields': [{'name':'reading','label':'Reading','type':'number','unit':'bar','required':True,'precision':4}],
                    'tables':[{'name':'runs','label':'Runs','columns':[{'name':'reference','type':'number','required':True}], 'required':True}]},
            status='approved', approved_by=self.lead, approved_at=timezone.now(), **service.actor_fields(self.lead))
        self.web = TestClient()
        self.web.force_login(self.user)

    def make_person(self, username, permissions):
        user = User.objects.create_user(username=username, password='TestPasswordOnly123')
        role = Role.objects.create(name='Workflow ' + username)
        for codename in permissions:
            permission, _ = Permission.objects.get_or_create(codename=codename, defaults={'name': codename})
            role.permissions.add(permission)
        person = Personnel.objects.create(user=user, name=username.title(), role=role, qualifications='Test qualification', years_of_experience=5)
        if username != 'stranger':
            PersonnelCompetency.objects.create(personnel=person, scope=self.scope, method_reference='TEST-METHOD', assessed_on=timezone.localdate(), result='competent', **service.actor_fields(user))
            for activity in ('perform', 'review', 'approve'):
                PersonnelAuthorization.objects.create(personnel=person, scope=self.scope, activity=activity, effective_from=timezone.localdate(), **service.actor_fields(user))
        return user, person

    def job_payload(self, job_type='field', request_types=('calibration',), actor=None):
        return {'creation_key': str(uuid.uuid4()), 'client': self.customer.pk, 'po': '', 'job_type': job_type,
            'description':'Actual test scope work', 'received_date':str(timezone.localdate()),
            'start_date':str(timezone.localdate()), 'end_date':str(timezone.localdate()+timedelta(days=2)),
            'uuts':[{'name':'Instrument '+str(i), 'serial_number':'SN-'+str(i), 'request_type': kind,
                'scope': self.scope.pk, 'template': self.template.pk} for i, kind in enumerate(request_types)]}

    def draft(self, job_type='field', request_types=('calibration',), actor=None):
        return service.save_job(actor or self.user, self.job_payload(job_type, request_types))

    def upload(self, job, kind, actor=None):
        return service.upload_document(actor or self.user, job, SimpleUploadedFile(kind+'.pdf', b'%PDF-1.4\n%%EOF', content_type='application/pdf'), kind)

    def decide(self,job,branch,comment='',assignee_id=None):
        from .models import JobReview
        if branch!='revert' and not JobReview.objects.filter(submission=job.submissions.first()).exists():
            service.decide_job(self.job_reviewer,job,branch,comment,assignee_id)
        return service.decide_job(self.lead,job,branch,comment,assignee_id)

    def approved_job(self, job_type='field'):
        job = self.draft(job_type)
        if job_type == 'lab':
            self.upload(job, 'service_request'); self.upload(job, 'waybill')
        service.submit_job(self.user, job)
        self.decide( job, 'sender')
        service.commence_job(self.user, job)
        job.refresh_from_db()
        return job

    def test_lab_draft_allowed_submission_requires_both_actual_documents(self):
        job = self.draft('lab')
        self.assertEqual(job.approval_state, 'draft'); self.assertIsNone(job.submitted_at)
        self.assertEqual(Notification.objects.count(), 0)
        with self.assertRaises(ValidationError): service.submit_job(self.user, job)
        self.upload(job, 'service_request')
        with self.assertRaises(ValidationError): service.submit_job(self.user, job)
        self.upload(job, 'waybill'); service.submit_job(self.user, job); job.refresh_from_db()
        self.assertEqual(job.approval_state, 'awaiting'); self.assertTrue(Notification.objects.filter(recipient=self.lead).exists())

    def test_field_requires_no_documents_and_optional_request_after_approval(self):
        job = self.draft()
        with self.assertRaises(ValidationError): self.upload(job, 'service_request')
        service.submit_job(self.user, job)
        self.assertFalse(job.job_documents.exists())
        self.decide( job, 'sender'); job.refresh_from_db()
        self.upload(job, 'service_request')
        with self.assertRaises(ValidationError): self.upload(job, 'waybill')

    def test_save_and_commence_only_requests_approval(self):
        job = self.draft(); service.submit_job(self.user, job); job.refresh_from_db()
        self.assertEqual(job.record_status, 'scheduled'); self.assertIsNone(job.commenced_at)
        with self.assertRaises(PermissionDenied): service.commence_job(self.user, job)

    def test_48_hour_clock_is_from_submission_and_idempotent(self):
        job = self.draft(); draft = self.draft(); service.submit_job(self.user, job)
        now = timezone.now(); Job.objects.filter(pk=job.pk).update(submitted_at=now-timedelta(hours=48)+timedelta(minutes=1))
        service.evaluate_clocks(now); job.refresh_from_db(); self.assertEqual(job.record_status,'scheduled')
        service.evaluate_clocks(now+timedelta(minutes=1)); job.refresh_from_db(); self.assertEqual(job.record_status,'pending')
        self.assertEqual(job.approval_state,'awaiting'); draft.refresh_from_db(); self.assertEqual(draft.approval_state,'draft')
        count = Notification.objects.count(); service.evaluate_clocks(now+timedelta(minutes=1)); self.assertEqual(Notification.objects.count(),count)

    def test_approved_jobs_do_not_escalate(self):
        job = self.approved_job(); Job.objects.filter(pk=job.pk).update(submitted_at=timezone.now()-timedelta(days=3))
        service.evaluate_clocks(); job.refresh_from_db(); self.assertEqual(job.record_status,'ongoing')

    def test_sender_approval_and_authorized_commencement(self):
        job=self.draft();service.submit_job(self.user,job);self.decide(job,'sender');job.refresh_from_db()
        self.assertEqual(job.primary_assignee_id,self.person.pk)
        with self.assertRaises(PermissionDenied):service.commence_job(self.other,job)
        service.commence_job(self.user,job); count=WorkflowAudit.objects.filter(job=job,event='commenced').count()
        service.commence_job(self.user,job);self.assertEqual(WorkflowAudit.objects.filter(job=job,event='commenced').count(),count)

    def test_comment_branch_requires_comment_and_resolved_assignment(self):
        job=self.draft();service.submit_job(self.user,job)
        with self.assertRaises(ValidationError):self.decide(job,'comment','',self.person.pk)
        self.decide(job,'comment','Use the approved method',self.person.pk);job.refresh_from_db()
        self.assertEqual(job.approval_state,'approved');self.assertTrue(Notification.objects.filter(recipient=self.user,message__contains='Use the approved method').exists())

    def test_reversion_resubmission_preserves_history_and_restarts_clock(self):
        job=self.draft();service.submit_job(self.user,job)
        with self.assertRaises(ValidationError):self.decide(job,'revert','')
        self.decide(job,'revert','Fix description');job.refresh_from_db()
        with self.assertRaises(PermissionDenied):service.commence_job(self.user,job)
        service.submit_job(self.user,job);self.assertEqual(job.submissions.count(),2)
        self.assertEqual(job.submissions.last().decision.comment,'Fix description')

    def test_assign_another_notifies_sender_and_new_assignee(self):
        job=self.draft();service.submit_job(self.user,job);self.decide(job,'another','',self.other_person.pk);job.refresh_from_db()
        self.assertEqual(job.primary_assignee_id,self.other_person.pk)
        self.assertFalse(job.assigned_to.filter(user=self.user).exists())
        self.assertTrue(Notification.objects.filter(recipient=self.other).exists());self.assertTrue(Notification.objects.filter(recipient=self.user).exists())
        with self.assertRaises(PermissionDenied):service.commence_job(self.user,job)

    def test_supervisor_save_assign_and_self_assignment_policy(self):
        data=self.job_payload();data.update(assignees=[self.person.pk],primary_assignee=self.person.pk)
        job=service.save_job(self.lead,data);service.submit_job(self.lead,job,assign_now=True);job.refresh_from_db();self.assertEqual(job.approval_state,'awaiting')
        self.assertFalse(job.assigned_to.exists())
        data=self.job_payload();data.update(assignees=[self.lead_person.pk],primary_assignee=self.lead_person.pk)
        job=service.save_job(self.lead,data)
        service.submit_job(self.lead,job,assign_now=True)
        with self.assertRaises(PermissionDenied):service.decide_job(self.lead,job,'sender')
        job.refresh_from_db();self.assertEqual(job.approval_state,'awaiting')

    def test_mixed_uuts_and_parent_cannot_complete_prematurely(self):
        job=self.draft(request_types=('calibration','repair','calibration_repair'))
        self.assertEqual(job.request_type_label,'Mixed');service.submit_job(self.user,job);self.decide(job,'sender');service.commence_job(self.user,job)
        with self.assertRaises(ValidationError):service.close_uut(self.lead,job.uuts.first())
        job.refresh_from_db();self.assertEqual(job.record_status,'ongoing')

    def test_hand_filled_and_digital_sources_keep_actual_author(self):
        job=self.approved_job();uut=job.uuts.first()
        doc=service.upload_document(self.other,job,SimpleUploadedFile('assessment.pdf',b'%PDF-1.4\n%%EOF'),'assessment',uut=uut)
        record=service.start_record(self.other,uut,'assessment','upload',doc)
        self.assertEqual(record.created_by_id,self.other.pk);self.assertEqual(doc.created_by_id,self.other.pk)
        self.assertFalse(job.assigned_to.filter(user=self.other).exists())
        digital=service.start_record(self.user,uut,'calibration');self.assertEqual(digital.created_by_id,self.user.pk)

    def test_autosave_recovery_conflict_and_approved_lock(self):
        job=self.approved_job();record=service.start_record(self.user,job.uuts.first(),'calibration')
        saved=service.save_record(self.user,record,{'reading':'2.4','runs':[{'reference':'2.5'}], 'reference_standards':[self.standard.pk]},0)
        restored=TechnicalRecord.objects.get(pk=record.pk);self.assertEqual(restored.data['reading'],'2.4')
        with self.assertRaises(service.Conflict):service.save_record(self.user,record,{'reading':'9'},0)
        service.transition_record(self.user,saved,'submit');service.transition_record(self.lead,saved,'review');service.transition_record(self.lead,saved,'approve')
        with self.assertRaises(ValidationError):service.save_record(self.user,saved,{'reading':'9'},saved.version)

    def test_template_required_and_existing_run_keeps_version(self):
        job=self.approved_job();uut=job.uuts.first();uut.template=None;uut.save()
        with self.assertRaises(ValidationError):service.start_record(self.user,uut,'calibration')
        uut.template=self.template;uut.save();record=service.start_record(self.user,uut,'calibration')
        CalibrationLogTemplate.objects.create(scope=self.scope,title='New method version',version=2,schema={'fields':[]},status='approved',**service.actor_fields(self.lead))
        record.refresh_from_db();self.assertEqual(record.template_id,self.template.pk)

    def test_overdue_proposal_requires_reason_and_approval_preserves_original(self):
        job=self.approved_job();yesterday=timezone.localdate()-timedelta(days=1)
        Job.objects.filter(pk=job.pk).update(scheduled_end_date=yesterday,original_end_date=yesterday);job.refresh_from_db()
        service.evaluate_clocks();count=Notification.objects.count();service.evaluate_clocks();self.assertEqual(Notification.objects.count(),count)
        with self.assertRaises(ValidationError):service.propose_schedule(self.user,job,str(timezone.localdate()+timedelta(days=3)),'')
        proposal=service.propose_schedule(self.user,job,str(timezone.localdate()+timedelta(days=3)),'Awaiting spare parts');job.refresh_from_db();self.assertEqual(job.scheduled_end_date,yesterday)
        service.decide_schedule(self.lead,proposal,True);job.refresh_from_db();self.assertEqual(job.original_end_date,yesterday);self.assertGreater(job.scheduled_end_date,yesterday)

    def test_personnel_active_competency_authorization_and_protected_qr(self):
        job=self.draft();self.assertTrue(service.eligible(self.person,job))
        self.person.user.is_active=False;self.person.user.save();self.assertFalse(service.eligible(self.person,job));self.assertFalse(service.active_personnel().filter(pk=self.person.pk).exists())
        self.web.force_login(self.stranger);response=self.web.get(reverse('ei_personnel_verify',args=[self.person.verification_token]));self.assertEqual(response.status_code,200)
        self.assertNotContains(response,'Test qualification')

    def test_direct_wrong_user_approval_commencement_read_and_upload_denied(self):
        job=self.draft();service.submit_job(self.user,job)
        self.web.force_login(self.stranger)
        for action in ('decide','commence'):
            response=self.web.post(reverse('ei_job_action',args=[job.pk,action]),json.dumps({'decision':'sender'}),content_type='application/json');self.assertEqual(response.status_code,403)
        self.assertEqual(self.web.get(reverse('job_detail',args=[job.pk])).status_code,403)
        response=self.web.post(reverse('ei_job_upload',args=[job.pk]),{'kind':'supporting','file':SimpleUploadedFile('x.pdf',b'%PDF-1.4\n%%EOF')});self.assertEqual(response.status_code,403)

    def test_notifications_and_decisions_idempotent(self):
        job=self.draft();service.submit_job(self.user,job);count=Notification.objects.count();service.submit_job(self.user,job);self.assertEqual(Notification.objects.count(),count)
        self.decide(job,'sender');count=Notification.objects.count();self.decide(job,'sender');self.assertEqual(Notification.objects.count(),count)

    def test_invalid_documents_and_interrupted_upload_leave_no_orphan(self):
        job=self.draft('lab')
        with self.assertRaises(ValidationError):service.upload_document(self.user,job,SimpleUploadedFile('fake.pdf',b'not a pdf'),'waybill')
        with patch.object(JobDocument,'save',side_effect=RuntimeError('interrupted')):
            with self.assertRaises(RuntimeError):self.upload(job,'waybill')
        self.assertFalse(list(Path(self.files.name).glob('*')));self.assertEqual(JobDocument.objects.count(),0)

    def test_private_drafts_and_duplicate_submission_key(self):
        data=self.job_payload();a=service.save_job(self.user,data);b=service.save_job(self.user,data);self.assertEqual(a.pk,b.pk)
        self.assertTrue(service.visible_jobs(self.lead).filter(pk=a.pk).exists())
        self.assertFalse(service.visible_jobs(self.other).filter(pk=a.pk).exists())

    def test_numeric_validation_and_date_order(self):
        data=self.job_payload();data['end_date']=str(timezone.localdate()-timedelta(days=1))
        with self.assertRaises(ValidationError):service.save_job(self.user,data)
        record=service.start_record(self.user,self.approved_job().uuts.first(),'calibration')
        with self.assertRaises(ValidationError):service.save_record(self.user,record,{'reading':'NaN'},0)

    def test_csrf_and_anonymous_api(self):
        client=TestClient(enforce_csrf_checks=True);response=client.post(reverse('ei_job_save'),'{}',content_type='application/json');self.assertEqual(response.status_code,403)
        response=TestClient().post(reverse('ei_job_save'),'{}',content_type='application/json');self.assertEqual(response.status_code,401)

    def test_main_screens_render_and_form_hidden_until_request(self):
        response=self.web.get(reverse('job_list'));self.assertEqual(response.status_code,200);self.assertContains(response,'JOB RECORDS');self.assertContains(response,'ei-wizard" hidden')
        job=self.approved_job();self.assertEqual(self.web.get(reverse('job_detail',args=[job.pk])).status_code,200)
        self.assertEqual(self.web.get(reverse('ei_uut_workspace',args=[job.uuts.first().pk])).status_code,200)
        record=service.start_record(self.user,job.uuts.first(),'calibration');self.assertEqual(self.web.get(reverse('ei_record_view',args=[record.pk])).status_code,200)
        self.web.force_login(self.lead);self.assertEqual(self.web.get(reverse('ei_scope_templates',args=[self.scope.pk])).status_code,200)
        self.assertEqual(self.web.get(reverse('ei_personnel_workspace',args=[self.person.pk])).status_code,200)

    def test_lagos_timezone_display(self):
        now=timezone.now();self.assertEqual(timezone.localtime(now).utcoffset(),timedelta(hours=1))

    def test_old_qr_does_not_switch_viewer_identity(self):
        token=TimestampSigner().sign(self.user.pk)
        self.web.force_login(self.stranger)
        response=self.web.get(reverse('qr_login',args=[token]));self.assertEqual(response.status_code,403)
        self.assertEqual(int(self.web.session['_auth_user_id']),self.stranger.pk)
        self.assertEqual(TestClient().get(reverse('qr_login',args=[token])).status_code,302)

    def test_authorization_memo_and_revision_retained(self):
        self.web.force_login(self.lead)
        response=self.web.post(reverse('ei_personnel_action',args=[self.person.pk,'evidence']),{
            'kind':'authorization_memo','title':'Signed laboratory authorization',
            'file':SimpleUploadedFile('memo.pdf',b'%PDF-1.4\n%%EOF')})
        self.assertEqual(response.status_code,200);memo=response.json()['id']
        verifier,_=Permission.objects.get_or_create(codename='verify_personnel_records',defaults={'name':'Verify'})
        self.other_person.role.permissions.add(verifier)
        self.web.force_login(self.other)
        response=self.web.post(reverse('ei_personnel_action',args=[self.person.pk,'verify']),{'id':memo});self.assertEqual(response.status_code,200)
        self.web.force_login(self.lead)
        data={'activity':'perform','scope':self.scope.pk,'effective':str(timezone.localdate()),'memo':memo,'reason':'Authorization revision'}
        response=self.web.post(reverse('ei_personnel_action',args=[self.person.pk,'authorization']),json.dumps(data),content_type='application/json')
        self.assertEqual(response.status_code,200)
        new=PersonnelAuthorization.objects.get(pk=response.json()['id']);self.assertEqual(new.memo_id,memo);self.assertTrue(new.previous.revoked_at);self.assertEqual(new.previous.reason,'')
        self.assertEqual(new.version,2)

    def test_record_api_restores_and_rejects_stale_version(self):
        record=service.start_record(self.user,self.approved_job().uuts.first(),'calibration')
        url=reverse('ei_record_save',args=[record.pk])
        response=self.web.post(url,json.dumps({'version':0,'data':{'reading':'2.4','runs':[{'reference':'2.5'}],'reference_standards':[self.standard.pk]}}),content_type='application/json')
        self.assertEqual(response.status_code,200);self.assertEqual(self.web.get(url).json()['data']['reading'],'2.4')
        response=self.web.post(url,json.dumps({'version':0,'data':{'reading':'9'}}),content_type='application/json')
        self.assertEqual(response.status_code,409);self.assertEqual(self.web.get(url).json()['data']['reading'],'2.4')

    def test_template_authority_and_unsafe_formula_rejected(self):
        url=reverse('ei_template_create',args=[self.scope.pk,'create'])
        schema={'fields':[{'name':'x','type':'number'}],'formulas':[{'name':'bad','expression':'__import__("os")'}]}
        self.assertEqual(self.web.post(url,json.dumps({'schema':schema}),content_type='application/json').status_code,403)
        self.web.force_login(self.lead)
        self.assertEqual(self.web.post(url,json.dumps({'schema':schema}),content_type='application/json').status_code,400)

    def test_parent_completes_only_after_all_three_uuts_approved(self):
        job=self.draft(request_types=('calibration','repair','calibration_repair'));service.submit_job(self.user,job);self.decide(job,'sender');service.commence_job(self.user,job)
        for index,uut in enumerate(job.uuts.all()):
            kinds=['assessment']+(['repair'] if uut.request_type!='calibration' else [])+(['calibration'] if uut.request_type!='repair' else [])
            for kind in kinds:
                document=service.upload_document(self.user,job,SimpleUploadedFile(kind+'.pdf',b'%PDF-1.4\n%%EOF'),kind,uut=uut)
                record=service.start_record(self.user,uut,kind,'upload',document)
                service.transition_record(self.user,record,'submit');service.transition_record(self.lead,record,'review');service.transition_record(self.lead,record,'approve')
            service.close_uut(self.lead,uut);job.refresh_from_db()
            self.assertEqual(job.progress_status,'completed' if index==2 else 'ongoing')
        self.assertIsNotNone(job.completed_at)

    def test_controlled_amendment_preserves_previous_until_approval(self):
        job=self.approved_job();uut=job.uuts.first();doc=service.upload_document(self.user,job,SimpleUploadedFile('assessment.pdf',b'%PDF-1.4\n%%EOF'),'assessment',uut=uut)
        record=service.start_record(self.user,uut,'assessment','upload',doc)
        service.transition_record(self.user,record,'submit');service.transition_record(self.lead,record,'review');service.transition_record(self.lead,record,'approve');record.refresh_from_db()
        amendment=service.amend_record(self.lead,record,'Correct uploaded page');record.refresh_from_db();self.assertEqual(record.status,'approved');self.assertEqual(amendment.previous_id,record.pk)
        service.transition_record(self.lead,amendment,'submit')
        # A second authorized person must review an amendment authored by the supervisor.
        self.other_person.role=Role.objects.get(name='Chief Metrologist');self.other_person.save(update_fields=['role'])
        service.transition_record(self.other,amendment,'review');service.transition_record(self.other,amendment,'approve');record.refresh_from_db();self.assertEqual(record.status,'superseded')

    def test_reference_snapshot_and_report_keep_original_certificate(self):
        record=service.start_record(self.user,self.approved_job().uuts.first(),'calibration')
        record=service.save_record(self.user,record,{'reading':'2.4','runs':[{'reference':'2.5'}],'reference_standards':[self.standard.pk]},0)
        self.standard.certificate_no='CHANGED-CERT';self.standard.save()
        response=self.web.get(reverse('ei_record_report',args=[record.pk]));self.assertContains(response,'TEST-CERT');self.assertNotContains(response,'CHANGED-CERT');self.assertContains(response,'TEST-METHOD')

    def test_uploaded_excel_is_viewed_as_html_with_original_download(self):
        from openpyxl import Workbook
        import io
        workbook=Workbook();workbook.active.append(['Service Request','E&I test']);output=io.BytesIO();workbook.save(output)
        job=self.draft('lab');doc=service.upload_document(self.user,job,SimpleUploadedFile('request.xlsx',output.getvalue()),'service_request')
        response=self.web.get(reverse('ei_document_view',args=[doc.token]));self.assertEqual(response.status_code,200);self.assertContains(response,'Service Request');self.assertContains(response,'Download Original')

    def test_pdf_and_excel_exports_are_retained_idempotently(self):
        from .ei_reports import retained_export
        from openpyxl import load_workbook
        import io
        record=service.start_record(self.user,self.approved_job().uuts.first(),'calibration')
        record=service.save_record(self.user,record,{'reading':'2.4','runs':[{'reference':'2.5'}],'reference_standards':[self.standard.pk]},0)
        pdf=retained_export(self.user,record,'pdf');self.assertTrue(pdf.file.read().startswith(b'%PDF-'));pdf.file.close()
        self.assertEqual(retained_export(self.user,record,'pdf').pk,pdf.pk)
        excel=retained_export(self.user,record,'xlsx')
        with excel.file.open('rb') as handle:
            workbook=load_workbook(io.BytesIO(handle.read()));values=[str(cell.value) for row in workbook.active for cell in row]
            self.assertIn('TEST-METHOD',values);self.assertIn('TEST-CERT',values);workbook.close()
        self.assertEqual(JobDocument.objects.filter(kind='record_report').count(),2)

    def test_new_personnel_has_explicit_credentials_and_name_edit_keeps_login(self):
        self.web.force_login(self.lead)
        data={'name':'Registered Instrument Engineer','login_username':'instrument.engineer','temporary_password':'Initial-Lab-Password-2026!',
              'role':self.person.role_id,'years_of_experience':5,'qualifications':'Test engineering qualification',
              'employment_status':'active','scope_of_expertise':[self.scope.pk]}
        response=self.web.post(reverse('personnel_create'),data);self.assertEqual(response.status_code,302)
        person=Personnel.objects.get(user__username='instrument.engineer');self.assertTrue(person.user.check_password(data['temporary_password']))
        data['name']='Updated Professional Name';data.pop('temporary_password');data.pop('login_username')
        response=self.web.post(reverse('personnel_update',args=[person.pk]),data);self.assertEqual(response.status_code,302)
        person.refresh_from_db();self.assertEqual(person.user.username,'instrument.engineer');self.assertEqual(person.name,'Updated Professional Name');self.assertEqual(person.profile_history.count(),2)

    def test_retry_after_lost_autosave_ack_does_not_duplicate_revision(self):
        record=service.start_record(self.user,self.approved_job().uuts.first(),'calibration')
        data={'reading':'2.4','runs':[{'reference':'2.5'}],'reference_standards':[self.standard.pk]}
        saved=service.save_record(self.user,record,data.copy(),0)
        retried=service.save_record(self.user,record,data.copy(),0)
        self.assertEqual(retried.version,saved.version);self.assertEqual(record.revisions.count(),1)

    def test_suggested_personnel_are_not_assigned_or_given_private_draft_access(self):
        data=self.job_payload();data.update(assignees=[self.other_person.pk],primary_assignee=self.other_person.pk)
        job=service.save_job(self.user,data)
        self.assertTrue(job.suggested_personnel.filter(pk=self.other_person.pk).exists());self.assertFalse(job.assigned_to.exists())
        self.assertFalse(service.visible_jobs(self.other).filter(pk=job.pk).exists());self.assertIsNone(job.primary_assignee_id)

    def test_scope_register_accepts_minimal_form_and_retains_archived_records(self):
        permission,_=Permission.objects.get_or_create(codename='view_settings',defaults={'name':'View settings'})
        self.lead_person.role.permissions.add(permission)
        self.web.force_login(self.lead)
        url=reverse('settings_tab',args=['scopes'])
        response=self.web.post(url,{'name':'Temperature Review Scope','abbreviation':'TEMP','description':'Maintained scope'})
        self.assertEqual(response.status_code,302)
        scope=CalibrationType.objects.get(name='Temperature Review Scope')
        self.assertEqual(scope.validity_months,12)
        response=self.web.post(url,{'delete_scope':scope.pk});self.assertEqual(response.status_code,302)
        scope.refresh_from_db();self.assertFalse(scope.active)
        self.web.post(url,{'restore_scope':scope.pk});scope.refresh_from_db();self.assertTrue(scope.active)
        self.assertEqual(WorkflowAudit.objects.filter(after__scope_id=scope.pk).count(),3)
        page=self.web.get(url)
        self.assertContains(page,'Scope Records')
        self.assertContains(page,'Calibration Templates')
        self.assertContains(page,'ei-scope-form')

    def test_job_retains_scope_after_draft_reopen(self):
        job=self.draft('field');job.refresh_from_db()
        self.assertEqual(job.uuts.first().scope_id,self.scope.pk)
        self.assertContains(self.web.get(reverse('job_list')),self.scope.name)

    def test_archived_scope_cannot_be_selected_for_new_job(self):
        self.scope.active=False;self.scope.save(update_fields=['active'])
        with self.assertRaises(ValidationError):
            service.save_job(self.user,self.job_payload())

    def test_primary_responsible_is_selected_independently_from_registered_personnel(self):
        data=self.job_payload();data['primary_assignee']=self.other_person.pk
        job=service.save_job(self.user,data)
        self.assertEqual(job.suggested_primary_id,self.other_person.pk)
        self.assertFalse(job.assigned_to.exists())
        service.submit_job(self.user,job)
        self.decide(job,'another',assignee_id=self.other_person.pk)
        job.refresh_from_db();self.assertEqual(job.primary_assignee_id,self.other_person.pk)

    def test_admin_role_only_signoff_does_not_fabricate_analyst_competence(self):
        self.stranger.is_superuser=True;self.stranger.save()
        self.assertFalse(service.authorized(self.stranger_person,'perform',self.scope,'TEST-METHOD'))
        self.assertTrue(service.authorized(self.stranger_person,'approve',self.scope,'TEST-METHOD'))
        self.web.force_login(self.stranger)
        url=reverse('ei_personnel_action',args=[self.person.pk,'authorization'])
        response=self.web.post(url,{'activity':'perform','scope':self.scope.pk,'effective':str(timezone.localdate()),'reason':'Admin only'})
        self.assertEqual(response.status_code,200)

    def test_officer_permissions_are_separate_and_self_approval_blocked(self):
        self.web.force_login(self.user)
        self.assertEqual(self.web.post(reverse('ei_personnel_action',args=[self.person.pk,'competency']),{}).status_code,403)
        self.assertEqual(self.web.post(reverse('ei_personnel_action',args=[self.person.pk,'evidence']),{'title':'Unpermitted'}).status_code,403)
        self.web.force_login(self.lead)
        self.assertEqual(self.web.post(reverse('ei_personnel_action',args=[self.lead_person.pk,'authorization']),{}).status_code,403)

    def test_evidence_status_verification_permanent_qualification_and_revision(self):
        from .personnel_records import compliance
        self.web.force_login(self.lead)
        url=reverse('ei_personnel_action',args=[self.person.pk,'evidence'])
        upload=lambda:SimpleUploadedFile('certificate.pdf',b'%PDF-1.4\n%%EOF')
        response=self.web.post(url,{'kind':'qualification','title':'Academic degree','file':upload()})
        self.assertEqual(response.status_code,200);pk=response.json()['id']
        self.assertEqual(compliance(self.person)['qualification']['code'],'pending')
        verify=reverse('ei_personnel_action',args=[self.person.pk,'verify'])
        self.assertEqual(self.web.post(verify,{'id':pk}).status_code,403)
        permission=Permission.objects.get(codename='verify_personnel_records');self.other_person.role.permissions.add(permission)
        self.web.force_login(self.other);self.assertEqual(self.web.post(verify,{'id':pk}).status_code,200)
        self.assertEqual(compliance(self.person)['qualification']['code'],'active')
        self.assertIsNone(self.person.evidence.get(pk=pk).expires_on)
        self.assertEqual(self.web.post(verify,{'id':pk}).status_code,400)
        self.web.force_login(self.lead)
        response=self.web.post(url,{'kind':'qualification','title':'Degree correction','previous':pk,'history':'Correct issuer name','file':upload()})
        self.assertEqual(response.status_code,200);revision=response.json()['id']
        self.assertEqual(self.person.evidence.get(pk=pk).status,'verified')
        self.web.force_login(self.other);self.assertEqual(self.web.post(verify,{'id':revision}).status_code,200)
        self.assertEqual(self.person.evidence.get(pk=pk).status,'superseded')
        self.assertEqual(self.person.evidence.count(),2)

    def test_expiry_suspension_and_reassessment_change_actual_authorization(self):
        from .personnel_records import compliance
        grant=self.person.authorizations.filter(activity='perform').first()
        self.web.force_login(self.lead)
        response=self.web.post(reverse('ei_personnel_action',args=[self.person.pk,'suspend']),{'id':grant.pk,'reason':'Reassessment needed'})
        self.assertEqual(response.status_code,200);self.assertFalse(service.authorized(self.person,'perform',self.scope,'TEST-METHOD'))
        self.assertEqual(compliance(self.person)['authorization']['code'],'expired')
        assessment=self.person.competencies.first();assessment.expires_on=timezone.localdate()-timedelta(days=1);assessment.save()
        self.assertEqual(compliance(self.person)['competency']['code'],'expired')

    def test_qr_exposes_only_minimal_verification_and_files_stay_private(self):
        from .models import PersonnelEvidence
        evidence=PersonnelEvidence.objects.create(personnel=self.person,kind='qualification',title='Secret CV',details={'history':'Confidential history'},**service.actor_fields(self.lead))
        self.person.contact_phone='Private telephone';self.person.save()
        self.web.force_login(self.stranger)
        response=self.web.get(reverse('ei_personnel_verify',args=[self.person.verification_token]))
        self.assertContains(response,'Personnel Verification');self.assertNotContains(response,'Secret CV');self.assertNotContains(response,'Private telephone');self.assertNotContains(response,'Test qualification')
        self.assertEqual(self.web.get(reverse('ei_personnel_file',args=[evidence.pk])).status_code,403)
        self.assertEqual(self.web.get(reverse('ei_personnel_workspace',args=[self.person.pk])).status_code,403)
        self.assertEqual(TestClient().get(reverse('ei_personnel_verify',args=[self.person.verification_token])).status_code,302)

    def test_staff_correction_requests_do_not_mutate_approved_records(self):
        before=self.person.qualifications
        response=self.web.post(reverse('ei_personnel_action',args=[self.person.pk,'correction']),{'details':'Please correct qualification wording'})
        self.assertEqual(response.status_code,200);self.person.refresh_from_db();self.assertEqual(self.person.qualifications,before)
        request=self.person.correction_requests.get();self.assertEqual(request.created_by_id,self.user.pk)
        self.assertTrue(Notification.objects.filter(notification_type='personnel_review').exists())

    def test_review_reminders_are_deduplicated_and_do_not_expire_degrees(self):
        from .personnel_records import evaluate_personnel_reviews
        assessment=self.person.competencies.first();assessment.expires_on=timezone.localdate()+timedelta(days=7);assessment.save()
        evaluate_personnel_reviews();count=Notification.objects.filter(notification_type='personnel_review').count()
        self.assertGreater(count,0);evaluate_personnel_reviews();self.assertEqual(Notification.objects.filter(notification_type='personnel_review').count(),count)

    def test_personnel_directory_search_status_and_read_only_history_counts(self):
        self.web.force_login(self.user)
        response=self.web.get(reverse('personnel_list'),{'search':'Analyst','category':'qualification','status':'missing'})
        self.assertContains(response,'Not Uploaded');self.assertContains(response,'Personnel Name / ID')
        job=self.approved_job();record=service.start_record(self.user,job.uuts.first(),'calibration')
        self.web.force_login(self.lead)
        response=self.web.get(reverse('ei_personnel_workspace',args=[self.person.pk]))
        self.assertEqual(response.context['counts']['authored_records'],1)
        self.assertEqual(response.context['counts']['assigned_jobs'],1)
        self.assertEqual(record.created_by_id,self.user.pk)

    def test_designated_assessment_replaces_old_competency_without_erasing_history(self):
        self.web.force_login(self.lead)
        prior=self.person.competencies.first()
        data={'scope':self.scope.pk,'method':'TEST-METHOD','date':str(timezone.localdate()),'result':'reassess'}
        response=self.web.post(reverse('ei_personnel_action',args=[self.person.pk,'competency']),data)
        self.assertEqual(response.status_code,200)
        current=self.person.competencies.get(pk=response.json()['id'])
        self.assertEqual(current.previous_id,prior.pk);self.assertEqual(current.created_by_id,self.lead.pk)
        self.assertFalse(service.authorized(self.person,'perform',self.scope,'TEST-METHOD'))
        self.assertEqual(self.person.competencies.count(),2)

    def test_no_arbitrary_expiry_for_academic_degree_and_training_expiry_is_actual(self):
        from .personnel_records import compliance
        from .models import PersonnelEvidence
        self.web.force_login(self.lead)
        response=self.web.post(reverse('ei_personnel_action',args=[self.person.pk,'evidence']),{'kind':'qualification','title':'Permanent degree','expires':str(timezone.localdate()+timedelta(days=1))})
        self.assertEqual(response.status_code,400)
        PersonnelEvidence.objects.create(personnel=self.person,kind='training',title='Expired training certificate',status='verified',expires_on=timezone.localdate()-timedelta(days=1),**service.actor_fields(self.lead))
        self.assertEqual(compliance(self.person)['training']['code'],'expired')

    def test_suspended_personnel_cannot_be_assigned_to_restricted_job(self):
        job=self.draft();service.submit_job(self.user,job)
        grant=self.other_person.authorizations.filter(activity='perform').first();grant.suspended_at=timezone.now();grant.save()
        with self.assertRaises(ValidationError):self.decide(job,'another',assignee_id=self.other_person.pk)

    def test_legacy_analyst_validation_uses_personnel_not_auth_user_ids(self):
        from .personnel_records import validate_analyst
        validate_analyst(self.person,[self.scope])
        self.stranger.is_superuser=True;self.stranger.save()
        with self.assertRaises(PermissionDenied):validate_analyst(self.stranger_person,[self.scope])

    def test_legacy_log_privileges_and_actual_roles_are_not_confused(self):
        from .models import Waybill,WaybillItem,CalibrationLog
        bill=Waybill.objects.create(waybill_no='PERSONNEL-TEST',waybill_type='incoming')
        item=WaybillItem.objects.create(waybill=bill,equipment_description='Test UUT',quantity=1,serial_no='ROLE-TEST',scope=self.scope)
        log=CalibrationLog.objects.create(waybill_item=item,scope=self.scope,calibrated_at_date=timezone.localdate(),calibrated_by=self.person,recorded_by=self.other,status='pending_review')
        self.stranger.is_superuser=True;self.stranger.save();self.web.force_login(self.stranger)
        # Role-only signing is allowed, but a profile signature is still required.
        self.assertEqual(self.web.post(reverse('calibration_log_review',args=[log.pk])).status_code,302)
        self.assertEqual(self.web.post(reverse('api:calibrationlog-review',args=[log.pk])).status_code,400)
        log.refresh_from_db();self.assertEqual(log.status,'pending_review');self.assertEqual(log.recorded_by_id,self.other.pk);self.assertEqual(log.calibrated_by_id,self.person.pk)
        self.web.force_login(self.lead)
        response=self.web.get(reverse('ei_personnel_workspace',args=[self.person.pk]))
        self.assertEqual(response.context['counts']['legacy_logs'],1)
        self.assertContains(response,'Analyst: Analyst')
        self.assertContains(response,'Other')

    def test_renewal_after_revocation_continues_authorization_history(self):
        self.web.force_login(self.lead)
        old=self.person.authorizations.filter(activity='perform').first()
        self.web.post(reverse('ei_personnel_action',args=[self.person.pk,'revoke']),{'id':old.pk,'reason':'Old authorization revoked'})
        response=self.web.post(reverse('ei_personnel_action',args=[self.person.pk,'authorization']),{'activity':'perform','scope':self.scope.pk,'effective':str(timezone.localdate()),'reason':'Reassessed and renewed'})
        self.assertEqual(response.status_code,200)
        new=self.person.authorizations.get(pk=response.json()['id']);old.refresh_from_db()
        self.assertEqual(new.previous_id,old.pk);self.assertEqual(new.version,old.version+1);self.assertEqual(old.revocation_reason,'Old authorization revoked')

    def test_environment_navigation_redesigned_records_and_existing_export(self):
        from .models import EnvironmentRecord
        record=EnvironmentRecord.objects.create(date=timezone.localdate(),time='08',temperature='23.50',relative_humidity='45.00',recorded_by=self.user)
        response=self.web.get(reverse('environment_list'),{'from_date':str(timezone.localdate()),'to_date':str(timezone.localdate())})
        self.assertEqual(response.status_code,200);self.assertContains(response,'Environmental Logs');self.assertContains(response,'ei-env-control-grid');self.assertContains(response,'tempChart');self.assertContains(response,'humidityChart');self.assertContains(response,'23.50')
        self.assertEqual(EnvironmentRecord.objects.get(pk=record.pk).recorded_by_id,self.user.pk)
        response=self.web.get(reverse('environment_export'),{'from_date':str(timezone.localdate()),'to_date':str(timezone.localdate())})
        self.assertEqual(response.status_code,200);self.assertIn('spreadsheet',response['Content-Type'])

    def test_environment_entry_retains_permissions_and_actual_recorder(self):
        from .models import EnvironmentRecord
        permission,_=Permission.objects.get_or_create(codename='record_environment',defaults={'name':'Record environment'})
        self.person.role.permissions.add(permission)
        response=self.web.post(reverse('environment_create'),{'date':str(timezone.localdate()),'time':'12','temperature':'24.10','relative_humidity':'50.20'})
        self.assertEqual(response.status_code,302)
        record=EnvironmentRecord.objects.get(time='12');self.assertEqual(record.recorded_by_id,self.user.pk)
        self.web.force_login(self.stranger)
        response=self.web.post(reverse('environment_create'),{'date':str(timezone.localdate()),'time':'17','temperature':'24.10','relative_humidity':'50.20'})
        self.assertEqual(response.status_code,302);self.assertFalse(EnvironmentRecord.objects.filter(time='17').exists())

    def capa_record(self):
        from .models import CAPARecord
        data={'ncr_number':'ei-ncr-001','title':'Test nonconformity','source':'Internal audit','description':'Test controlled issue','owner':self.person.pk,'issued_date':str(timezone.localdate()),'target_close_date':str(timezone.localdate()+timedelta(days=5)),'issued_ncr':SimpleUploadedFile('issued.pdf',b'%PDF-1.4\n%%EOF')}
        response=self.web.post(reverse('ei_capa_create',args=['save']),data)
        self.assertEqual(response.status_code,200,response.content)
        return CAPARecord.objects.get(pk=response.json()['id'])

    def test_capa_requires_evidence_and_unique_ncr_identifier(self):
        from .models import CAPARecord
        response=self.web.post(reverse('ei_capa_create',args=['save']),{'ncr_number':'EI-NCR-X','title':'Title','source':'Audit','description':'Details','owner':self.person.pk,'issued_date':str(timezone.localdate()),'target_close_date':str(timezone.localdate())})
        self.assertEqual(response.status_code,400);self.assertEqual(CAPARecord.objects.count(),0)
        record=self.capa_record();self.assertEqual(record.ncr_number,'EI-NCR-001')
        response=self.web.post(reverse('ei_capa_action',args=[record.pk,'submit']),{'version':record.version,'reviewer':self.lead_person.pk})
        self.assertEqual(response.status_code,400)
        response=self.web.post(reverse('ei_capa_create',args=['save']),{'ncr_number':'ei-ncr-001','title':'Duplicate','source':'Audit','description':'D','owner':self.person.pk,'issued_date':str(timezone.localdate()),'target_close_date':str(timezone.localdate())})
        self.assertEqual(response.status_code,409);self.assertEqual(CAPARecord.objects.count(),1)

    def test_capa_independent_closure_lock_and_retained_documents(self):
        record=self.capa_record()
        data={'version':record.version,'owner':self.person.pk,'target_close_date':str(record.target_close_date),'immediate_correction':'Containment','root_cause':'Verified cause','corrective_action':'Corrective action','effectiveness':'Effectiveness verified','closeout_report':SimpleUploadedFile('closeout.pdf',b'%PDF-1.4\n%%EOF')}
        response=self.web.post(reverse('ei_capa_action',args=[record.pk,'save']),data);self.assertEqual(response.status_code,200,response.content);record.refresh_from_db()
        response=self.web.post(reverse('ei_capa_action',args=[record.pk,'submit']),{'version':record.version,'reviewer':self.lead_person.pk});self.assertEqual(response.status_code,200);record.refresh_from_db()
        response=self.web.post(reverse('ei_capa_action',args=[record.pk,'approve']),{'version':record.version});self.assertEqual(response.status_code,403)
        self.web.force_login(self.lead)
        self.assertEqual(self.web.get(reverse('ei_capa_detail',args=[record.pk])).status_code,200)
        response=self.web.post(reverse('ei_capa_action',args=[record.pk,'approve']),{'version':record.version});self.assertEqual(response.status_code,200);record.refresh_from_db()
        self.assertEqual(record.status,'closed');self.assertEqual(record.approved_by_id,self.lead.pk)
        self.assertEqual(record.documents.count(),2)
        self.web.force_login(self.user)
        response=self.web.post(reverse('ei_capa_action',args=[record.pk,'save']),{'version':record.version});self.assertEqual(response.status_code,400)
        self.web.force_login(self.stranger)
        self.assertEqual(self.web.get(reverse('ei_capa_document',args=[record.documents.first().pk])).status_code,404)

    def test_capa_concurrency_archive_and_dashboard_notifications(self):
        record=self.capa_record()
        response=self.web.post(reverse('ei_capa_action',args=[record.pk,'save']),{'version':0});self.assertEqual(response.status_code,409)
        self.web.force_login(self.lead)
        response=self.web.post(reverse('ei_capa_action',args=[record.pk,'archive']),{'version':record.version});self.assertEqual(response.status_code,200);record.refresh_from_db();self.assertTrue(record.archived);self.assertTrue(record.documents.exists())
        response=self.web.post(reverse('ei_capa_action',args=[record.pk,'restore']),{'version':record.version});self.assertEqual(response.status_code,200);record.refresh_from_db();self.assertFalse(record.archived)

    def test_dashboard_role_visibility_and_attention_queue(self):
        self.draft(actor=self.other)
        response=self.web.get(reverse('dashboard'));self.assertEqual(response.status_code,200)
        self.assertContains(response,'Requires Attention');self.assertContains(response,'Notifications')
        self.assertFalse(response.context['visible_jobs'].filter(creator=self.other).exists())
        job=self.draft();service.submit_job(self.user,job)
        self.web.force_login(self.lead)
        response=self.web.get(reverse('dashboard'));self.assertEqual(response.status_code,200)
        self.assertContains(response,'Jobs awaiting approval')

    def test_equipment_programme_30day_reminders_are_supervisor_only_and_deduplicated(self):
        from .equipment_records import evaluate_programme,due_status
        self.standard.calibration_due_date=timezone.localdate()+timedelta(days=30);self.standard.save()
        self.assertEqual(due_status(self.standard),'Due within 30 days')
        evaluate_programme();count=Notification.objects.filter(notification_type='primary_due_reminder').count()
        self.assertGreater(count,0);self.assertTrue(Notification.objects.filter(recipient=self.lead,notification_type='primary_due_reminder').exists())
        self.assertFalse(Notification.objects.filter(recipient=self.stranger,notification_type='primary_due_reminder').exists())
        evaluate_programme();self.assertEqual(Notification.objects.filter(notification_type='primary_due_reminder').count(),count)
        self.web.force_login(self.lead);response=self.web.get(reverse('dashboard'));self.assertContains(response,'Equipment calibration due within 30 days')
        self.assertEqual(self.web.get(reverse('ei_calibration_programme')).status_code,200)

    def test_equipment_calibration_independent_approval_updates_existing_asset(self):
        from .models import EquipmentRecord
        permission,_=Permission.objects.get_or_create(codename='view_equipment',defaults={'name':'View equipment'});self.person.role.permissions.add(permission)
        old_due=self.standard.calibration_due_date
        response=self.web.post(reverse('ei_equipment_record_add',args=[self.standard.pk,'add']),{'kind':'calibration','performed_on':str(timezone.localdate()),'provider':'External calibration laboratory','certificate_no':'CERT-EI-NEW','interval_months':12,'expanded_uncertainty':'0.005','uncertainty_unit':'bar','coverage_factor':'2','file':SimpleUploadedFile('certificate.pdf',b'%PDF-1.4\n%%EOF')})
        self.assertEqual(response.status_code,200,response.content);record=EquipmentRecord.objects.get(pk=response.json()['id']);self.standard.refresh_from_db();self.assertEqual(self.standard.calibration_due_date,old_due)
        self.assertEqual(self.web.post(reverse('ei_equipment_record_action',args=[self.standard.pk,record.pk,'approve']),{}).status_code,403)
        self.web.force_login(self.lead);response=self.web.post(reverse('ei_equipment_record_action',args=[self.standard.pk,record.pk,'approve']),{});self.assertEqual(response.status_code,200,response.content)
        self.standard.refresh_from_db();record.refresh_from_db();self.assertEqual(self.standard.certificate_no,'CERT-EI-NEW');self.assertEqual(self.standard.calibration_due_date,record.due_date);self.assertEqual(record.details['uncertainty_unit'],'bar')
        self.assertEqual(self.web.get(reverse('ei_equipment_detail',args=[self.standard.pk]),{'tab':'uncertainty'}).status_code,200)
        self.assertEqual(self.web.post(reverse('ei_equipment_record_action',args=[self.standard.pk,record.pk,'approve']),{}).status_code,400)
        self.web.force_login(self.stranger);self.assertEqual(self.web.get(reverse('ei_equipment_file',args=[record.pk])).status_code,403)

    def test_retained_impartiality_records_private_and_view_only(self):
        self.web.force_login(self.lead)
        response=self.web.post(reverse('ei_personnel_action',args=[self.person.pk,'evidence']),{'kind':'impartiality','title':'Impartiality assessment retained','file':SimpleUploadedFile('impartiality.pdf',b'%PDF-1.4\n%%EOF')})
        self.assertEqual(response.status_code,200)
        response=self.web.get(reverse('ei_retained_records'));self.assertContains(response,'Impartiality assessment retained');self.assertContains(response,'View-only register')
        self.web.force_login(self.stranger);response=self.web.get(reverse('ei_retained_records'));self.assertNotContains(response,'Impartiality assessment retained')

    def test_equipment_month_interval_handles_month_end(self):
        from datetime import date
        from .equipment_records import add_months
        self.assertEqual(add_months(date(2026,1,31),1),date(2026,2,28))

    def test_equipment_registration_unique_id_and_automatic_calendar_due_date(self):
        self.web.force_login(self.lead)
        data={'asset_number':'ei-inst-001','name':'E&I Signal Source','equipment_type':'Signal Source','laboratory_location':'Instrumentation Lab','scope':self.scope.pk,'include_programme':'on','calibration_frequency':12,'last_calibrated':str(timezone.localdate()),'status':'active','coverage_factor':'2'}
        response=self.web.post(reverse('primary_create'),data);self.assertEqual(response.status_code,302,response.context and response.context['form'].errors)
        asset=Primary.objects.get(asset_number='EI-INST-001');self.assertFalse(asset.is_reference);self.assertIsNotNone(asset.calibration_due_date)
        response=self.web.post(reverse('primary_create'),data);self.assertEqual(response.status_code,200);self.assertEqual(Primary.objects.filter(asset_number='EI-INST-001').count(),1)

    def test_support_equipment_cannot_be_used_as_reference_standard(self):
        self.standard.is_reference=False;self.standard.save(update_fields=['is_reference'])
        record=service.start_record(self.user,self.approved_job().uuts.first(),'calibration')
        with self.assertRaises(ValidationError):service.save_record(self.user,record,{'reading':'2.4','runs':[{'reference':'2.5'}],'reference_standards':[self.standard.pk]},0)

    def test_equipment_retained_reports_survive_deactivation(self):
        from .models import EquipmentRecord
        self.web.force_login(self.lead)
        response=self.web.post(reverse('primary_delete',args=[self.standard.pk]),{})
        self.assertEqual(response.status_code,302);self.standard.refresh_from_db();self.assertEqual(self.standard.status,'inactive')
        self.assertTrue(Primary.objects.filter(pk=self.standard.pk).exists())

    def test_retained_technical_record_opens_without_edit_controls(self):
        record=service.start_record(self.user,self.approved_job().uuts.first(),'calibration')
        response=self.web.get(reverse('ei_record_view',args=[record.pk]),{'readonly':'1'})
        self.assertFalse(response.context['can_edit']);self.assertFalse(response.context['can_review']);self.assertFalse(response.context['can_approve'])

    def test_dashboard_equipment_calibration_alert_is_visible_to_lab_supervisor(self):
        self.standard.calibration_due_date=timezone.localdate()-timedelta(days=1);self.standard.save()
        self.web.force_login(self.lead);response=self.web.get(reverse('dashboard'))
        self.assertContains(response,'Equipment calibration overdue / maintenance');self.assertGreater(response.context['urgent_total'],0)

    def test_retained_competency_and_authorization_snapshots_are_read_only_and_private(self):
        assessment=self.person.competencies.first();grant=self.person.authorizations.first()
        response=self.web.get(reverse('ei_personnel_snapshot',args=['competency',assessment.pk]));self.assertEqual(response.status_code,200);self.assertContains(response,'View Only');self.assertContains(response,'TEST-METHOD');self.assertNotContains(response,'Record / Revise Assessment')
        response=self.web.get(reverse('ei_personnel_snapshot',args=['authorization',grant.pk]));self.assertEqual(response.status_code,200);self.assertContains(response,'Technical Authorization');self.assertNotContains(response,'Approve / Revise Authorization')
        self.web.force_login(self.stranger);self.assertEqual(self.web.get(reverse('ei_personnel_snapshot',args=['competency',assessment.pk])).status_code,403)

    def test_environment_graph_uses_dark_palette_and_explicit_range_selection(self):
        self.web.force_login(self.lead)
        response = self.web.get(reverse('environment_list'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'ei-env-trends')
        self.assertContains(response, "return '#cde1e6'")
        self.assertContains(response, 'data-days="30" aria-pressed="true"')
        self.assertNotContains(response, 'event?.target')

    def test_lab_equipment_and_client_uut_registers_remain_separate(self):
        from .models import Waybill, WaybillItem
        permission, _ = Permission.objects.get_or_create(codename='view_equipment', defaults={'name':'View equipment'})
        self.lead_person.role.permissions.add(permission)
        incoming = Waybill.objects.create(waybill_no='CLIENT-UUT-ONLY', waybill_type='incoming')
        uut = WaybillItem.objects.create(waybill=incoming, equipment_description='Client test transmitter', serial_no='CLIENT-ONLY', quantity=1)
        self.web.force_login(self.lead)
        response = self.web.get(reverse('equipment_list'))
        self.assertContains(response, 'Lab Equipment Records')
        self.assertEqual([asset.pk for asset in response.context['assets']], [self.standard.pk])
        self.assertNotContains(response, uut.equipment_description)
        self.assertNotContains(response, 'Received Equipment / UUT')
        response = self.web.get(reverse('ei_received_equipment'))
        self.assertContains(response, 'Client Equipment (UUT)')
        self.assertContains(response, uut.equipment_description)
        self.assertNotContains(response, self.standard.name)
        response = self.web.get(reverse('ei_calibration_programme'))
        self.assertContains(response, 'Lab Calibration Programme')
        self.assertNotContains(response, uut.equipment_description)
