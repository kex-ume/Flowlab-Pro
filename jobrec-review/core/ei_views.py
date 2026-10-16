"""E&I-specific screens and CSRF-protected JSON endpoints."""
import base64
from functools import wraps
import io
import json
from pathlib import Path
import zipfile
from xml.etree import ElementTree

import qrcode
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError, OperationalError, transaction
from django.db.models import Max, Q
from django.http import FileResponse, HttpResponse, JsonResponse, Http404
from django.shortcuts import get_object_or_404, render
from django.urls import reverse
from django.utils import timezone

from .auth_utils import has_perm
from . import ei_workflow as service
from .models import (Job, JobUUT, JobDocument, JobScheduleRevision, TechnicalRecord,
    CalibrationType, CalibrationLogTemplate, CalibrationLog, Personnel,
    PersonnelEvidence, PersonnelCompetency, PersonnelAuthorization,
    PersonnelProfileRevision, Client)
from .models import Primary
from .personnel_records import flags as personnel_flags, designated, compliance, authorization_status
from .models import PersonnelCorrectionRequest


def api(methods=('POST',)):
    def decorate(fn):
        @wraps(fn)
        def wrapped(request, *args, **kwargs):
            if not request.user.is_authenticated:
                return JsonResponse({'error': 'Sign in to E&I Lab.'}, status=401)
            if not request.user.is_active:
                return JsonResponse({'error': 'Your E&I account is inactive.'}, status=403)
            if request.method not in methods:
                return JsonResponse({'error': 'Method not allowed.'}, status=405)
            request._ei_pending_files = []
            committed = False
            try:
                result = fn(request, *args, **kwargs)
                committed = True
                return result if isinstance(result, HttpResponse) else JsonResponse(result)
            except service.Conflict as exc:
                return JsonResponse({'error': '; '.join(exc.messages), 'conflict': True}, status=409)
            except PermissionDenied as exc:
                return JsonResponse({'error': str(exc)}, status=403)
            except (ValidationError, ValueError, TypeError) as exc:
                return JsonResponse({'error': '; '.join(exc.messages) if isinstance(exc, ValidationError) else str(exc)}, status=400)
            except IntegrityError:
                return JsonResponse({'error': 'A duplicate identifier or concurrent action was detected. Reload the existing record.'}, status=409)
            except OperationalError as exc:
                if 'database is locked' not in str(exc).lower() and 'database is busy' not in str(exc).lower(): raise
                return JsonResponse({'error':'Another user is saving. Your changes were not committed; reload the current version and retry.','conflict':True},status=409)
            except Http404:
                return JsonResponse({'error': 'The requested E&I record was not found.'}, status=404)
            finally:
                if not committed:
                    for storage, name in request._ei_pending_files:
                        storage.delete(name)
        return wrapped
    return decorate


def payload(request):
    if 'application/json' in request.content_type:
        result = json.loads(request.body or '{}')
        if not isinstance(result, dict):
            raise ValidationError('Request body must be a JSON object.')
        return result
    return request.POST.dict()


def job_result(job):
    job.refresh_from_db()
    return {'id': job.pk, 'public_id': str(job.public_id), 'version': job.lock_version,
            'status': job.record_status, 'approval': job.approval_state,
            'url': reverse('job_detail', args=[job.pk]), 'snapshot': service.snapshot(job)}


@api()
def eligible_lookup(request):
    if not has_perm(request.user, 'create_jobs'):
        raise PermissionDenied('Job entry permission is required.')
    rows = payload(request).get('uuts', [])
    people = []
    for person in service.active_personnel():
        allowed = True
        for row in rows:
            scope = CalibrationType.objects.filter(pk=row.get('scope')).first() if row.get('scope') else None
            permission = 'repair_equipment' if row.get('request_type') == 'repair' else 'edit_calibration'
            template = CalibrationLogTemplate.objects.filter(pk=row.get('template'), scope=scope,status='approved').first() if row.get('template') else None
            method = template.schema.get('method_reference') if template else None
            if not has_perm(person.user, permission) or (row.get('request_type') == 'calibration_repair' and not has_perm(person.user,'repair_equipment')) or (scope and not service.authorized(person, 'perform', scope, method)):
                allowed = False; break
        if allowed:
            people.append({'id':person.pk, 'name':person.name})
    return {'personnel':people}


def bootstrap(user, job=None):
    data = {'save_url': reverse('ei_job_save'), 'eligible_url': reverse('ei_eligible_lookup'), 'is_supervisor': service.supervisor(user),
            'creator': user.get_full_name() or user.username,
            'clients': list(Client.objects.exclude(status='inactive').values('id', 'name')),
            'personnel': [{'id': p.pk, 'name': p.name, 'role': p.role.name if p.role else '',
                           'supervisor': service.supervisor(p.user)} for p in service.active_personnel()],
            'scopes': list(CalibrationType.objects.filter(active=True).values('id', 'name')),
            'templates': list(CalibrationLogTemplate.objects.filter(status='approved').values('id', 'scope_id', 'title', 'version'))}
    if job:
        data['job'] = job_result(job)
        data['initial'] = {'client': job.client_id, 'po': job.purchase_order_number or '', 'description': job.job_description,
            'job_type': job.job_type, 'received_date': str(job.received_started_date or ''),
            'start_date': str(job.planned_start_date or ''), 'end_date': str(job.scheduled_end_date or ''),
            'remarks': job.remarks or '', 'assignees': list(job.suggested_personnel.values_list('id', flat=True)),
            'primary_assignee': job.suggested_primary_id, 'supervisor': job.supervisor_id,
            'uuts': [dict(id=x.pk, name=x.name, equipment_type=x.equipment_type, manufacturer=x.manufacturer,
                model=x.model, serial_number=x.serial_number, equipment_identifier=x.equipment_identifier,
                request_type=x.request_type, scope=x.scope_id, template=x.template_id, duplicate_reason=x.duplicate_reason)
                for x in job.uuts.filter(archived=False)]}
        data['documents'] = [{'kind': d.kind, 'name': d.original_name, 'url': reverse('ei_document_view', args=[d.token])} for d in job.job_documents.all()]
    return data


@api()
def job_save(request, pk=None):
    job = get_object_or_404(Job, pk=pk, workflow_managed=True) if pk else None
    job = service.save_job(request.user, payload(request), job)
    return job_result(job)


@api()
def job_action(request, pk, action):
    job = get_object_or_404(Job, pk=pk, workflow_managed=True)
    data = payload(request)
    if action in ('submit', 'assign_submit'):
        job = service.submit_job(request.user, job, action == 'assign_submit')
    elif action == 'decide':
        job = service.decide_job(request.user, job, data.get('decision'), data.get('comment', ''), data.get('assignee'))
    elif action == 'commence':
        job = service.commence_job(request.user, job)
    elif action == 'reassign':
        if not service.supervisor(request.user):
            raise PermissionDenied('Lab supervisor permission is required.')
        with transaction.atomic():
            job = Job.objects.select_for_update().get(pk=job.pk)
            service.require_job_read(request.user, job)
            if job.approval_state != 'approved' or job.progress_status == 'completed':
                raise ValidationError('Only approved, open jobs may be reassigned.')
            person = service.active_personnel().filter(pk=data.get('assignee')).first()
            if not person or not service.eligible(person, job) or not data.get('reason', '').strip():
                raise ValidationError('Select an eligible assignee and provide the reassignment reason.')
            old = service.job_recipients(job)
            before = service.snapshot(job)
            job.assigned_to.set([person]); job.primary_assignee = person; job.lock_version += 1; job.save()
            service.audit(request.user, job, 'reassigned', before, {'assignee': person.pk, 'reason': data['reason']})
            service.notify_job(job, old + service.job_recipients(job), f'Job reassigned by {request.user.username} to {person.name}: {data["reason"]}')
    elif action == 'reopen':
        if not service.supervisor(request.user):
            raise PermissionDenied('An authorized supervisor must reopen a completed job.')
        with transaction.atomic():
            job = Job.objects.select_for_update().get(pk=job.pk)
            service.require_job_read(request.user, job)
            if job.progress_status != 'completed' or not data.get('reason', '').strip():
                raise ValidationError('A completed job and controlled reopening reason are required.')
            previous = {'completed_at': job.completed_at.isoformat() if job.completed_at else None}
            job.progress_status = 'ongoing'; job.completed_at = None; job.lock_version += 1; job.save()
            job.uuts.filter(archived=False).update(status='ongoing')
            service.audit(request.user, job, 'controlled_job_reopen', previous, {'reason': data['reason']})
            service.notify_job(job, service.job_recipients(job), f'Job reopened by {request.user.username}: {data["reason"]}')
    else:
        raise ValidationError('Unknown job action.')
    return job_result(job)


@api()
def job_upload(request, pk):
    job = get_object_or_404(Job, pk=pk, workflow_managed=True)
    upload = request.FILES.get('file')
    if not upload:
        raise ValidationError('Choose a document to upload.')
    kind = request.POST.get('kind', 'supporting')
    if kind not in ('service_request', 'waybill', 'supporting', 'certificate'):
        raise ValidationError('Unknown document category.')
    doc = service.upload_document(request.user, job, upload, kind, request.POST.get('title', ''))
    return {'id': doc.pk, 'name': doc.original_name, 'url': reverse('ei_document_view', args=[doc.token])}


@login_required
def document_view(request, token):
    doc = get_object_or_404(JobDocument, token=token)
    service.require_job_read(request.user, doc.job)
    if request.GET.get('download') != '1' and Path(doc.original_name).suffix.lower() in ('.xlsx', '.docx'):
        return office_preview(request, doc.file, doc.original_name, reverse('ei_document_view', args=[doc.token]))
    response = FileResponse(doc.file.open('rb'), content_type=doc.mime_type,
                            as_attachment=request.GET.get('download') == '1', filename=doc.original_name)
    response['X-Content-Type-Options'] = 'nosniff'
    response['Content-Security-Policy'] = "default-src 'none'; sandbox"
    return response


def office_preview(request, file, name, url):
    sheets = []; paragraphs = []
    if name.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        with file.open('rb') as handle:
            workbook = load_workbook(handle, read_only=True, data_only=True)
            try:
                for sheet in workbook.worksheets:
                    rows = []
                    for row in sheet.iter_rows(max_row=min(sheet.max_row, 500), max_col=min(sheet.max_column, 30), values_only=True):
                        rows.append(['' if value is None else str(value) for value in row])
                    sheets.append({'name':sheet.title, 'rows':rows})
            finally:
                workbook.close()
    else:
        with file.open('rb') as handle, zipfile.ZipFile(handle) as archive:
            root = ElementTree.fromstring(archive.read('word/document.xml'))
            ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            paragraphs = [''.join(p.itertext()) for p in root.findall('.//w:p', ns)]
    return render(request, 'core/jobs/document_preview.html', {'name':name, 'sheets':sheets, 'paragraphs':paragraphs, 'download_url':url+'?download=1'})


@api()
def schedule_action(request, pk, revision=None):
    job = get_object_or_404(Job, pk=pk, workflow_managed=True)
    data = payload(request)
    if revision:
        change = get_object_or_404(JobScheduleRevision, pk=revision, job=job)
        change = service.decide_schedule(request.user, change, data.get('decision') == 'approve', data.get('comment', ''))
    else:
        change = service.propose_schedule(request.user, job, data.get('proposed_date'), data.get('reason', ''))
    return {'id': change.pk, 'status': change.status, 'current_end': str(job.scheduled_end_date)}


@login_required
def uut_workspace(request, pk):
    uut = get_object_or_404(JobUUT.objects.select_related('job', 'scope', 'template'), pk=pk, archived=False)
    service.require_job_read(request.user, uut.job)
    available = bool(uut.job.commenced_at and uut.job.approval_state == 'approved')
    return render(request, 'core/jobs/uut_workspace.html', {'uut': uut,
        'records': uut.technical_records.select_related('created_by', 'document', 'template'),
        'can_enter': available and uut.status != 'completed' and service.can_contribute(request.user, uut),
        'can_close': available and service.can_contribute(request.user, uut, 'approve'),
        'legacy_logs': CalibrationLog.objects.filter(waybill_item=uut.legacy_item, status='approved') if uut.legacy_item_id else [],
        'approved_templates': CalibrationLogTemplate.objects.filter(scope=uut.scope, status='approved')})


@api()
def uut_action(request, pk, action):
    uut = get_object_or_404(JobUUT, pk=pk, archived=False)
    data = payload(request)
    if action == 'template':
        service.require_work(request.user, uut)
        if uut.technical_records.filter(kind='calibration').exists():
            raise ValidationError('A recorded run retains its template; create a controlled new run/version.')
        template = get_object_or_404(CalibrationLogTemplate, pk=data.get('template'), scope=uut.scope, status='approved')
        uut.template = template; uut.save()
        service.audit(request.user, uut.job, 'uut_template_selected', after={'template': template.pk}, uut=uut)
        return {'url': reverse('ei_uut_workspace', args=[uut.pk])}
    if action == 'close':
        service.close_uut(request.user, uut)
        return {'url': reverse('job_detail', args=[uut.job_id])}
    if action == 'link_uncertainty':
        service.require_work(request.user, uut)
        log = get_object_or_404(CalibrationLog, pk=data.get('calibration'), waybill_item_id=uut.legacy_item_id, status='approved')
        if not uut.legacy_item_id or log.expanded_uncertainty is None:
            raise ValidationError('Choose an approved linked E&I calculation with a saved uncertainty result.')
        with transaction.atomic():
            record, created = TechnicalRecord.objects.get_or_create(uut=uut, kind='uncertainty', legacy_calibration=log,
                defaults={'source': 'link', 'status': 'approved', 'data': {'expanded_uncertainty': str(log.expanded_uncertainty),
                    'unit': '% (as stored by the E&I calculation)', 'coverage_factor': None,
                    'method': log.procedure.manual_reference if log.procedure_id else '',
                    'calculation_version': f'legacy-log-{log.pk}/{log.updated_at.isoformat()}',
                    'coverage_note': 'Historical coverage factor was not saved in a calculation snapshot.'}, **service.actor_fields(request.user)})
            if created:
                service.audit(request.user, uut.job, 'uncertainty_associated', after={'record': record.pk, 'calculation': log.pk}, uut=uut)
        return {'url': reverse('ei_record_view', args=[record.pk])}
    if action != 'record':
        raise ValidationError('Unknown UUT action.')
    doc = None
    if data.get('source') == 'upload':
        upload = request.FILES.get('file')
        if not upload:
            raise ValidationError('Choose the hand-filled record file.')
        service.require_work(request.user, uut)
        doc = service.upload_document(request.user, uut.job, upload, data.get('kind', 'calibration'), uut=uut)
    record = service.start_record(request.user, uut, data.get('kind'), data.get('source', 'digital'), doc, data.get('calibration_phase', 'as_left'))
    return {'id': record.pk, 'url': reverse('ei_record_view', args=[record.pk])}


def record_result(record):
    return {'id': record.pk, 'version': record.version, 'status': record.status, 'data': record.data,
            'saved_at': timezone.localtime(record.updated_at).strftime('%d %b %Y %H:%M'),
            'url': reverse('ei_record_view', args=[record.pk])}


@login_required
def record_view(request, pk):
    record = get_object_or_404(TechnicalRecord.objects.select_related('uut__job', 'template', 'document', 'created_by'), pk=pk)
    service.require_job_read(request.user, record.uut.job)
    read_only=request.GET.get('readonly')=='1'
    can_edit = not read_only and record.created_by_id == request.user.pk and record.status in ('draft', 'rework') and service.can_contribute(request.user, record.uut)
    schema = record.template.schema if record.template_id else {'fields': [
        {'name': name, 'label': name.replace('_', ' ').title(), 'type': 'text', 'required': True}
        for name in (('condition', 'as_found', 'notes') if record.kind == 'assessment' else ('diagnosis', 'actions', 'parts', 'verification', 'notes'))]}
    if record.kind == 'uncertainty':
        schema = {'fields':[{'name':key,'label':key.replace('_',' ').title(),'type':'text'} for key in record.data]}
    if record.kind == 'calibration':
        schema = dict(schema, fields=[{'name':'calibration_phase','label':'Calibration phase','type':'select','choices':['as_found','as_left']},
            {'name':'method_reference','label':'Method reference','type':'text','required':True}] + [f for f in schema.get('fields',[]) if f['name'] not in ('calibration_phase','method_reference')])
    if record.source == 'upload':
        schema = {'fields':[]}
    info = dict(record_result(record), schema=schema, editable=can_edit,
                save_url=reverse('ei_record_save', args=[record.pk]))
    if record.kind == 'calibration' and record.source == 'digital':
        info['reference_standards'] = list(Primary.objects.filter(scope=record.uut.scope, status='active',is_reference=True).values('id','name','serial_no','certificate_no'))
    return render(request, 'core/jobs/technical_record.html', {'record': record, 'record_bootstrap': info,
        'can_edit': can_edit, 'can_review': not read_only and record.created_by_id != request.user.pk and service.can_contribute(request.user, record.uut, 'review'),
        'can_approve': not read_only and record.created_by_id != request.user.pk and service.can_contribute(request.user, record.uut, 'approve')})


@api(('GET', 'POST'))
def record_save(request, pk):
    record = get_object_or_404(TechnicalRecord, pk=pk)
    service.require_job_read(request.user, record.uut.job)
    if request.method == 'POST':
        data = payload(request)
        record = service.save_record(request.user, record, data.get('data', {}), data.get('version', -1))
    return record_result(record)


@login_required
def record_report(request, pk):
    record = get_object_or_404(TechnicalRecord.objects.select_related('uut__job__client', 'template'), pk=pk)
    service.require_job_read(request.user, record.uut.job)
    if request.GET.get('format') in ('pdf','xlsx'):
        from .ei_reports import retained_export
        doc = retained_export(request.user,record,request.GET['format'])
        return FileResponse(doc.file.open('rb'),as_attachment=True,filename=doc.original_name,content_type=doc.mime_type)
    fields = record.template.schema.get('fields', []) + record.template.schema.get('formulas', []) if record.template_id else []
    labels = {f['name']: f for f in fields}
    tables = record.template.schema.get('tables', []) if record.template_id else []
    table_names = {t['name'] for t in tables}
    values = [{'label':labels.get(key, {}).get('label', key.replace('_', ' ').title()),
               'value':value, 'unit':labels.get(key, {}).get('unit', '')} for key,value in record.data.items() if key not in table_names and key not in ('reference_standards','reference_standards_snapshot')]
    rows = [{'label':table.get('label', table['name']), 'columns':table['columns'],
             'rows':[[row.get(column['name'], '') for column in table['columns']] for row in record.data.get(table['name'], [])]} for table in tables]
    return render(request, 'core/jobs/record_report.html', {'record':record, 'values':values, 'tables':rows, 'standards':record.data.get('reference_standards_snapshot',[])})


@api()
def record_action(request, pk, action):
    record = get_object_or_404(TechnicalRecord, pk=pk)
    data = payload(request)
    if action == 'replace_upload':
        service.require_work(request.user, record.uut)
        if record.source != 'upload' or record.created_by_id != request.user.pk or record.status not in ('draft', 'rework'):
            raise PermissionDenied('Only the upload draft author may replace the source document.')
        upload = request.FILES.get('file')
        if not upload:
            raise ValidationError('Choose the revised record document.')
        document = service.upload_document(request.user, record.uut.job, upload, record.kind, uut=record.uut)
        with transaction.atomic():
            record = TechnicalRecord.objects.select_for_update().get(pk=record.pk)
            record.document = document; record.version += 1; record.save()
            service.audit(request.user, record.uut.job, 'upload_record_revised', after={'record':record.pk,'document':document.pk,'version':record.version}, uut=record.uut)
    elif action == 'amend':
        record = service.amend_record(request.user, record, data.get('reason', ''))
    else:
        record = service.transition_record(request.user, record, action, data.get('reason', ''))
    return record_result(record)


@login_required
def template_workspace(request, scope_id=None):
    if not service.supervisor(request.user):
        raise PermissionDenied('Authorized method manager access is required.')
    scope = get_object_or_404(CalibrationType, pk=scope_id) if scope_id else None
    return render(request, 'core/jobs/log_templates.html', {'scope': scope,
        'scopes': CalibrationType.objects.all(), 'templates': CalibrationLogTemplate.objects.filter(scope=scope) if scope else CalibrationLogTemplate.objects.all()})


@api()
def template_action(request, scope_id, action, pk=None):
    if not service.supervisor(request.user):
        raise PermissionDenied('Authorized method manager access is required.')
    scope = get_object_or_404(CalibrationType, pk=scope_id)
    data = payload(request)
    with transaction.atomic():
        CalibrationType.objects.select_for_update().get(pk=scope.pk)
        if action == 'create':
            schema = service.validate_schema(data.get('schema', {}))
            version = (scope.log_templates.aggregate(n=Max('version'))['n'] or 0) + 1
            template = CalibrationLogTemplate.objects.create(scope=scope, title=str(data.get('title', scope.name))[:200], version=version, schema=schema, **service.actor_fields(request.user))
        elif action == 'approve':
            template = get_object_or_404(CalibrationLogTemplate.objects.select_for_update(), pk=pk, scope=scope)
            if template.status == 'approved':
                return {'id': template.pk, 'status': template.status}
            service.validate_schema(template.schema)
            template.status = 'approved'; template.approved_by = request.user; template.approved_at = timezone.now(); template.save()
        else:
            raise ValidationError('Approved templates are immutable; create a new version.')
        service.audit(request.user, None, 'template_' + action, after={'template': template.pk, 'scope': scope.pk, 'version': template.version})
    return {'id': template.pk, 'status': template.status, 'url': reverse('ei_scope_templates', args=[scope.pk])}


def require_personnel(user, person, edit=False):
    officer = any(designated(user, code) for code in ('verify_personnel_records','assess_personnel_competency','authorize_personnel'))
    if not user.is_active or not (user.pk == person.user_id and not edit or has_perm(user, 'manage_personnel') or service.supervisor(user) or officer):
        raise PermissionDenied('This personnel profile is protected.')


@login_required
def personnel_workspace(request, pk=None, token=None):
    person = get_object_or_404(Personnel, verification_token=token) if token else get_object_or_404(Personnel, pk=pk)
    if token:
        if not request.user.is_active:
            raise PermissionDenied('Sign in with an active account.')
        rows=[]
        if service.lab_authority(person.user):
            rows.append({'scope':'All E&I scopes (role-based signing policy)','activity':'Technical review / approval',
                         'status':'Active' if person.employment_status=='active' and person.user.is_active else 'Inactive'})
        for grant in person.authorizations.select_related('scope').all():
            if person.authorizations.filter(previous=grant).exists(): continue
            state = authorization_status(person,grant)
            rows.append({'scope':grant.scope.name if grant.scope else 'All authorized scopes','activity':grant.get_activity_display(),'status':state})
        return render(request,'core/personnel/verification.html',{'person':person,'grants':rows})
    require_personnel(request.user, person)
    link = request.build_absolute_uri(reverse('ei_personnel_verify', args=[person.verification_token]))
    image = qrcode.make(link); output = io.BytesIO(); image.save(output, format='PNG')
    counts = {'assigned_jobs': person.assigned_jobs.count(),
              'authored_records': TechnicalRecord.objects.filter(created_by=person.user).count(),
              'legacy_logs': CalibrationLog.objects.filter(Q(recorded_by=person.user)|Q(calibrated_by=person)|Q(reviewed_by=person.user)|Q(approved_by=person.user)).distinct().count(),
              'approved_records': TechnicalRecord.objects.filter(approved_by=person.user).count()}
    authorization_rows=list(person.authorizations.select_related('scope','memo','created_by'))
    for grant in authorization_rows: grant.current_status=authorization_status(person,grant)
    return render(request, 'core/personnel/competency.html', {'person': person, 'counts': counts,
        'qr_image': base64.b64encode(output.getvalue()).decode(), 'scopes': CalibrationType.objects.all(), 'authorization_rows':authorization_rows,
        'statuses':compliance(person), 'jobs':Job.objects.filter(Q(assigned_to=person)|Q(uuts__technical_records__created_by=person.user)|Q(uuts__technical_records__reviewed_by=person.user)|Q(uuts__technical_records__approved_by=person.user)).distinct(),
        'legacy_logs':CalibrationLog.objects.filter(Q(recorded_by=person.user)|Q(calibrated_by=person)|Q(reviewed_by=person.user)|Q(approved_by=person.user)).distinct(),
        'technical_records':TechnicalRecord.objects.filter(Q(created_by=person.user)|Q(reviewed_by=person.user)|Q(approved_by=person.user)).select_related('uut__job'),
        'evidence_sections':[{'title':'Qualifications','kind':'qualification','records':person.evidence.filter(kind='qualification')},
            {'title':'Training','kind':'training','records':person.evidence.filter(kind='training')},
            {'title':'Supporting Documents / Experience','kind':'supporting','records':person.evidence.exclude(kind__in=['qualification','training'])}],
        **personnel_flags(request.user)})


@api()
def personnel_action(request, pk, action):
    person = get_object_or_404(Personnel, pk=pk)
    require_personnel(request.user, person)
    required = {'competency':'assess_personnel_competency','authorization':'authorize_personnel','revoke':'authorize_personnel','suspend':'authorize_personnel','verify':'verify_personnel_records'}
    if action in required and not designated(request.user, required[action]):
        raise PermissionDenied('A designated personnel officer permission is required for this action.')
    if action in ('evidence','profile','resolve_correction') and not (personnel_flags(request.user)['may_manage'] or designated(request.user,'verify_personnel_records')):
        raise PermissionDenied('Only a personnel supervisor or authorized records officer can maintain these records.')
    data = payload(request)
    fields = service.actor_fields(request.user)
    with transaction.atomic():
        Personnel.objects.select_for_update().get(pk=person.pk)
        if action == 'evidence':
            kind=str(data.get('kind','qualification'))
            if kind not in ('qualification','training','experience','authorization_memo','supporting','impartiality','confidentiality','job_description','lab_role'):
                raise ValidationError('Select a valid evidence category.')
            previous=person.evidence.filter(pk=data.get('previous')).first() if data.get('previous') else None
            if data.get('previous') and (not previous or previous.kind != kind):
                raise ValidationError('Revision must refer to evidence of the same category and person.')
            if previous and previous.revisions.filter(status__in=['pending','verified']).exists():
                raise ValidationError('This evidence already has a current revision.')
            if previous and not str(data.get('history','')).strip():
                raise ValidationError('A revision reason is required.')
            expires=service.parse_date(data.get('expires'),'Evidence expiry date')
            issued=service.parse_date(data.get('date'),'Evidence date')
            if kind == 'qualification' and expires:
                raise ValidationError('Permanent academic qualifications have no expiry; use Training for time-limited certificates.')
            if expires and issued and expires < issued:
                raise ValidationError('Expiry cannot precede the record date.')
            evidence = PersonnelEvidence(personnel=person, kind=str(data.get('kind', 'qualification'))[:25], title=str(data.get('title', ''))[:200],
                details={'issuer': data.get('issuer', ''), 'date': data.get('date', ''), 'history': data.get('history', '')}, expires_on=expires,previous=previous, **fields)
            if not evidence.title.strip(): raise ValidationError('Evidence title is required.')
            upload = request.FILES.get('file')
            if upload:
                evidence.mime_type, _ = service.sniff_upload(upload)
                evidence.original_name = Path(upload.name.replace('\\', '/')).name[:200]
                evidence.file.save(evidence.original_name, upload, save=False)
                request._ei_pending_files.append((evidence.file.storage, evidence.file.name))
            try:
                evidence.save()
            except Exception:
                if evidence.file.name:
                    evidence.file.storage.delete(evidence.file.name)
                raise
            result = {'id': evidence.pk}
        elif action == 'competency':
            if person.user_id == request.user.pk:
                raise PermissionDenied('An assessor cannot assess their own competency.')
            scope = get_object_or_404(CalibrationType, pk=data.get('scope'))
            assessed = service.parse_date(data.get('date'), 'Assessment date', True)
            expires = service.parse_date(data.get('expires'), 'Reassessment date')
            if assessed > timezone.localdate() or expires and expires < assessed:
                raise ValidationError('Assessment cannot be future dated; expiry must follow assessment.')
            evidence = person.evidence.filter(pk=data.get('evidence')).first() if data.get('evidence') else None
            if data.get('evidence') and not evidence:
                raise ValidationError('Assessment evidence must belong to this person.')
            if data.get('result') not in ('competent', 'reassess') or not data.get('method', '').strip():
                raise ValidationError('A method reference and assessment result are required.')
            prior=person.competencies.filter(scope=scope,method_reference=data['method'][:150],revisions__isnull=True).order_by('-assessed_on','-pk').first()
            row = PersonnelCompetency.objects.create(personnel=person, scope=scope, method_reference=data['method'][:150], assessed_on=assessed, previous=prior,
                expires_on=expires, result=data['result'], evidence=evidence, **fields)
            result = {'id': row.pk}
        elif action == 'authorization':
            if person.user_id == request.user.pk:
                raise PermissionDenied('An officer cannot approve their own technical authorization.')
            activity = data.get('activity')
            if activity not in ('perform', 'review', 'approve', 'supervise'):
                raise ValidationError('Choose a valid laboratory activity.')
            scope = get_object_or_404(CalibrationType, pk=data['scope']) if data.get('scope') else None
            start = service.parse_date(data.get('effective'), 'Effective date', True)
            end = service.parse_date(data.get('expires'), 'Expiry date')
            if end and end < start:
                raise ValidationError('Expiry cannot precede the effective date.')
            memo = person.evidence.filter(pk=data.get('memo')).first() if data.get('memo') else None
            if data.get('memo') and not memo:
                raise ValidationError('Authorization memo must belong to this personnel profile.')
            if memo and (memo.kind != 'authorization_memo' or memo.status != 'verified'):
                raise ValidationError('Select a verified authorization letter belonging to this person.')
            if not str(data.get('reason','')).strip(): raise ValidationError('An authority reference / revision reason is required.')
            method = str(data.get('method_reference','')).strip()[:150]
            prior = person.authorizations.filter(activity=activity, scope=scope, method_reference=method).order_by('-version','-pk').first()
            if prior and not prior.revoked_at:
                prior.revoked_at = timezone.now(); prior.revoked_by = request.user; prior.revocation_reason = 'Superseded by formal authorization revision'; prior.save()
            row = PersonnelAuthorization.objects.create(personnel=person, scope=scope, activity=activity, method_reference=method, effective_from=start,
                expires_on=end, memo=memo, previous=prior, version=(prior.version + 1 if prior else 1), reason=data.get('reason', ''), **fields)
            result = {'id': row.pk}
        elif action in ('revoke','suspend'):
            row = get_object_or_404(PersonnelAuthorization, pk=data.get('id'), personnel=person)
            if not data.get('reason', '').strip():
                raise ValidationError('Revocation reason is required.')
            if row.revoked_at: raise ValidationError('A revoked authorization cannot be changed; issue a controlled revision.')
            service.audit(request.user,None,'personnel_authorization_'+action+'_before',before={'authorization':row.pk,'suspended_at':str(row.suspended_at),'revoked_at':str(row.revoked_at)})
            if action == 'suspend':
                if row.suspended_at: raise ValidationError('Already suspended; issue a controlled revision to restore authorization.')
                row.suspended_at=timezone.now();row.suspended_by=request.user;row.suspension_reason=data['reason']
            else:
                row.revoked_at = timezone.now(); row.revoked_by = request.user; row.revocation_reason = data['reason']
            row.save()
            from .models import Notification
            recipients=set(Personnel.objects.filter(user__is_active=True,role__permissions__codename='authorize_personnel').values_list('user_id',flat=True))
            recipients.add(person.user_id)
            for recipient in recipients:
                Notification.objects.create(recipient_id=recipient,notification_type='personnel_review',
                    message=f'Personnel {person.pk} · {person.name}: authorization {action}d. Reason: {data["reason"]}',link=reverse('ei_personnel_verify',args=[person.verification_token]))
            result = {'id': row.pk}
        elif action == 'verify':
            row=get_object_or_404(PersonnelEvidence,pk=data.get('id'),personnel=person)
            if row.status != 'pending': raise ValidationError('Verified evidence is view-only; create a revision.')
            if row.created_by_id == request.user.pk: raise PermissionDenied('The uploader cannot verify their own evidence.')
            if not row.file: raise ValidationError('Upload supporting evidence before verification.')
            issued=service.parse_date(row.details.get('date'),'Evidence date')
            if issued and issued > timezone.localdate(): raise ValidationError('Future-dated qualifications / training cannot be verified as completed.')
            row.status='verified';row.verified_by=request.user;row.verified_at=timezone.now();row.save()
            if row.previous:
                row.previous.status='superseded';row.previous.save(update_fields=['status'])
            result={'id':row.pk}
        elif action == 'correction':
            details=str(data.get('details','')).strip()
            if not details: raise ValidationError('Describe the requested correction.')
            row=PersonnelCorrectionRequest.objects.create(personnel=person,details=details,**fields)
            from .models import Notification
            for officer in Personnel.objects.filter(user__is_active=True,role__permissions__codename__in=['manage_personnel','verify_personnel_records']).distinct():
                Notification.objects.create(recipient=officer.user,notification_type='personnel_review',message=f'Personnel {person.pk} · {person.name}: correction request awaiting review.',link=reverse('ei_personnel_workspace',args=[person.pk]))
            result={'id':row.pk}
        elif action == 'resolve_correction':
            row=get_object_or_404(PersonnelCorrectionRequest,pk=data.get('id'),personnel=person,status='pending')
            resolution=str(data.get('resolution','')).strip()
            if not resolution: raise ValidationError('Resolution is required; make record changes through the controlled revision actions.')
            row.status='resolved';row.resolution=resolution;row.resolved_by=request.user;row.resolved_at=timezone.now();row.save()
            result={'id':row.pk}
        elif action == 'profile':
            require_personnel(request.user, person, edit=True)
            before = {name: str(getattr(person, name)) for name in ('name', 'position', 'department', 'qualifications', 'years_of_experience', 'employment_status')}
            for name in ('position', 'department', 'contact_phone', 'qualifications'):
                value = str(data.get(name, getattr(person, name)))
                maximum = person._meta.get_field(name).max_length
                if maximum and len(value) > maximum:
                    raise ValidationError(f'{name} exceeds the allowed length of {maximum}.')
                setattr(person, name, value)
            if data.get('employment_status') in ('active', 'inactive'):
                person.employment_status = data['employment_status']
            if 'years_of_experience' in data:
                person.years_of_experience = int(data['years_of_experience'])
            if person.years_of_experience < 0:
                raise ValidationError('Years of experience cannot be negative.')
            person.experience_as_of = service.parse_date(data.get('as_of'), 'Experience as-of date')
            person.save()
            PersonnelProfileRevision.objects.create(personnel=person, snapshot={'before': before, 'after': data}, **fields)
            result = {'id': person.pk}
        else:
            raise ValidationError('Unknown personnel action.')
        service.audit(request.user, None, 'personnel_' + action, after={'personnel': person.pk, **result})
    return dict(result, url=reverse('ei_personnel_workspace', args=[person.pk]))


@login_required
def personnel_file(request, pk):
    evidence = get_object_or_404(PersonnelEvidence, pk=pk)
    require_personnel(request.user, evidence.personnel)
    if not evidence.file:
        raise ValidationError('This entry has no uploaded file.')
    if request.GET.get('download') != '1' and Path(evidence.original_name).suffix.lower() in ('.xlsx', '.docx'):
        return office_preview(request, evidence.file, evidence.original_name, reverse('ei_personnel_file', args=[evidence.pk]))
    response = FileResponse(evidence.file.open('rb'), content_type=evidence.mime_type,
        filename=evidence.original_name, as_attachment=request.GET.get('download') == '1')
    response['X-Content-Type-Options'] = 'nosniff'
    response['Content-Security-Policy'] = "default-src 'none'; sandbox"
    return response
