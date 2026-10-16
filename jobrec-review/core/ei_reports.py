"""Controlled E&I exports from recorded inputs; no new uncertainty evaluation."""
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from django.conf import settings
from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone

from .models import JobDocument, TechnicalRecord, CompanyProfile
from . import ei_workflow as service


def report_data(record):
    schema = record.template.schema if record.template_id else {}
    fields = {f['name']: f for f in schema.get('fields', []) + schema.get('formulas', [])}
    tables = schema.get('tables', [])
    excluded = {t['name'] for t in tables} | {'reference_standards','reference_standards_snapshot'}
    values = [(fields.get(key, {}).get('label', key.replace('_',' ').title()), value,
               fields.get(key, {}).get('unit', '')) for key,value in record.data.items() if key not in excluded]
    return values, tables, record.data.get('reference_standards_snapshot', [])


def xlsx_bytes(record):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    workbook = Workbook(); sheet = workbook.active; sheet.title = 'Controlled Record'
    rows = [
        ['Engineering Automation Technology Limited'], [record.get_kind_display() + ' Report'],
        ['Report', f'EI-REC-{record.pk}', 'Revision', record.version], ['State', record.get_status_display()],
        ['Job', record.uut.job.job_code, 'Client', record.uut.job.client.name],
        ['PO', record.uut.job.purchase_order_number or ''], ['Equipment', record.uut.name, 'Serial', record.uut.serial_number],
        ['Scope', record.uut.scope.name if record.uut.scope_id else ''], ['Analyst / uploader', record.actor_name],
        ['Template', str(record.template) if record.template_id else record.get_source_display()],
        ['Reviewer', record.reviewed_by.get_full_name() or record.reviewed_by.username if record.reviewed_by_id else 'Not reviewed'],
        ['Approver', record.approved_by.get_full_name() or record.approved_by.username if record.approved_by_id else 'Not approved'],
        [], ['Recorded information / result', 'Value', 'Unit'],
    ]
    values,tables,standards = report_data(record)
    rows += [[label, str(value) if not isinstance(value,(int,float)) else value, unit] for label,value,unit in values]
    for row in rows:
        sheet.append(row)
    for table in tables:
        sheet.append([]); sheet.append([table.get('label', table['name'])]); sheet.append([c.get('label',c['name']) + (' ('+c['unit']+')' if c.get('unit') else '') for c in table['columns']])
        for row in record.data.get(table['name'], []):
            sheet.append([row.get(c['name'], '') for c in table['columns']])
    if standards:
        sheet.append([]); sheet.append(['Reference Standards and Traceability']); sheet.append(['Reference', 'Serial', 'Certificate', 'Calibration date', 'Due date'])
        for standard in standards:
            sheet.append([standard[k] for k in ('name','serial','certificate','calibrated','due')])
    sheet.append([]); sheet.append(['EATL · E&I Laboratory · Controlled technical record']); sheet.append(['FlowLab Pro'])
    company = CompanyProfile.objects.first()
    logo = Path(company.logo.path) if company and company.logo else Path(settings.BASE_DIR)/'core/static/core/images/image2.jpeg'
    if logo.is_file():
        from openpyxl.drawing.image import Image as SheetImage
        sheet.insert_rows(1,4)
        image = SheetImage(str(logo)); image.width = 420; image.height = 81; sheet.add_image(image,'A1')
        for index in range(1,5):sheet.row_dimensions[index].height=20
        sheet['A5']=''
        sheet['D1']=f'EI-REC-{record.pk} · Rev {record.version}'
        sheet['D2']=timezone.localtime(record.updated_at).strftime('%d %b %Y')
    for row in sheet:
        for cell in row:
            # User-entered labels/values must never become spreadsheet formulas.
            if cell.data_type == 'f' or isinstance(cell.value,str) and cell.value[:1] in ('=','+','-','@'):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.font = Font(name='Calibri', size=11, color='153344')
    for index in ([6,18] if logo.is_file() else [1,2,14]):
        for cell in sheet[index]:
            cell.font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='153344')
    for column in range(1,sheet.max_column+1):
        sheet.column_dimensions[get_column_letter(column)].width = 30
    sheet.freeze_panes = 'A15'; sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A4; sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
    sheet.oddFooter.center.text = '&IFlowLab Pro'; sheet.oddFooter.right.text = 'Page &P of &N'
    output = BytesIO(); workbook.save(output); return output.getvalue()


def pdf_bytes(record):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, Image
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    output = BytesIO(); styles = getSampleStyleSheet()
    styles.add(ParagraphStyle('RecordText', fontName='Helvetica', fontSize=8, leading=11, textColor=colors.HexColor('#153344')))
    font = Path('C:/Windows/Fonts/arial.ttf')
    if font.is_file():
        if 'EATLRecordUnicode' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('EATLRecordUnicode',str(font)))
        styles['RecordText'].fontName = 'EATLRecordUnicode'
    styles.add(ParagraphStyle('RecordRight', parent=styles['RecordText'], alignment=TA_RIGHT))
    def text(value, style='RecordText'):
        return Paragraph(escape(str(value)).replace('\n','<br/>'), styles[style])
    width = A4[0] - 30*mm
    company = CompanyProfile.objects.first()
    logo = Path(company.logo.path) if company and company.logo else Path(settings.BASE_DIR)/'core/static/core/images/image2.jpeg'
    logo_item = Image(str(logo), width=90*mm, height=18*mm, kind='proportional') if logo.is_file() else text('Engineering Automation Technology Limited')
    header = Table([[logo_item, text(f'EI-REC-{record.pk} · Rev {record.version}\n{timezone.localtime(record.updated_at):%d %b %Y}', 'RecordRight')]], colWidths=[width*.66,width*.34])
    header.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'), ('LEFTPADDING',(0,0),(-1,-1),0), ('RIGHTPADDING',(0,0),(-1,-1),0)]))
    story = [header, Spacer(1,8*mm), text(record.get_kind_display()+' Report','Title'), text(record.get_status_display()+' · '+record.get_source_display()), Spacer(1,4*mm)]
    facts = [('Job',record.uut.job.job_code),('Client',record.uut.job.client.name),('PO',record.uut.job.purchase_order_number or 'Not recorded'),
             ('Equipment / serial',record.uut.name+' / '+record.uut.serial_number),('Scope',record.uut.scope.name if record.uut.scope_id else 'Not recorded'),
             ('Analyst / uploader',record.actor_name),('Template',str(record.template) if record.template_id else 'Uploaded / linked source'),
             ('Reviewer', record.reviewed_by.get_full_name() or record.reviewed_by.username if record.reviewed_by_id else 'Not reviewed'),
             ('Approver', record.approved_by.get_full_name() or record.approved_by.username if record.approved_by_id else 'Not approved')]
    def grid(rows, widths=None):
        table = Table([[text(value) for value in row] for row in rows],colWidths=widths,repeatRows=1,hAlign='LEFT',splitInRow=1)
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#edf4f7')),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#cfdee5')),
            ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
        story.append(table); story.append(Spacer(1,4*mm))
    grid([['Job / instrument details','Recorded value']]+facts,[width*.3,width*.7])
    values,tables,standards = report_data(record)
    if values:
        grid([['Recorded information / result','Value','Unit']]+values,[width*.4,width*.45,width*.15])
    if record.document_id:
        story.append(text('Retained original: '+record.document.original_name+' · uploaded by '+record.document.actor_name))
    for table in tables:
        story.append(text(table.get('label',table['name']),'Heading2'))
        rows = [[c.get('label',c['name'])+(' ('+c['unit']+')' if c.get('unit') else '') for c in table['columns']]]
        rows += [[row.get(c['name'],'') for c in table['columns']] for row in record.data.get(table['name'],[])]
        grid(rows,[width/len(table['columns'])]*len(table['columns']))
    if standards:
        story.append(text('Reference Standards and Traceability','Heading2'))
        grid([['Reference','Serial','Certificate','Calibration date','Due date']]+[[s[k] for k in ('name','serial','certificate','calibrated','due')] for s in standards],[width/5]*5)
    def footer(canvas, doc):
        canvas.saveState();canvas.setStrokeColor(colors.HexColor('#cfdee5'));canvas.line(15*mm,14*mm,A4[0]-15*mm,14*mm)
        canvas.setFont('Helvetica',7);canvas.setFillColor(colors.HexColor('#587184'));canvas.drawString(15*mm,10*mm,'EATL · E&I Laboratory · Controlled technical record')
        canvas.setFont('Helvetica-Oblique',6);canvas.drawRightString(A4[0]-15*mm,10*mm,f'FlowLab Pro · {doc.page}');canvas.restoreState()
    document = SimpleDocTemplate(output,pagesize=A4,leftMargin=15*mm,rightMargin=15*mm,topMargin=15*mm,bottomMargin=20*mm)
    document.build(story,onFirstPage=footer,onLaterPages=footer); return output.getvalue()


def retained_export(user, record, format):
    service.require_job_read(user,record.uut.job)
    if format not in ('pdf','xlsx'):
        raise ValueError('Choose PDF or Excel.')
    with transaction.atomic():
        record = TechnicalRecord.objects.select_for_update().select_related('uut__job','template').get(pk=record.pk)
        name = f'EI-REC-{record.pk}-Rev-{record.version}.{format}'
        existing = JobDocument.objects.filter(job=record.uut.job,uut=record.uut,kind='record_report',original_name=name).first()
        if existing:
            return existing
        raw = pdf_bytes(record) if format=='pdf' else xlsx_bytes(record)
        content = ContentFile(raw,name=name)
        mime,digest = service.sniff_upload(content)
        doc = JobDocument(job=record.uut.job,uut=record.uut,kind='record_report',title=name,original_name=name,mime_type=mime,sha256=digest,revision=record.version,**service.actor_fields(user))
        try:
            doc.file.save(name,content,save=False);doc.save()
            service.audit(user,record.uut.job,'report_exported',after={'record':record.pk,'version':record.version,'document':doc.pk,'format':format,'sha256':digest},uut=record.uut)
        except Exception:
            if doc.file.name:doc.file.storage.delete(doc.file.name)
            raise
        return doc
