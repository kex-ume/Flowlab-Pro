import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from werkzeug.datastructures import FileStorage

import app.modules.decision_rules as decision_rules


class DecisionRuleTests(unittest.TestCase):
    def sample_data(self):
        return {
            "assessment_uid": "DR-TEST-001", "assessment_id": 1, "status": "Approved",
            "job_id": 1, "job_number": "JOB-TEST-001", "customer_name": "Test Customer",
            "address": "1 Test Road", "city": "Lagos", "state": "Lagos", "country": "Nigeria",
            "meter_type": "Coriolis", "mut_description": "Flow meter", "manufacturer": "Test OEM",
            "model_number": "M-1", "mut_serial_number": "SN-1", "meter_size": "4 in",
            "flow_min": 500, "flow_max": 6000, "flow_unit": "L/min", "method_name": "Gravimetric",
            "service_request_number": "LSR-001", "service_request_revision": "04",
            "calculation_uid": "UB-001", "calculation_revision": 1,
            "rule_code": "DR-02", "rule_revision": "TBC", "specification_reference": "OEM specification",
            "specification_revision": "4", "specification_clause": "5.2", "tolerance_unit": "%",
            "statement_text": "All reported points conform.", "overall_decision": "PASS",
            "approved_by": "Chief Test", "approved_at": "2026-09-13",
            "points": [{"flow_value": 500, "flow_unit": "L/min", "error_percent": .12,
                "expanded_uncertainty": .08, "lower_tolerance": -.5, "upper_tolerance": .5,
                "lower_acceptance": -.42, "upper_acceptance": .42, "decision": "PASS"}],
        }

    def test_guarded_acceptance_uses_expanded_uncertainty(self):
        self.assertEqual(decision_rules.calculate_dr02_point(.25, .09, -.5, .5)["decision"], "PASS")
        result = decision_rules.calculate_dr02_point(.43, .10, -.5, .5)
        self.assertEqual(result["decision"], "FAIL")
        self.assertAlmostEqual(result["upper_acceptance"], .4)
        self.assertEqual(decision_rules.parse_flow_range("500 - 6,000 L/min"), (500.0, 6000.0, "L/min"))

    def test_service_request_uses_explicit_form_cells_not_adjacent_labels(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "REQUEST"
        sheet["A8"] = "Company Name:"
        sheet["A13"] = "Meter Type:"
        sheet["A16"] = "Meter Serial Number:"
        sheet["A18"] = "Meter OEM:"
        sheet["A21"] = "Required Number of Flowrates:"
        sheet["D21"] = "Required Number of Runs:"
        sheet["A31"] = "Specify Required flow range/flow rate:"
        sheet["A32"] = "Statement of Conformity Required?"
        sheet["A33"] = "If yes, state decision rule method:"
        sheet["A34"] = "If yes, state the acceptance limits:"
        sheet["D66"] = "Rev 02"
        sheet["B8"] = "Example Customer"
        sheet["E8"] = "Date: 14/09/2026"
        sheet["C13"] = "Coriolis"
        sheet["C16"] = "SN-100"
        sheet["C18"] = "Example OEM"
        sheet["D18"] = "Meter Model: X100"
        sheet["C21"] = 5
        sheet["E21"] = 3
        sheet["D31"] = "500 - 6,000 L/min"
        sheet["D32"] = "Yes"
        sheet["D33"] = "Option 2 guarded acceptance"
        sheet["D34"] = "±0.5%"
        stream = BytesIO(); workbook.save(stream); stream.seek(0)
        result = decision_rules.parse_service_request(FileStorage(stream=stream, filename="request.xlsx"))
        self.assertEqual(result["company_name"], "Example Customer")
        self.assertEqual(result["serial_number"], "SN-100")
        self.assertEqual(result["model_number"], "X100")
        self.assertEqual(result["flow_point_count"], 5)
        self.assertEqual(result["run_count"], 3)
        self.assertEqual((result["flow_min"], result["flow_max"], result["flow_unit"]), (500, 6000, "L/min"))
        self.assertEqual((result["lower_tolerance"], result["upper_tolerance"]), (-0.5, 0.5))

    def test_controlled_template_and_pdf_are_generated(self):
        original_directory = decision_rules._report_directory
        with tempfile.TemporaryDirectory() as directory:
            decision_rules._report_directory = lambda: Path(directory)
            try:
                template = Path("data/templates/report_cover.xlsx")
                excel = decision_rules.generate_excel_cover(self.sample_data(), template)
                pdf = decision_rules.generate_pdf_cover(self.sample_data())
                workbook = load_workbook(excel, data_only=True)
                sheet = workbook[workbook.sheetnames[0]]
                self.assertEqual(sheet["F4"].value, "JOB-TEST-001")
                self.assertEqual(sheet["B32"].value, "CALIBRATION RESULT AND CONFORMITY")
                self.assertGreater(pdf.stat().st_size, 1000)
            finally:
                decision_rules._report_directory = original_directory


if __name__ == "__main__":
    unittest.main()
